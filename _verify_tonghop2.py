# -*- coding: utf-8 -*-
import os, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl
FIN = os.path.join(os.path.dirname(os.path.abspath(__file__)), "finance_data")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_verify_tonghop2_out.txt")
L = []
def log(s=""):
    L.append(str(s)); open(OUT,"w",encoding="utf-8").write("\n".join(L)+"\n"); print(s)

# 在 TỔNG HỢP 全文找引用 Chi thực tế!E3 的单元格，以及含该大数的单元格
wb = openpyxl.load_workbook(os.path.join(FIN,"merge_report.xlsx"), data_only=False)
ws = wb["TỔNG HỢP"]
log("===== TỔNG HỢP 里引用 Chi thực tế!E3 的单元格（证明顶层含中方） =====")
for r in range(1, ws.max_row+1):
    for c in range(1, ws.max_column+1):
        v = ws.cell(r,c).value
        if isinstance(v,str) and "Chi thực tế'!E3" in v:
            log("  %s%d = %s" % (openpyxl.utils.get_column_letter(c), r, v))
log("")
log("===== TỔNG HỢP 列C 中出现的较大数值（找顶层合计） =====")
big = []
for r in range(1, ws.max_row+1):
    v = ws.cell(r,3).value
    if isinstance(v,(int,float)) and abs(v) > 1e9:
        big.append((r,v))
for r,v in big:
    log("  行%-3d C = %s" % (r, v))
wb.close()
log("done")
