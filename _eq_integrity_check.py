#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""完整性体检：equity_report.xlsx / cn_report.xlsx / vn_report.xlsx 是否是【能正常读的 zip】。

背景：服务器日志出现
  [fin-datefmt] 打不开 ...equity_report.xlsx.tmp.xlsx: Bad CRC-32 for file 'docProps/core.xml'
  [fin-keep] cell stats failed (...equity_report.xlsx): Bad CRC-32 for file 'docProps/core.xml'
必须确认现网文件没被写坏，并对比我留的手工备份。
"""
import os
import sys
import zipfile

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(ROOT, "finance_data")
BK = os.path.join(FIN, "backups")

TARGETS = [
    ("equity_report.xlsx", os.path.join(FIN, "equity_report.xlsx")),
    ("eq 手工备份", os.path.join(BK, "eq_report_before_finkeep_20261016.xlsx.bak")),
    ("cn_report.xlsx", os.path.join(FIN, "cn_report.xlsx")),
    ("vn_report.xlsx", os.path.join(FIN, "vn_report.xlsx")),
    ("merge_report.xlsx", os.path.join(FIN, "merge_report.xlsx")),
]


def report(name, p):
    if not os.path.isfile(p):
        print("  %-18s (不存在)" % name)
        return
    st = os.stat(p)
    print("  %-18s %7d B  %s" % (name, st.st_size,
                                 __import__("datetime").datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S")))
    try:
        with zipfile.ZipFile(p) as z:
            bad = z.testzip()
            names = z.namelist()
        if bad:
            print("      ✗ ZIP 损坏，第一个坏条目: %s" % bad)
        else:
            print("      ✓ ZIP 完好（%d 个条目）" % len(names))
    except Exception as e:
        print("      ✗ 不是有效 zip: %s" % e)
        return
    try:
        import openpyxl
        wb = openpyxl.load_workbook(p, data_only=False)
        info = []
        for ws in wb.worksheets:
            nz = 0
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, 3000),
                                    max_col=min(ws.max_column or 0, 80)):
                for c in row:
                    if c.value not in (None, ""):
                        nz += 1
            info.append("%s(%d格)" % (ws.title, nz))
        wb.close()
        print("      ✓ openpyxl 可读: " + ", ".join(info[:8]))
    except Exception as e:
        print("      ✗ openpyxl 读不了: %s" % e)


print("=" * 78)
for n, p in TARGETS:
    report(n, p)
print("-" * 78)
# 临时/残留文件
for fn in sorted(os.listdir(FIN)):
    if fn.endswith(".tmp") or ".tmp." in fn:
        print("  残留临时文件:", fn, os.path.getsize(os.path.join(FIN, fn)), "B")
hd = os.path.join(BK, "eq_history")
if os.path.isdir(hd):
    fs = sorted(os.listdir(hd))
    print("  eq_history 共 %d 份，最近 3 份: %s" % (len(fs), fs[-3:]))
print("=" * 78)
