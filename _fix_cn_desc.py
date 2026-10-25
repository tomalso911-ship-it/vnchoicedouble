# -*- coding: utf-8 -*-
"""修复 cn_report.xlsx「实际支出」：
   1) 摘要列（C=VI / M=ZH / W=EN）被写成数字 0 的数据行 -> 回填摘要
      来源优先级：① 最新备份里的原文 ② vn_report.xlsx 同笔对应语言块
   2) 空模板行里残留的孤立 0 摘要 -> 清空
   3) 数据行的 ZH/EN 块日期被写成字符串 -> 转成与上方行一致的序列号 + 同一数字格式
   金额列（D/E/N/O/X/Y）与核销列（AC）绝对不碰。
   落盘：先备份 -> 写 .fix.tmp -> os.replace 原子替换。
"""
import datetime
import glob
import os
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
CN = os.path.join(FIN, "cn_report.xlsx")
VN = os.path.join(FIN, "vn_report.xlsx")
BKD = os.path.join(FIN, "backups")
OUT = os.path.join(HERE, "_fix_cn_desc_out.txt")
VNSHEET = "Chi th\u1ef1c t\u1ebf"

DESC = ((3, "VI"), (13, "ZH"), (23, "EN"))
DATE_COLS = (1, 11, 21)
AMT_COLS = (5, 15, 25)          # 支出列 E/O/Y（收入列不参与对账）
OTHER_COLS = (1, 2, 4, 5, 11, 12, 14, 15, 21, 22, 24, 25)

L = []


def log(s=""):
    L.append(str(s))
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    print(s)


def usable(v):
    if v is None:
        return False
    if isinstance(v, str):
        t = v.strip()
        return bool(t) and t != "0"
    return False


def dstr(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        return v.strip()
    return ""


def fmt(v):
    if v is None:
        return "-"
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


# ---------- 1. 备份候选（新 -> 旧） ----------
bks = []
for pat in ("cn_report*.bak", "cn_report*.xlsx.bak", "cn_report*.xlsx"):
    bks += glob.glob(os.path.join(BKD, "**", pat), recursive=True)
bks = [p for p in bks if os.path.abspath(p) != os.path.abspath(CN)]
bks = sorted(set(bks), key=lambda p: os.path.getmtime(p), reverse=True)
log("=== 备份候选 %d 个（新 -> 旧） ===" % len(bks))
for p in bks[:10]:
    log("   %s   %s" % (datetime.datetime.fromtimestamp(os.path.getmtime(p)).strftime("%m-%d %H:%M"), os.path.basename(p)))
log("")

_bk_cache = {}


def bk_cell(path, r, c, date_s=None, voucher=None):
    """取备份单元格；若给了 (date_s, voucher)，则要求该备份里同行的日期/凭证完全一致，
       避免踩到历史上表结构不同的备份（行号错位会取到别人的文字）。"""
    wb = _bk_cache.get(path, "NA")
    if wb == "NA":
        try:
            wb = openpyxl.load_workbook(path, data_only=False)
        except Exception:
            wb = False
        _bk_cache[path] = wb
    if not wb:
        return None
    ws = wb[wb.sheetnames[0]]
    if date_s is not None and voucher is not None:
        if dstr(ws.cell(r, 1).value) != date_s:
            return None
        if str(ws.cell(r, 2).value or "").strip().upper() != str(voucher).strip().upper():
            return None
    return ws.cell(r, c).value


# ---------- 2. 装载 ----------
wbc = openpyxl.load_workbook(CN, data_only=False)
wsc = wbc[wbc.sheetnames[0]]
wbv = openpyxl.load_workbook(VN, data_only=False)
wsv = wbv[VNSHEET]

vin = {}
for r in range(5, 420):
    v = str(wsv.cell(r, 2).value or "").strip().upper()
    if v:
        vin.setdefault(v, []).append(r)


def vn_match(date_s, voucher, amount, date_obj):
    rows = vin.get(voucher.upper(), [])
    if not rows:
        return None, ""
    exact = [r for r in rows if dstr(wsv.cell(r, 1).value) == date_s]
    if exact:
        return exact[0], "日期+凭证一致"
    same_amt = [r for r in rows if wsv.cell(r, 5).value == amount]
    if same_amt:
        return same_amt[0], "仅凭证+金额一致"
    if date_obj:
        rows = sorted(rows, key=lambda r: abs((wsv.cell(r, 1).value - date_obj).days)
                      if isinstance(wsv.cell(r, 1).value, datetime.datetime) else 99999)
    return rows[0], "仅凭证一致(取日期最近)"


def amt_total(ws):
    t = 0
    for r in range(5, 205):
        for c in AMT_COLS:
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                t += v
    return t


total_before = amt_total(wsc)
log("=== 修改前：支出合计(E/O/Y, 行5-204) = %s ===" % format(total_before, ","))

data_rows = [r for r in range(5, 205) if wsc.cell(r, 2).value is not None]
log("数据行 = %s" % data_rows)
log("")

# ---------- 3. 回填摘要 ----------
log("=== 摘要回填 ===")
changed = []
for r in data_rows:
    d_obj = wsc.cell(r, 1).value
    date_s = dstr(d_obj)
    voucher = str(wsc.cell(r, 2).value or "").strip()
    amount = wsc.cell(r, 5).value
    if not isinstance(d_obj, (datetime.datetime, datetime.date)):
        d_obj = None
    vr, why = vn_match(date_s, voucher, amount, d_obj)
    log("行%-3d %s %-9s 金额=%s  -> vn 行 %s (%s)" % (r, date_s, voucher, fmt(amount), vr, why or "无同笔"))
    for c, lang in DESC:
        cur = wsc.cell(r, c).value
        if usable(cur):
            log("      %s 摘要已有：%s" % (lang, str(cur)[:40]))
            continue
        text = None
        src = ""
        for p in bks:                      # ① 备份原文（要求同行日期+凭证一致）
            v = bk_cell(p, r, c, date_s, voucher)
            if usable(v):
                text = str(v).strip()
                src = "备份 " + os.path.basename(p)
                break
        if text is None and vr is not None:  # ② vn 同笔同语言块
            v = wsv.cell(vr, c).value
            if usable(v):
                text = str(v).strip()
                src = "vn 行%d" % vr
        if text is None:
            log("      %s 摘要：当前=%s，备份/vn 都取不到 -> 不动" % (lang, fmt(cur)))
            continue
        wsc.cell(r, c).value = text
        wsc.cell(r, c).number_format = "General"
        changed.append((r, lang, text, src))
        log("      %s 摘要：%s -> %s   [%s]" % (lang, fmt(cur), text, src))
log("")

# ---------- 4. 数据行 ZH/EN 日期字符串 -> 序列号 ----------
log("=== 数据行 ZH/EN 日期格式统一 ===")
for r in data_rows:
    for c in (11, 21):
        v = wsc.cell(r, c).value
        if isinstance(v, str) and v.strip():
            try:
                d = datetime.datetime.strptime(v.strip(), "%Y-%m-%d")
            except Exception:
                continue
            serial = (d - datetime.datetime(1899, 12, 30)).days
            ref = wsc.cell(5, c)
            wsc.cell(r, c).value = serial
            if ref.number_format:
                wsc.cell(r, c).number_format = ref.number_format
            log("   行%d 列%d  %s -> %d（格式=%s）" % (r, c, v, serial, ref.number_format))
log("")

# ---------- 5. 空模板行里残留的孤立 0 ----------
log("=== 空行残留 0 清理 ===")
cleared = []
for r in range(16, 205):
    if any(wsc.cell(r, c).value is not None for c in OTHER_COLS):
        continue
    hit = []
    for c, lang in DESC:
        v = wsc.cell(r, c).value
        if v is not None and (v == 0 or (isinstance(v, str) and v.strip() == "0")):
            wsc.cell(r, c).value = None
            hit.append(lang)
            cleared.append((r, lang))
    if hit:
        log("   行%-3d 清掉残留 0：%s" % (r, "/".join(hit)))
if not cleared:
    log("   无")
log("")

# ---------- 6. 备份 + 原子落盘 ----------
ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
bak = os.path.join(BKD, "cn_report_%s_before_descbackfill.xlsx.bak" % ts)
shutil.copy2(CN, bak)
log("=== 已备份原文件 -> %s ===" % bak)

tmp = CN + ".fix.tmp"
wbc.save(tmp)
wbc.close()
wbv.close()
os.replace(tmp, CN)
log("=== 已原子落盘：%s ===" % CN)

# ---------- 7. 复核 ----------
wb2 = openpyxl.load_workbook(CN, data_only=False)
ws2 = wb2[wb2.sheetnames[0]]
total_after = amt_total(ws2)
log("")
log("=== 复核 ===")
log("支出合计 改前=%s  改后=%s  %s" % (format(total_before, ","), format(total_after, ","),
                                     "一致 ✓" if total_before == total_after else "!!! 不一致 !!!"))
log("行3 合计公式 E = %s" % ws2.cell(3, 5).value)
for r in [r for r in range(5, 205) if ws2.cell(r, 2).value is not None]:
    log("行%-3d %s %-9s | VI=%s" % (r, dstr(ws2.cell(r, 1).value), ws2.cell(r, 2).value, str(ws2.cell(r, 3).value)[:34]))
for c, lang in DESC:
    bad = [r for r in range(5, 205) if ws2.cell(r, c).value is not None and not usable(ws2.cell(r, c).value)]
    bad = [r for r in bad if ws2.cell(r, 2).value is not None]
    log("%s 摘要仍无效的行 = %s" % (lang, bad or "无 ✓"))
log("")
log("共回填 %d 格，清理 %d 格" % (len(changed), len(cleared)))
log("done")
