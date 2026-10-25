# -*- coding: utf-8 -*-
"""逐格比较 equity_report.xlsx 与历史备份，确认数据没被验收粘贴改脏。"""
import shutil
import sys
import tempfile

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

CUR = 'finance_data/equity_report.xlsx'
BAKS = [
    'finance_data/backups/eq_history/eq_20261018_125846_653_presave.xlsx.bak',
    'finance_data/backups/eq_history/eq_20261018_130308_309_presave.xlsx.bak',
    'finance_data/backups/eq_history/eq_20261018_130311_128_presave.xlsx.bak',
]


def values(path):
    wb = load_workbook(path, data_only=False)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        grid = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    grid[(c.row, c.column)] = str(c.value)
        out[name] = grid
    return out


cur = values(CUR)
for b in BAKS:
    tmp = tempfile.mktemp(suffix='.xlsx')
    shutil.copyfile(b, tmp)
    try:
        old = values(tmp)
    except Exception as e:
        print(b, 'ERR', e)
        continue
    diffs = []
    for name in set(list(cur.keys()) + list(old.keys())):
        g1, g2 = cur.get(name, {}), old.get(name, {})
        for k in set(list(g1.keys()) + list(g2.keys())):
            if g1.get(k) != g2.get(k):
                diffs.append((name, k, g2.get(k), g1.get(k)))
    print(b.split('/')[-1], '-> diffs =', len(diffs))
    for d in diffs[:10]:
        print('   ', d)
