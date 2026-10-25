# -*- coding: utf-8 -*-
"""根因排查：上传的 vn_report.xlsx 为什么没有被三语化识别/翻译。
复刻前端 fin_isTrilingualSheet() 的判定规则：
  头 20 行内，对 c in 0..7，若 row[c] 与 row[c+10] 都非空 → hits++；hits>=3 判为三语并列表。
"""
import os
import sys
import json
import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(BASE, "finance_data", "vn_report.xlsx")
TITLES = os.path.join(BASE, "finance_data", "fin_sheet_titles.json")
CACHE = os.path.join(BASE, "finance_data", "fin_trans_cache.json")

st = os.stat(FIN)
print("=== 文件 ===")
print("vn_report.xlsx  size=%d  mtime=%s" % (st.st_size, __import__("datetime").datetime.fromtimestamp(st.st_mtime)))
for p in (TITLES, CACHE):
    if os.path.isfile(p):
        s2 = os.stat(p)
        print("%-28s size=%d mtime=%s" % (os.path.basename(p), s2.st_size,
              __import__("datetime").datetime.fromtimestamp(s2.st_mtime)))
    else:
        print("%-28s <缺失>" % os.path.basename(p))

print()
print("=== 表名映射（fin_sheet_titles.json）===")
try:
    titles = json.load(open(TITLES, encoding="utf-8"))
    for k, v in titles.items():
        print("  %r -> %s" % (k, v))
except Exception as e:
    print("  读取失败:", e)
    titles = {}

wb = openpyxl.load_workbook(FIN, data_only=False)
print()
print("=== 各表判定（复刻 fin_isTrilingualSheet）===")
for ws in wb.worksheets:
    hits = 0
    hits_all = 0
    for r in range(0, min(20, ws.max_row)):
        for c in range(0, 8):
            a = ws.cell(row=r + 1, column=c + 1).value
            b = ws.cell(row=r + 1, column=c + 11).value
            if a not in (None, "") and b not in (None, ""):
                hits += 1
    for r in range(0, min(200, ws.max_row)):
        for c in range(0, 8):
            a = ws.cell(row=r + 1, column=c + 1).value
            b = ws.cell(row=r + 1, column=c + 11).value
            if a not in (None, "") and b not in (None, ""):
                hits_all += 1
    tri = hits >= 3
    in_titles = ws.title in titles
    print("  %-24r max_col=%d  前20行配对hits=%d (前200行=%d)  三语判定=%-5s 表名在映射表=%s"
          % (ws.title, ws.max_column, hits, hits_all, tri, in_titles))

print()
print("=== 前 6 行原始内容（A..AB = 列1..28）===")
for ws in wb.worksheets[:2]:
    print("--- %r" % ws.title)
    for r in range(1, min(7, ws.max_row + 1)):
        vals = []
        for c in range(1, 29):
            v = ws.cell(row=r, column=c).value
            vals.append("" if v is None else str(v)[:18])
        print("  r%-2d | %s" % (r, " | ".join(vals)))
    print()

print("=== fin_trans_cache.json 条目数 ===")
try:
    tc = json.load(open(CACHE, encoding="utf-8"))
    print("  %d 条" % len(tc))
    ks = list(tc.keys())[:5]
    for k in ks:
        print("   %r -> %r" % (k, tc[k]))
except Exception as e:
    print("  读取失败:", e)
