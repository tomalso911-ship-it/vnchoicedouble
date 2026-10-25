# -*- coding: utf-8 -*-
"""对比新旧文件 汇总 表：A..H 源块 与 K..R 中文块 到底填了什么。"""
import os
import sys
import shutil
import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
FD = os.path.join(BASE, "finance_data")


def load(p):
    if not p.lower().endswith(".xlsx"):
        t = os.path.join(BASE, "_tmp_d.xlsx")
        shutil.copyfile(p, t)
        p = t
    return openpyxl.load_workbook(p)


def dump(path, label, sheet_idx=0):
    print("=" * 100)
    print(label, "|", os.path.basename(path))
    wb = load(path)
    ws = wb.worksheets[sheet_idx]
    print("sheet=%r rows=%d cols=%d" % (ws.title, ws.max_row, ws.max_column))
    for r in range(1, min(12, ws.max_row) + 1):
        src = []
        for c in range(1, 9):
            v = ws.cell(r, c).value
            src.append("." if v in (None, "") else str(v)[:20])
        zh = []
        for c in range(11, 19):
            v = ws.cell(r, c).value
            zh.append("." if v in (None, "") else str(v)[:20])
        print("  r%-2d A-H[%s]" % (r, " | ".join(src)))
        print("       K-R[%s]" % (" | ".join(zh)))


dump(os.path.join(FD, "vn_report.xlsx.bak_before_postfix"), "[旧-昨天可用]")
dump(os.path.join(FD, "vn_report.xlsx"), "[新-现在]")
