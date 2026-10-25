#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""诊断：股东权益「中国投资数据」表头到底是什么语言 + 服务端中译越是否可用。"""
import json
import os
import sys
import urllib.parse
import urllib.request

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = os.path.dirname(os.path.abspath(__file__))
P = os.path.join(ROOT, "finance_data", "equity_report.xlsx")

wb = openpyxl.load_workbook(P, data_only=False)
for ws in wb.worksheets:
    if ws.title != "Đầu tư Trung Quốc":
        continue
    print("工作表:", ws.title, " dims:", ws.dimensions)
    for r in range(2, 5):
        vals = []
        for c in range(1, 14):
            v = ws.cell(r, c).value
            if v not in (None, ""):
                vals.append("%s%d=%r" % (ws.cell(r, c).column_letter, r, v))
        print("  行%d: %s" % (r, " | ".join(vals)))
wb.close()

print("-" * 60)
for q in ("序号", "年份", "公司", "款项来自", "款项去处", "日期", "现金/转账", "备注", "金额", "汇率"):
    u = "http://127.0.0.1:5050/api/translate?q=%s&target=vi" % urllib.parse.quote(q)
    try:
        with urllib.request.urlopen(u, timeout=25) as r:
            print("  %-10s -> %s" % (q, r.read().decode("utf-8")[:160]))
    except Exception as e:
        print("  %-10s ERR %s" % (q, e))
print("=" * 60)
