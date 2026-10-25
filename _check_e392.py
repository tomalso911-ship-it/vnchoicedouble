# -*- coding: utf-8 -*-
import os, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl
FIN = os.path.join(os.path.dirname(os.path.abspath(__file__)), "finance_data")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_check_e392_out.txt")
L = []
def log(s=""):
    L.append(str(s)); open(OUT,"w",encoding="utf-8").write("\n".join(L)+"\n"); print(s)

p = os.path.join(FIN,"merge_report.xlsx")
wbf = openpyxl.load_workbook(p, data_only=False)["Chi thực tế"]
wbd = openpyxl.load_workbook(p, data_only=True)["Chi thực tế"]
tot = 0.0; big = []
for r in range(5, 393):
    v = wbf.cell(r,5).value
    d = wbd.cell(r,5).value
    val = d if isinstance(d,(int,float)) else (v if isinstance(v,(int,float)) else None)
    if isinstance(val,(int,float)):
        tot += val
    if isinstance(val,(int,float)) and abs(val) > 1e8:
        big.append((r, val, wbf.cell(r,3).value))
log("merge E5:E392 求和 = %s" % tot)
log("区间内 >1e8 的单元格: %s" % big)
log("E3 (表头合计) = %r" % wbf.cell(3,5).value)
log("done")
