# -*- coding: utf-8 -*-
"""看那份 cn 每日备份里到底是【测试数据】还是【用户真实数据】。"""
import json
import os
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
src = os.path.join(HERE, "finance_data", "backups", "cn_report_2026-10-13.xlsx.bak")
tmp = os.path.join(HERE, "_tmp_bak.xlsx")
shutil.copyfile(src, tmp)
try:
    wb = openpyxl.load_workbook(tmp)
    ws = wb[wb.sheetnames[0]]
    print("sheets =", json.dumps(wb.sheetnames, ensure_ascii=True))
    print("dims   =", ws.dimensions)
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=12, max_col=6):
        vals = [c.value for c in row]
        if any(v not in (None, "") for v in vals):
            rows.append(vals)
    print("rows   =", json.dumps(rows, ensure_ascii=True, default=str))
    ac = [ws["AC%d" % r].value for r in (4, 5, 6, 7)]
    print("AC4..7 =", json.dumps(ac, ensure_ascii=True, default=str))
finally:
    try:
        os.remove(tmp)
    except Exception:
        pass
