# -*- coding: utf-8 -*-
"""对比 equity_report.xlsx 与备份，确认验收粘贴有没有写脏数据。"""
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

CUR = 'finance_data/equity_report.xlsx'
BAKS = [
    'finance_data/backups/eq_history/eq_20261018_125846_653_presave.xlsx.bak',
    'finance_data/backups/eq_history/eq_20261018_130308_309_presave.xlsx.bak',
    'finance_data/backups/eq_history/eq_20261018_130311_128_presave.xlsx.bak',
]


def dump(path, label):
    try:
        wb = load_workbook(path, data_only=False)
    except Exception as e:
        print(label, 'LOAD_ERR', e)
        return None
    print('===', label, 'sheets=', wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    print('  dims', ws.max_row, 'x', ws.max_column)
    for r in (1, 2, 3, 4, 21, 22):
        vals = [ws.cell(row=r, column=c).value for c in range(1, 16)]
        print('  r%d' % r, vals)
    return ws


cur = dump(CUR, 'CURRENT')
for b in BAKS:
    dump(b, b.split('/')[-1])
