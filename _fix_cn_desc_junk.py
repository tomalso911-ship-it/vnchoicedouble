# -*- coding: utf-8 -*-
"""CN 摘要列出现「垃圾值」（0 / #VALUE! / 空）时自动修复：
   数据行（有日期或凭证）→ 从备份原文 / vn 同笔三语块恢复；
   空模板行 → 直接清成空。
   金额、日期、凭证一律不碰。备份 + 原子落盘 + 复核。
"""
import datetime
import os
import shutil
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
CN = os.path.join(FIN, "cn_report.xlsx")
VN = os.path.join(FIN, "vn_report.xlsx")
BKD = os.path.join(FIN, "backups")
SRC_BAK = os.path.join(BKD, "cn_report_20261014_095127_before_datefmt.xlsx.bak")
OUT = os.path.join(HERE, "_fix_cn_desc_junk_out.txt")
DESC = ((3, "VI"), (13, "ZH"), (23, "EN"))
JUNK = ("", "0", "#VALUE!", "#N/A", "#REF!", "#DIV/0!", "None", "nan")
L = []


def log(s=""):
    L.append(str(s))
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    print(s)


def fmt(v):
    if v is None:
        return "-"
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


def dstr(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        return v.strip()
    return ""


def usable(v):
    if not isinstance(v, str):
        return False
    t = v.strip()
    return bool(t) and t not in JUNK


def is_junk(v):
    if v is None:
        return False                      # 本来就是空 -> 不算垃圾，不主动填
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v) == 0
    if isinstance(v, str):
        return v.strip() in JUNK
    return False


def as_xlsx(p):
    t = os.path.join(tempfile.gettempdir(), os.path.basename(p) + ".junk.xlsx")
    shutil.copy2(p, t)
    return t


wbc = openpyxl.load_workbook(CN, data_only=False)
wsc = wbc[wbc.sheetnames[0]]
src = openpyxl.load_workbook(as_xlsx(SRC_BAK), data_only=False)
wss = src[src.sheetnames[0]]
wbv = openpyxl.load_workbook(VN, data_only=False)
wsv = wbv["Chi th\u1ef1c t\u1ebf"]

vin = {}
for r in range(5, 420):
    v = str(wsv.cell(r, 2).value or "").strip().upper()
    if v:
        vin.setdefault(v, []).append(r)


def vn_rows(date_s, voucher, amount, date_obj):
    rows = vin.get(voucher.upper(), [])
    if not rows:
        return None
    exact = [r for r in rows if dstr(wsv.cell(r, 1).value) == date_s]
    if exact:
        return exact[0]
    same = [r for r in rows if wsv.cell(r, 5).value == amount]
    if same:
        return same[0]
    if date_obj:
        rows = sorted(rows, key=lambda r: abs((wsv.cell(r, 1).value - date_obj).days)
                      if isinstance(wsv.cell(r, 1).value, datetime.datetime) else 99999)
    return rows[0]


log("=== 扫描摘要列垃圾值 ===")
fixes = []
clears = []
for r in range(1, 210):
    a, b = wsc.cell(r, 1).value, wsc.cell(r, 2).value
    is_data = (a is not None or b is not None) and r >= 5
    for c, lang in DESC:
        cur = wsc.cell(r, c).value
        if not is_junk(cur):
            if cur is not None and not isinstance(cur, str):
                pass
            continue
        if not is_data:
            clears.append((r, c, lang, cur))
            continue
        # 数据行：找原文
        text, src_tag = None, ""
        v = wss.cell(r, c).value
        if usable(v):
            text, src_tag = v.strip(), "备份 095127"
        else:
            date_s = dstr(a)
            vr = vn_rows(date_s, str(b or ""), wsc.cell(r, 5).value,
                         a if isinstance(a, datetime.datetime) else None)
            if vr is not None and usable(wsv.cell(vr, c).value):
                text, src_tag = str(wsv.cell(vr, c).value).strip(), "vn 行%d" % vr
        if text is None:
            log("   行%-3d %s 摘要=%s，备份/vn 都取不到 -> 清空" % (r, lang, fmt(cur)))
            clears.append((r, c, lang, cur))
        else:
            fixes.append((r, c, lang, cur, text, src_tag))
            log("   行%-3d %s 摘要 %s -> %s   [%s]" % (r, lang, fmt(cur), text, src_tag))

log("")
log("数据行待恢复 %d 格 / 待清空 %d 格" % (len(fixes), len(clears)))
if not fixes and not clears:
    log("没有需要修的，未落盘。")
    sys.exit(0)

ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
bak = os.path.join(BKD, "cn_report_%s_before_descjunkfix.xlsx.bak" % ts)
shutil.copy2(CN, bak)
log("已备份 -> %s" % bak)

for r, c, _l, _cur, text, _tag in fixes:
    wsc.cell(r, c).value = text
    wsc.cell(r, c).number_format = "General"
for r, c, _l, _cur in clears:
    wsc.cell(r, c).value = None

tmp = CN + ".fix.tmp"
wbc.save(tmp)
wbc.close()
src.close()
wbv.close()
os.replace(tmp, CN)
log("已原子落盘：%s" % CN)

wb2 = openpyxl.load_workbook(CN, data_only=False)
ws2 = wb2[wb2.sheetnames[0]]
log("")
log("=== 复核：现在所有数据行 ===")
tot = 0
for r in range(5, 210):
    if ws2.cell(r, 2).value is None and ws2.cell(r, 1).value is None:
        continue
    log("行%-3d %s | %s | %s" % (r, dstr(ws2.cell(r, 1).value), fmt(ws2.cell(r, 2).value),
                                 str(ws2.cell(r, 3).value)[:34]))
    for c in (5, 15, 25):
        v = ws2.cell(r, c).value
        if isinstance(v, (int, float)):
            tot += v
log("支出合计(三块相加) = %s   （每块 %s）" % (format(tot, ","), format(tot // 3, ",")))
bad = [(r, lang) for r in range(5, 210) for c, lang in DESC if is_junk(ws2.cell(r, c).value)]
log("仍有垃圾摘要的数据行 = %s" % (bad or "无 ✓"))
wb2.close()
log("done")
