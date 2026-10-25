# -*- coding: utf-8 -*-
"""只读：对比 vn 与 merge 在 实际支出合计 和 TỔNG HỢP 总表 的实际数值。
   看合并报表的「实际支出」是否真的比越南报表大，且总表是否跟着变大。"""
import os, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl
FIN = os.path.join(os.path.dirname(os.path.abspath(__file__)), "finance_data")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_audit_tonghop_diff_out.txt")
L = []
def log(s=""):
    L.append(str(s)); open(OUT,"w",encoding="utf-8").write("\n".join(L)+"\n"); print(s)

def load(fn):
    return openpyxl.load_workbook(os.path.join(FIN, fn), data_only=True)

def chi_total(wb):
    ws = wb["Chi thực tế"]
    # 找静态合计行：E 列最大数值（通常是 E3 / 末尾合计行）
    best = 0; bestr = 0
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 5).value
        if isinstance(v, (int, float)) and abs(v) > abs(best):
            best = v; bestr = r
    return best, bestr

for fn in ("vn_report.xlsx", "merge_report.xlsx"):
    wb = load(fn)
    tot, r = chi_total(wb)
    log("%-18s Chi thực tế 实际支出合计(E列最大) = %s  (行%d)" % (fn, tot, r))
    wb.close()

log("")
log("===== TỔNG HỢP 全部数值单元格（vn vs merge） =====")
wv = load("vn_report.xlsx")["TỔNG HỢP"]
wm = load("merge_report.xlsx")["TỔNG HỢP"]
rows = max(wv.max_row, wm.max_row)
diff = 0
for r in range(1, rows + 1):
    for c in range(1, max(wv.max_column, wm.max_column) + 1):
        vv = wv.cell(r, c).value
        vm = wm.cell(r, c).value
        nv = vv if isinstance(vv, (int, float)) else (vv if isinstance(vv, str) else None)
        nm = vm if isinstance(vm, (int, float)) else (vm if isinstance(vm, str) else None)
        if isinstance(vv,(int,float)) or isinstance(vm,(int,float)):
            if vv != vm:
                lab = wv.cell(r,2).value or wm.cell(r,2).value
                log("  行%-3d 列%s  标签=%r  vn=%s  merge=%s" % (r, openpyxl.utils.get_column_letter(c), lab, vv, vm))
                diff += 1
log("TỔNG HỢP 数值差异数 = %d" % diff)
wv.parent.close(); wm.parent.close()
log("done")
