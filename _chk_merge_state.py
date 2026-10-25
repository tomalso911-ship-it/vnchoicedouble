# -*- coding: utf-8 -*-
"""只读检查：vn_report.xlsx「Chi thực tế」当前是否已含合并进来的中国记录。"""
import io
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
VN = os.path.join(HERE, "finance_data", "vn_report.xlsx")
SHEET = "Chi th\u1ef1c t\u1ebf"
MARK_COL = 40   # AN
OUT = os.path.join(HERE, "_chk_merge_state_out.txt")
L = []


def log(s):
    L.append(str(s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(s)


wb = openpyxl.load_workbook(VN, data_only=False)
log("sheets = %s" % wb.sheetnames)
ws = wb[SHEET]

cn_rows = []
last_data = 0
for r in range(5, 400):
    vals = [ws.cell(r, c).value for c in (1, 2, 3, 4, 5, 11, 14, 21, 24)]
    nonempty = any(v is not None and str(v).strip() != "" for v in vals)
    if nonempty:
        last_data = r
    mk = ws.cell(r, MARK_COL).value
    if isinstance(mk, str) and mk.startswith("CN:"):
        font = ws.cell(r, 1).font
        col = None
        try:
            col = font.color.rgb if (font and font.color) else None
        except Exception:
            col = "?"
        cn_rows.append((r, mk, ws.cell(r, 1).value, ws.cell(r, 2).value,
                        ws.cell(r, 5).value, col))

log("最后一条有数据的行 = %d" % last_data)
log("CN 标记行数 = %d" % len(cn_rows))
for x in cn_rows:
    log("   行%-4d %-22s 日期=%-12s 凭证=%-10s E=%-12s 字色=%s"
        % (x[0], x[1], str(x[2])[:12], str(x[3]), str(x[4]), x[5]))

log("行3 合计 : D=%s E=%s N=%s O=%s" % (ws.cell(3, 4).value, ws.cell(3, 5).value,
                                        ws.cell(3, 14).value, ws.cell(3, 15).value))
log("done")
