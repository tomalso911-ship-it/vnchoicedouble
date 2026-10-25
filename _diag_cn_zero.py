# -*- coding: utf-8 -*-
"""诊断中国财务报表「金额列出现 ¥0」：
   把 cn_report.xlsx 与最近的备份逐格对比，找出被写成 0 的格子与写入范围。
   关注列（三语块的金额列）：D E / N O / X Y（1 基 4 5 / 14 15 / 24 25）。
"""
import glob
import io
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
BAK = os.path.join(FIN, "backups")
CN = os.path.join(FIN, "cn_report.xlsx")
OUT = os.path.join(HERE, "_diag_cn_zero_out.txt")
L = []

MONEY_COLS = [(4, "D"), (5, "E"), (14, "N"), (15, "O"), (24, "X"), (25, "Y")]
NAME_COLS = [(2, "B"), (3, "C"), (12, "L"), (13, "M"), (22, "V"), (23, "W")]


def log(s):
    L.append(str(s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(s)


def probe(path):
    p = path
    tmp = None
    if not path.lower().endswith(".xlsx"):
        tmp = path + ".probe.xlsx"
        import shutil
        shutil.copyfile(path, tmp)
        p = tmp
    wb = openpyxl.load_workbook(p, data_only=False)
    ws = wb[wb.sheetnames[0]]
    rows = {}
    maxr = min(ws.max_row or 0, 260)
    for r in range(1, maxr + 1):
        vals = {}
        for c, _ in MONEY_COLS + NAME_COLS + [(1, "A"), (29, "AC")]:
            vals[c] = ws.cell(r, c).value
        rows[r] = vals
    wb.close()
    if tmp:
        try:
            os.remove(tmp)
        except OSError:
            pass
    return rows, maxr


log("=== 当前 cn_report.xlsx ===")
cur, maxr = probe(CN)
zero_cells = []
for r in range(5, maxr + 1):
    for c, nm in MONEY_COLS:
        v = cur[r][c]
        if v is None:
            continue
        if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v) == 0.0:
            zero_cells.append((r, nm, v))
log("金额列中【等于 0】的格子数 = %d" % len(zero_cells))
log("按列统计：%s" % {nm: sum(1 for x in zero_cells if x[1] == nm) for _, nm in MONEY_COLS})

log("")
log("前 40 个 0 格（含同行 A 日期 / B 凭证 / C 摘要）：")
for r, nm, v in zero_cells[:40]:
    d = cur[r][1]
    b = cur[r][2]
    cc = cur[r][3]
    log("   行%-4d %s=%-6s | A=%-12s B=%-10s C=%s"
        % (r, nm, v, str(d)[:12], str(b)[:10], str(cc)[:30]))

log("")
log("整行（A..E 全空）但金额列有 0 的行：")
empty_rows = []
for r in range(5, maxr + 1):
    base = [cur[r][c] for c in (1, 2, 3, 4, 5)]
    if all(x is None or str(x).strip() == "" for x in base):
        if any(isinstance(cur[r][c], (int, float)) and not isinstance(cur[r][c], bool)
               and float(cur[r][c]) == 0.0 for c, _ in MONEY_COLS):
            empty_rows.append(r)
log("   行号：%s（共 %d 行）" % (empty_rows[:60], len(empty_rows)))

# ---------- 与备份对比 ----------
baks = sorted(glob.glob(os.path.join(BAK, "cn_report_*.xlsx.bak")))
log("")
log("=== 备份（最近 6 份）对比：被写成 0 的格子，在备份里是什么 ===")
for bk in baks[-6:]:
    try:
        old, om = probe(bk)
    except Exception as e:
        log("   %s 读取失败：%s" % (os.path.basename(bk), e))
        continue
    diff = []
    for r, nm, v in zero_cells:
        if r in old:
            ov = old[r][[c for c, n2 in MONEY_COLS if n2 == nm][0]]
            if ov != v:
                diff.append((r, nm, ov, v))
    log("   %-58s 与当前不同且当前=0 的格子 = %d 个；样例 %s"
        % (os.path.basename(bk), len(diff), diff[:6]))
log("done")
