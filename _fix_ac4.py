# -*- coding: utf-8 -*-
"""把被一次保存抹掉的核销列表头 AC4 恢复成「核销」（只补这一个空值，不动其它任何格）。"""
import io
import os
import shutil
import sys
import time

sys.path.insert(0, "c:/vnchoicedouble")
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

OUT = io.open("c:/vnchoicedouble/_fix_ac4_out.txt", "w", encoding="utf-8", newline="\n")


def log(s):
    OUT.write(str(s) + "\n")
    print(s)


p = "c:/vnchoicedouble/finance_data/cn_report.xlsx"
wb = openpyxl.load_workbook(p)
ws = wb[wb.sheetnames[0]]
log("AC4 当前 = %r" % ws["AC4"].value)
if ws["AC4"].value in (None, ""):
    bak = os.path.join("c:/vnchoicedouble/finance_data/backups",
                       "cn_report_%s_before_ac4fix.xlsx.bak" % time.strftime("%Y%m%d_%H%M%S"))
    shutil.copy2(p, bak)
    log("已备份 -> %s" % bak)
    ws["AC4"] = "核销"
    wb.save(p)
    log("已恢复 AC4 = '核销'")
else:
    log("AC4 有值，不动")
OUT.close()
