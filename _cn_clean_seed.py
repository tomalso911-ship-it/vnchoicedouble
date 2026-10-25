# -*- coding: utf-8 -*-
"""把中国报表恢复成干净种子（清掉回归测试写进去的测试行），并确认核销列居中。"""
import io
import json
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import openpyxl
import app

P = app.CN_REPORT
junk = []
if os.path.isfile(P):
    wb = openpyxl.load_workbook(P)
    ws = wb[wb.sheetnames[0]]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and ("PC-TEST" in c.value or "端到端" in c.value or "PC-E2E" in c.value):
                junk.append((c.coordinate, c.value))
    print("test rows found =", json.dumps(junk[:6], ensure_ascii=True))

if junk:
    os.remove(P)
    app._cn_seed_if_missing()
    try:
        app._cn_cache_invalidate()
    except Exception:
        pass
    print("-> regenerated clean seed")

wb = openpyxl.load_workbook(P)
ws = wb[wb.sheetnames[0]]
print("sheets =", json.dumps(wb.sheetnames, ensure_ascii=True))
print("AC4    =", json.dumps({"v": ws["AC4"].value, "h": (ws["AC4"].alignment.horizontal if ws["AC4"].alignment else None)}, ensure_ascii=True))
print("AC5    =", json.dumps({"v": ws["AC5"].value, "h": (ws["AC5"].alignment.horizontal if ws["AC5"].alignment else None)}, ensure_ascii=True))
print("B5/C5/E5 =", json.dumps([ws["B5"].value, ws["C5"].value, ws["E5"].value], ensure_ascii=True))
print("dims   =", ws.dimensions, "| total row E205 =", ws["E205"].value)
# 再来一遍：确认幂等（第二次不应再改文件）
before = os.stat(P).st_mtime_ns
app._cn_normalize_ac_align()
print("idempotent (mtime unchanged) =", os.stat(P).st_mtime_ns == before)
