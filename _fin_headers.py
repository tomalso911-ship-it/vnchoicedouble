import openpyxl, json
wb = openpyxl.load_workbook('finance_data/vn_report.xlsx')
out = []
for ws in wb.worksheets:
    rows = []
    for r in [1,2,3,4,5]:
        vals = [ws.cell(r, c).value for c in range(1, 11)]
        rows.append(vals)
    out.append({"sheet": ws.title, "max_row": ws.max_row, "rows": rows})
with open('finance_data/_headers.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
print('written')
