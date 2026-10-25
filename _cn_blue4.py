# -*- coding: utf-8 -*-
"""端到端验证「粘贴进来的黑字 → 点保存 → 自动变蓝」（全在临时目录，不碰真实数据）：
   ① 种子里故意写一个【黑色】数据格 C5（模拟粘贴带进来的黑字）；
   ② 打开中国报表：确认页面上 finCnNeedsBlueReload() = true、C5 是黑的；
   ③ 点「保存」：确认  后端把文件刷蓝 + 前端自动重载 → C5 在屏幕和文件里都变蓝；
   ④ 再确认 finCnNeedsBlueReload() 回到 false（不会反复重载）。"""
import datetime
import io
import os
import shutil
import socket
import sys
import tempfile
import threading
import time

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "_cn_blue4_out.txt")
LOG = []


def log(s):
    LOG.append("[%s] %s" % (datetime.datetime.now().strftime("%H:%M:%S"), s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(LOG) + "\n")


def port_open(port):
    s = socket.socket()
    s.settimeout(0.6)
    try:
        s.connect(("127.0.0.1", port))
        return True
    except Exception:
        return False
    finally:
        s.close()


import app  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.styles import Color, Font  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402

TMP = tempfile.mkdtemp(prefix="cnblue4_")
app.FIN_DIR = TMP
app.CN_REPORT = os.path.join(TMP, "cn_report.xlsx")
app.CN_SNAP_KEY = os.path.join(TMP, "cn_report.snapshot.key")
app.CN_SNAP_PREFIX = "cn_report.snapshot."
app.CN_TITLES_PATH = os.path.join(TMP, "fin_sheet_titles_cn.json")
log("临时报表目录 = %s" % TMP)

app._cn_seed_if_missing()
app._cn_normalize_style()          # 数据区先整体蓝化

# 模拟"从别处粘贴"：C5 带来源的黑色字体
wb = openpyxl.load_workbook(app.CN_REPORT)
ws = wb[wb.sheetnames[0]]
ws["A5"] = datetime.date(2026, 3, 28)
ws["B5"] = "UNC00009"
c = ws["C5"]
c.value = "粘贴进来的黑字"
c.font = Font(color=Color(rgb="FF000000"))
ws["E5"] = 41400000
# 第 6 行放"今天"→ 它会接管屏幕高亮（红+粗），这样第 5 行就不会被红盖住，
# 黑字模拟才有意义（否则整行被高亮成红，看不出"粘贴黑字"这个场景）。
ws["A6"] = datetime.date.today()
ws["B6"] = "UNC00010"
ws["C6"] = "今天这行（高亮红）"
ws["E6"] = 2000000
wb.save(app.CN_REPORT)
app._cn_cache_invalidate()
ok, detail = app._cn_cache_build(force=True)
log("快照重建 ok=%s detail=%s" % (ok, detail))

wb2 = openpyxl.load_workbook(app.CN_REPORT)
ws2 = wb2[wb2.sheetnames[0]]


def fcol(cell):
    v = getattr(getattr(cell.font, "color", None), "rgb", None)
    return str(v).upper() if isinstance(v, str) else None


log("文件里 C5 字色 = %s (应为 FF000000 黑)" % fcol(ws2["C5"]))

PORT = 5159
BASE = "http://127.0.0.1:%d/lite" % PORT
if not port_open(PORT):
    threading.Thread(target=lambda: app.app.run(host="127.0.0.1", port=PORT, debug=False,
                                               use_reloader=False, threaded=True),
                     daemon=True).start()
    for _ in range(60):
        if port_open(PORT):
            break
        time.sleep(0.3)

CELL_JS = """(p) => {
  try {
    var sn = fin_currentSnapshot();
    var sid = sn.sheetOrder[0];
    var cc = ((sn.sheets[sid].cellData || {})[String(p.r)] || {})[String(p.c)];
    if (!cc) return null;
    var s = cc.s;
    if (typeof s === 'string') s = (sn.styles && sn.styles[s]) || null;
    return { v: cc.v, cl: (s && s.cl) ? (s.cl.rgb || s.cl) : null };
  } catch(e){ return 'ERR:' + e.message; }
}"""

errors = []
with sync_playwright() as p:
    b = p.chromium.launch(channel="chrome", headless=True, args=["--no-sandbox"])
    pg = b.new_page(viewport={"width": 1500, "height": 950})
    pg.on("pageerror", lambda e: errors.append("pageerror: " + str(e)[:220]))
    pg.on("console", lambda m: errors.append("console.error: " + m.text[:220]) if m.type == "error" else None)
    pg.goto(BASE, wait_until="domcontentloaded")
    pg.wait_for_timeout(1200)
    pg.fill("#loginUser", "tom")
    pg.fill("#loginPwd", "66668888")
    pg.click("#loginBtn")
    pg.wait_for_selector("#topbar:not(.hidden)", timeout=15000)
    pg.evaluate("() => { try { switchView('finance'); } catch(e){} }")
    pg.wait_for_timeout(400)
    pg.evaluate("() => { var t = document.querySelector('#fin-tabs .fin-tab[data-key=\\'cn\\']'); if (t) t.click(); }")
    for _ in range(60):
        pg.wait_for_timeout(400)
        if pg.evaluate("() => (typeof finReady !== 'undefined') && !!finReady"):
            break
    pg.wait_for_timeout(3500)

    log("")
    log("② 保存前：")
    log("   C5(黑字那格) = %s" % pg.evaluate(CELL_JS, {"r": 4, "c": 2}))
    log("   A5(蓝字那格) = %s" % pg.evaluate(CELL_JS, {"r": 4, "c": 0}))
    need = pg.evaluate("() => { try { return finCnNeedsBlueReload(); } catch(e){ return 'ERR:' + e.message; } }")
    log("   finCnNeedsBlueReload() = %s  (应为 True)" % need)

    log("")
    log("③ 点「保存」…")
    pg.click('#fin_workspace button[onclick="fin_autoSave()"]')
    pg.wait_for_timeout(6000)
    log("   保存后 C5 = %s  (应为 #0000FF)" % pg.evaluate(CELL_JS, {"r": 4, "c": 2}))
    need2 = pg.evaluate("() => { try { return finCnNeedsBlueReload(); } catch(e){ return 'ERR:' + e.message; } }")
    log("   finCnNeedsBlueReload() = %s  (应为 False，不会再重载)" % need2)

    wb3 = openpyxl.load_workbook(app.CN_REPORT)
    ws3 = wb3[wb3.sheetnames[0]]
    log("   文件里 C5 字色 = %s (应为 FF0000FF 蓝) | C5 值 = %r" % (fcol(ws3["C5"]), ws3["C5"].value))
    log("   文件里 A5 值 = %r | E5 = %r" % (ws3["A5"].value, ws3["E5"].value))
    log("")
    log("页面报错 = %s" % errors[:6])
    pg.screenshot(path=os.path.join(HERE, "_cn_blue4_after.png"))
    b.close()

shutil.rmtree(TMP, ignore_errors=True)
log("临时目录已清理")
print("\n".join(LOG))
