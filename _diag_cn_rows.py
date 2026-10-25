# -*- coding: utf-8 -*-
"""把 cn_report.xlsx「实际支出」的全部记录行完整打印（三语块 A-E / K-O / U-Y + AC/AN），
   并在 vn_report.xlsx 的「Chi thực tế」里按 (日期, 凭证) 找同一笔，逐列对比。
   目的：看清到底哪些列被写成了 0。
"""
import datetime
import io
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
CN = os.path.join(FIN, "cn_report.xlsx")
VN = os.path.join(FIN, "vn_report.xlsx")
OUT = os.path.join(HERE, "_diag_cn_rows_out.txt")
L = []

BLOCKS = {
    "VI(A-E)": [1, 2, 3, 4, 5],
    "ZH(K-O)": [11, 12, 13, 14, 15],
    "EN(U-Y)": [21, 22, 23, 24, 25],
}


def log(s):
    L.append(str(s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(s)


def fmt(v):
    if v is None:
        return "-"
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


wbc = openpyxl.load_workbook(CN, data_only=False)
wsc = wbc[wbc.sheetnames[0]]

log("=== cn_report.xlsx「%s」全部非空行 ===" % wsc.title)
cn_rows = []
for r in range(1, 220):
    vals = [wsc.cell(r, c).value for c in (1, 2, 3, 4, 5)]
    joined = " ".join(str(v) for v in vals if v is not None)
    if not joined.strip():
        continue
    if "T\u1ed4NG C\u1ed8NG" in joined or "=SUM" in joined:
        log("--- 行%d（合计行）：%s" % (r, [fmt(x) for x in vals]))
        continue
    cn_rows.append(r)
    log("行%-4d  AC(29)=%s  AN(40)=%s" % (r, fmt(wsc.cell(r, 29).value), fmt(wsc.cell(r, 40).value)))
    for nm, cols in BLOCKS.items():
        log("     %-9s %s" % (nm, " | ".join(fmt(wsc.cell(r, c).value) for c in cols)))

log("")
log("=== vn_report.xlsx「Chi thực tế」同笔对比（按 日期+凭证）===")
wbv = openpyxl.load_workbook(VN, data_only=False)
wsv = wbv["Chi th\u1ef1c t\u1ebf"]
vn_index = {}
for r in range(5, 420):
    d = wsv.cell(r, 1).value
    b = wsv.cell(r, 2).value
    if isinstance(d, (datetime.datetime, datetime.date)):
        vn_index.setdefault((d.strftime("%Y-%m-%d"), str(b or "").strip().upper()), []).append(r)

for r in cn_rows:
    d = wsc.cell(r, 1).value
    b = wsc.cell(r, 2).value
    key = (d.strftime("%Y-%m-%d") if isinstance(d, (datetime.datetime, datetime.date)) else "",
           str(b or "").strip().upper())
    hits = vn_index.get(key, [])
    if not hits:
        log("cn 行%-4d %-12s %-10s → vn 里找不到同笔" % (r, key[0], key[1]))
        continue
    vr = hits[0]
    log("cn 行%-4d ↔ vn 行%-4d  %-12s %-10s" % (r, vr, key[0], key[1]))
    for nm, cols in BLOCKS.items():
        cv = " | ".join(fmt(wsc.cell(r, c).value) for c in cols)
        vv = " | ".join(fmt(wsv.cell(vr, c).value) for c in cols)
        same = "   ✓同" if cv == vv else "   <<< 不同"
        log("     %-9s cn: %s%s" % (nm, cv, same))
        if cv != vv:
            log("     %-9s vn: %s" % ("", vv))

wbc.close()
wbv.close()
log("done")
