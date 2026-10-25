# -*- coding: utf-8 -*-
"""仅对 merge_report.xlsx 的 TỔNG HỢP / Chi thực tế 合计格写入静态求值结果。
   原因：合并报表是生成出来的派生报表，其总表里的跨表公式在前端 Univer 重算时
   不能可靠地反映中国实际支出。写入静态值后快照可直接显示，同时不改动越南报表。"""
import os, sys, re, shutil, datetime, glob
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
OUT = os.path.join(HERE, "_fix_merge_tonghop_out.txt")
L = []
def log(s=""):
    L.append(str(s))
    with open(OUT, "w", encoding="utf-8") as f: f.write("\n".join(L)+"\n")
    print(s)

_formula_re = re.compile(r'SUM\s*\((.*)\)\s*$', re.IGNORECASE)
_ref_re = re.compile(r'^\s*([A-Za-z]+)\s*(\d+)\s*$')
_arith_re = re.compile(r'^(.+?)\s*([+\-])\s*(.+)$')
_MAX_DEPTH = 32

def _col_to_idx(col):
    n = 0
    for ch in col.upper(): n = n*26 + (ord(ch)-64)
    return n

def _parse_ref(ref):
    if not isinstance(ref, str): return None, None, None
    ref = ref.strip().replace('$','')
    if '!' in ref:
        sheet, cell = ref.split('!',1)
        sheet = sheet.strip().strip("'")
    else:
        sheet, cell = None, ref
    m = _ref_re.match(cell.strip())
    if not m: return None, None, None
    return sheet, _col_to_idx(m.group(1)), int(m.group(2))

def _cell_val(wb, sheet, col, row, memo, visited, depth=0):
    if depth > _MAX_DEPTH: return None
    key = (sheet, col, row)
    if key in memo: return memo[key]
    if key in visited: return None
    visited.add(key)
    try: v = wb[sheet].cell(row, col).value
    except Exception: memo[key]=None; return None
    if isinstance(v, str) and v.startswith('='):
        ev = _eval_formula(v, wb, sheet, memo, visited, depth+1)
        memo[key]=ev; return ev
    memo[key]=v; return v

def _eval_formula(formula, wb, cur_sheet, memo, visited=None, depth=0):
    if visited is None: visited = set()
    if depth > _MAX_DEPTH: return None
    f = formula.strip()
    if not f.startswith('='): return None
    expr = f[1:].strip()
    # 数字字面量公式（如 =150000000）：直接返回数值，否则会被漏算
    try:
        return float(expr)
    except (ValueError, TypeError):
        pass
    m = _formula_re.match(expr)
    if m:
        total = 0.0
        for part in m.group(1).split(','):
            part = part.strip()
            if ':' in part:
                a,b = part.split(':')
                sa,ca,ra = _parse_ref(a); sb,cb,rb = _parse_ref(b)
                sheet = sa or sb or cur_sheet
                for r in range(ra, rb+1):
                    for c in range(ca, cb+1):
                        v = _cell_val(wb, sheet, c, r, memo, visited, depth+1)
                        if isinstance(v,(int,float)): total += v
            else:
                s,c,r = _parse_ref(part)
                v = _cell_val(wb, s or cur_sheet, c, r, memo, visited, depth+1)
                if isinstance(v,(int,float)): total += v
        return total
    m = _arith_re.match(expr)
    if m:
        left,op,right = m.group(1).strip(), m.group(2), m.group(3).strip()
        lv = _eval_formula(left if left.startswith('=') else '='+left, wb, cur_sheet, memo, visited, depth+1)
        rv = _eval_formula(right if right.startswith('=') else '='+right, wb, cur_sheet, memo, visited, depth+1)
        if isinstance(lv,(int,float)) and isinstance(rv,(int,float)):
            return lv+rv if op=='+' else lv-rv
    s,c,r = _parse_ref(expr)
    if s is None and c is None and r is None: return None
    v = _cell_val(wb, s or cur_sheet, c, r, memo, visited, depth+1)
    if isinstance(v,(int,float)): return v
    return None

def stamp_sheet(wb, sheet_name):
    if sheet_name not in wb.sheetnames: return 0
    ws = wb[sheet_name]
    done = 0
    if sheet_name == "TỔNG HỢP":
        for r in range(1, ws.max_row+1):
            for c in (3, 13, 23):
                v = ws.cell(r, c).value
                if isinstance(v, str) and v.startswith('='):
                    memo = {}
                    val = _eval_formula(v, wb, sheet_name, memo)
                    if isinstance(val,(int,float)):
                        ws.cell(r, c).value = val
                        done += 1
    if sheet_name == "Chi thực tế":
        for r in (3, ws.max_row):
            for c in (5, 15, 25):
                v = ws.cell(r, c).value
                if isinstance(v, str) and v.startswith('='):
                    memo = {}
                    val = _eval_formula(v, wb, sheet_name, memo)
                    if isinstance(val,(int,float)):
                        ws.cell(r, c).value = val
                        done += 1
    return done

def fix_merge():
    path = os.path.join(FIN, "merge_report.xlsx")
    wb = openpyxl.load_workbook(path, data_only=False)
    n1 = stamp_sheet(wb, "TỔNG HỢP")
    n2 = stamp_sheet(wb, "Chi thực tế")
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = os.path.join(FIN, "backups", "merge_report_%s_before_fix_tonghop_mergeonly.xlsx.bak" % ts)
    shutil.copyfile(path, bak)
    tmp = path + ".fixm.tmp"
    wb.save(tmp)
    os.replace(tmp, path)
    log("[%s] TỔNG HỢP=%d 格, Chi thực tế=%d 格 | 备份=%s" % (os.path.basename(path), n1, n2, os.path.basename(bak)))
    wb.close()

if __name__ == "__main__":
    fix_merge()
    for key in ("merge_report.snapshot.key",):
        p = os.path.join(FIN, key)
        try: os.remove(p); log("删快照键 %s" % key)
        except OSError: pass
    log("done")
