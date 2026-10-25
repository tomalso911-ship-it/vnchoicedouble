# -*- coding: utf-8 -*-
"""快速找：vn/merge 各表哪些单元格的值是单独一个「我」。"""
import os, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

FIN = os.path.join(os.path.dirname(os.path.abspath(__file__)), "finance_data")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_find_wo_out.txt")
L = []
def log(s=""):
    L.append(str(s)); open(OUT,"w",encoding="utf-8").write("\n".join(L)+"\n"); print(s)

for fn, tag in (("vn_report.xlsx","vn"), ("merge_report.xlsx","merge")):
    p = os.path.join(FIN, fn)
    wb = openpyxl.load_workbook(p, data_only=False)
    log("=== %s ===" % tag)
    for sh in wb.worksheets:
        found = []
        for r in range(1, min(sh.max_row, 50) + 1):  # 只扫前50行
            for c in range(1, sh.max_column + 1):
                v = sh.cell(r, c).value
                if isinstance(v, str) and v.strip() == "我":
                    found.append("%s!%s (行%d)" % (sh.title, sh.cell(r,c).coordinate, r))
        if found:
            log("[%s] %s" % (sh.title, ", ".join(found)))
    wb.close()
log("done")
