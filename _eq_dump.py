# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

wb = openpyxl.load_workbook("finance_data/equity_report.xlsx", data_only=False)
for ws in wb.worksheets:
    print("=== %s  rows=%s cols=%s ===" % (ws.title, ws.max_row, ws.max_column))
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        vals = []
        for c in row:
            v = c.value
            if v is None:
                vals.append("")
            else:
                vals.append(str(v))
        if any(vals):
            print(" r%-2d | %s" % (row[0].row, " | ".join(vals[:14])))
