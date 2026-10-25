# -*- coding: utf-8 -*-
"""把 cn_report.xlsx 行10 / 行14 回填来的「机翻味」摘要润成报表口径。
   只改文字，公司名/金额/日期一律不动；改前逐格校验当前值，不符就跳过该格。
   备份 + 原子落盘。
"""
import datetime
import os
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
CN = os.path.join(FIN, "cn_report.xlsx")
BKD = os.path.join(FIN, "backups")
OUT = os.path.join(HERE, "_polish_cn_desc_out.txt")
L = []


def log(s=""):
    L.append(str(s))
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    print(s)


# (行, 列) -> (必须包含的关键字, 新文字)
PLAN = {
    (10, 3):  ("ZHONGLI", "Công ty TNHH Thương mại Điện máy Zhongli"),
    (10, 13): ("中力电机", "中力电机贸易有限公司"),
    (10, 23): ("ZHONGLI", "Zhongli Electrical Machinery Trading Co., Ltd."),
    (14, 3):  ("Huy Ph", "Thanh toán tiền vật tư Công ty Huy Phượng"),
    (14, 13): ("Huy Phuong", "向 Huy Phượng 公司支付材料款"),
    (14, 23): ("Huy Phuong", "Payment for materials to Huy Phuong Company"),
}

wbc = openpyxl.load_workbook(CN, data_only=False)
wsc = wbc[wbc.sheetnames[0]]

log("=== 逐格校验 + 润色 ===")
todo = []
for (r, c), (key, new) in sorted(PLAN.items()):
    cur = wsc.cell(r, c).value
    s = str(cur or "")
    if key.lower() not in s.lower():
        log("  行%d 列%-2d 跳过：当前值 %r 不含关键字 %r（可能已被手动改过）" % (r, c, s[:40], key))
        continue
    todo.append((r, c, s, new))
    log("  行%d 列%-2d %s\n        -> %s" % (r, c, s, new))

if not todo:
    log("没有需要改的格子，未落盘。")
    sys.exit(0)

ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
bak = os.path.join(BKD, "cn_report_%s_before_descpolish.xlsx.bak" % ts)
shutil.copy2(CN, bak)
log("")
log("已备份 -> %s" % bak)

for r, c, _old, new in todo:
    wsc.cell(r, c).value = new
    wsc.cell(r, c).number_format = "General"

tmp = CN + ".fix.tmp"
wbc.save(tmp)
wbc.close()
os.replace(tmp, CN)
log("已原子落盘：%s" % CN)

wb2 = openpyxl.load_workbook(CN, data_only=False)
ws2 = wb2[wb2.sheetnames[0]]
log("")
log("=== 复核 ===")
for r in (10, 14):
    log("行%-3d VI=%s" % (r, ws2.cell(r, 3).value))
    log("     ZH=%s" % ws2.cell(r, 13).value)
    log("     EN=%s" % ws2.cell(r, 23).value)
    log("     金额 VI=%s ZH=%s EN=%s" % (ws2.cell(r, 5).value, ws2.cell(r, 15).value, ws2.cell(r, 25).value))
t = sum(v for r in range(5, 205) for c in (5, 15, 25)
        if isinstance((v := ws2.cell(r, c).value), (int, float)))
log("支出合计(三块相加) = %s  （每块 %s）" % (format(t, ","), format(t // 3, ",")))
wb2.close()
log("done")
