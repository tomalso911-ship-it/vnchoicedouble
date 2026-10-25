# -*- coding: utf-8 -*-
"""清理中国报表「摘要」列的脏值：文字列里的数字 0 + Excel 错误值(#NAME?/#VALUE! 及被译成 ＃姓名？ 的残渣)。

- 先备份；
- 只清「文字类列」(C/M/W = 摘要, G/Q/AA = 备注) 里的数值 0，以及任何错误值字符串；
- 保存后失效中国报表快照缓存，前端下次即读到干净数据。
用法: python _cn_clean_junk.py [--dry-run]
"""
import os
import re
import shutil
import sys
import time
import glob

import openpyxl

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

CN = 'finance_data/cn_report.xlsx'
TEXT_COLS = set()
for base in (3, 7):
    for off in (0, 10, 20):
        TEXT_COLS.add(base + off)

ERR_ASCII = re.compile(r'^#(VALUE|NAME|REF|DIV/0|N/A|NULL|NUM|GETTING_DATA|SPILL|CALC)[!?]?$', re.I)
ERR_ANY = re.compile(r'^\s*[#＃].*[?!？!]\s*$')          # 含全角 ＃姓名？ 这类"被翻译过的错误值"


def is_err(v):
    if not isinstance(v, str):
        return False
    s = v.strip()
    return bool(ERR_ASCII.match(s) or ERR_ANY.match(s))


def main():
    dry = '--dry-run' in sys.argv
    wb = openpyxl.load_workbook(CN)
    changed = []
    for ws in wb.worksheets:
        for r in range(1, ws.max_row + 1):
            for c in range(1, min(ws.max_column, 30) + 1):
                cell = ws.cell(r, c)
                v = cell.value
                if v is None:
                    continue
                hit = None
                if is_err(v):
                    hit = 'ERR'
                elif c in TEXT_COLS and isinstance(v, (int, float)) and not isinstance(v, bool) and v == 0:
                    hit = 'ZERO'
                if hit:
                    changed.append((ws.title, r, c, hit, repr(v)))
                    if not dry:
                        cell.value = None

    for x in changed[:40]:
        print('  %s r%d c%d %s %s' % x)
    print('总改动：%d 格' % len(changed))

    if dry or not changed:
        return

    ts = time.strftime('%Y%m%d_%H%M%S')
    bak = CN + '.bak_before_junk_clean.' + ts
    shutil.copyfile(CN, bak)
    print('已备份 ->', bak)

    tmp = CN + '.clean.tmp'
    wb.save(tmp)
    os.replace(tmp, CN)

    # 失效中国报表快照缓存（键 + 旧快照）
    for p in glob.glob('finance_data/cn_report.snapshot.*') + ['finance_data/cn_report.snapshot.key']:
        try:
            os.remove(p)
        except OSError:
            pass
    print('已失效中国报表快照缓存')


if __name__ == '__main__':
    main()
