# -*- coding: utf-8 -*-
"""只读排查：中国报表第 14 行（截图里没被翻译成中/英的那行）到底缺什么。
   ① 直接读 finance_data/cn_report.xlsx 的三语块（A/C 越南文 | K/M 中文 | U/W 英文）
   ② 把那一格原文拿去走服务端同一套翻译链（/api/translate，只 GET、不改任何数据）
   ③ 判断被哪一道闸挡住：fin_isCodeLike(长度/编码样式) / 保护词 / 翻译链全失败 / 译文=原文"""
import io
import json
import os
import sys
import urllib.parse
import urllib.request

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "_bug_c14_out.txt")
L = []


def log(s):
    L.append(str(s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")


import openpyxl  # noqa: E402

PATH = os.path.join(HERE, "finance_data", "cn_report.xlsx")
log("报表文件 = %s (存在=%s)" % (PATH, os.path.isfile(PATH)))

wb = openpyxl.load_workbook(PATH, data_only=True)
log("工作表 = %s" % wb.sheetnames)
ws = wb[wb.sheetnames[0]]
log("max_row = %s, max_col = %s" % (ws.max_row, ws.max_column))

COLS = ["A", "B", "C", "D", "E", "F", "K", "M", "U", "W", "AC"]
log("")
log("① 第 9~16 行三语块（1 基行号）：")
for r in range(9, 17):
    log("  --- 行 %d ---" % r)
    for c in COLS:
        v = ws["%s%d" % (c, r)].value
        if v is None or v == "":
            continue
        log("    %-3s = %r" % (c, v))

vi = ws["C14"].value
log("")
log("② 定位：C14（越南文/摘要）= %r" % vi)
log("   长度 = %s" % (len(vi) if isinstance(vi, str) else "N/A"))
log("   K/M/U/W 同行的中英块： K14=%r M14=%r U14=%r W14=%r"
    % (ws["K14"].value, ws["M14"].value, ws["U14"].value, ws["W14"].value))

if isinstance(vi, str) and vi.strip():
    log("")
    log("③ 前端那两道闸判定（与 choice_lite.html 同逻辑）：")
    s = vi.strip()
    log("   fin_isCodeLike: 长度>160 ? %s" % (len(s) > 160))
    import re
    log("   fin_isCodeLike: 纯数字/日期样式 ? %s" % bool(re.match(r"^[\d\s.,%+\-()/:]+$", s)))
    log("   fin_isCodeLike: 无空格+有数字+有字母(编码样式) ? %s"
        % ((" " not in s) and bool(re.search(r"\d", s)) and bool(re.search(r"[A-Za-zÀ-ỹ]", s))))
    log("   fin_hasNoTrans(保护词 dabaco) ? %s" % ("dabaco" in s.lower()))

    log("")
    log("④ 走服务端翻译链（若应用正在 5050 端口跑）：")
    for tgt in ("zh", "en"):
        u = "http://127.0.0.1:5050/api/translate?q=%s&target=%s" % (urllib.parse.quote(vi), tgt)
        try:
            with urllib.request.urlopen(u, timeout=40) as resp:
                j = json.loads(resp.read().decode("utf-8"))
            log("   target=%s -> ok=%s text=%r" % (tgt, j.get("ok"), j.get("text")))
            if j.get("text") == vi:
                log("      ★译文与原文一模一样 → 前端会判定「没译出来」，屏幕永远显示原文")
        except Exception as e:
            log("   target=%s -> 请求失败：%s（应用可能没在跑）" % (tgt, e))

    log("")
    log("⑤ 服务端持久译文缓存 finance_data/fin_trans_cache.json 里有没有这条：")
    try:
        with io.open(os.path.join(HERE, "finance_data", "fin_trans_cache.json"), encoding="utf-8") as fh:
            cache = json.load(fh)
        hits = [k for k in cache if vi in k or k in vi]
        log("   缓存条目数 = %d；匹配到 %d 条" % (len(cache), len(hits)))
        for k in hits[:5]:
            log("     %r -> %r" % (k, cache[k]))
    except Exception as e:
        log("   读缓存失败：%s" % e)

print("\n".join(L))
