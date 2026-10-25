# -*- coding: utf-8 -*-
"""细化：逐列统计非空格数，确认上传的中文报表到底有几个语言块。"""
import os
import sys
import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(BASE, "finance_data", "vn_report.xlsx")

wb = openpyxl.load_workbook(FIN, data_only=False)
for ws in wb.worksheets:
    counts = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 600), max_col=28):
        for cell in row:
            if cell.value not in (None, ""):
                counts[cell.column] = counts.get(cell.column, 0) + 1
    cols = sorted(counts)
    print("--- %r  行=%d 列=%d" % (ws.title, ws.max_row, ws.max_column))
    print("    有内容的列: " + ", ".join("%s(%d)" % (openpyxl.utils.get_column_letter(c), counts[c]) for c in cols))
    # 头 8 行 A..R
    for r in range(1, min(9, ws.max_row + 1)):
        vals = []
        for c in range(1, 19):
            v = ws.cell(row=r, column=c).value
            vals.append("·" if v in (None, "") else str(v)[:14])
        print("    r%-2d %s" % (r, " | ".join(vals)))
    print()

print("=== finance_data/backups 最近 6 份 ===")
import datetime
bd = os.path.join(BASE, "finance_data", "backups")
if os.path.isdir(bd):
    fs = sorted(((os.stat(os.path.join(bd, f)).st_mtime, f) for f in os.listdir(bd)), reverse=True)[:6]
    for t, f in fs:
        print("  %s  %s" % (datetime.datetime.fromtimestamp(t), f))

print()
print("=== finance_data 下 xlsx / bak 一览 ===")
for f in sorted(os.listdir(os.path.join(BASE, "finance_data"))):
    if f.endswith((".xlsx", ".bak", ".xlsx.bak")) or "bak" in f:
        p = os.path.join(BASE, "finance_data", f)
        if os.path.isfile(p):
            print("  %-46s %9d  %s" % (f, os.path.getsize(p), datetime.datetime.fromtimestamp(os.stat(p).st_mtime)))
