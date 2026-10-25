# -*- coding: utf-8 -*-
"""把「合并结果」从越南报表里拆出来（用户 2026-10-14 要求）：

  越南财务报表  = 干净版（不含任何中国记录）
  合并财务总报表 = 合并版（越南 + 中国记录，中国行蓝字）

做法（三步，全部先备份、可回退）：
  1) 当前 vn_report.xlsx（= 合并后的版本，含 CN 标记行）另存为 merge_report.xlsx；
     同时再留一份 backups/vn_report_<ts>_merged_result.xlsx.bak 作为合并结果的备份；
  2) 从 backups/ 里挑一份【CN 标记行 = 0】的 before_merge 备份，还原成 vn_report.xlsx；
  3) 校验两边（vn 必须是 0 行 CN；merge 必须 >0 行 CN），并清掉 vn 的快照缓存键。

用法：
  python _split_merge_report.py          # 预演（只检查、不写盘）
  python _split_merge_report.py --apply  # 真正执行
"""
import datetime
import io
import os
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
BAK = os.path.join(FIN, "backups")
VN = os.path.join(FIN, "vn_report.xlsx")
MERGE = os.path.join(FIN, "merge_report.xlsx")
SHEET = "Chi th\u1ef1c t\u1ebf"
MARK_COL = 40
OUT = os.path.join(HERE, "_split_merge_report_out.txt")
L = []


def log(s):
    L.append(str(s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(s)


def cn_mark_count(path):
    """读一个 xlsx，数「Chi thực tế」里 CN: 标记行；出错返回 (-1, -1)。
    ★openpyxl 是按【扩展名】判格式的：.bak 必须先复制成 .xlsx 才能读。"""
    probe = None
    src = path
    if not path.lower().endswith(".xlsx"):
        probe = path + ".probe.xlsx"
        shutil.copyfile(path, probe)
        src = probe
    try:
        wb = openpyxl.load_workbook(src, data_only=False)
        ws = wb[SHEET]
        n = 0
        for r in range(5, 420):
            mk = ws.cell(r, MARK_COL).value
            if isinstance(mk, str) and mk.startswith("CN:"):
                n += 1
        total_sheets = len(wb.sheetnames)
        wb.close()
        return n, total_sheets
    except Exception as e:
        log("   ! 读取失败 %s：%s" % (os.path.basename(path), e))
        return -1, -1
    finally:
        if probe:
            try:
                os.remove(probe)
            except OSError:
                pass


def main():
    apply = "--apply" in sys.argv
    log("模式 = %s" % ("★真正写盘" if apply else "预演（不写盘）"))
    log("")

    # ---- 0) 现状 ----
    vn_cnt, vn_sheets = cn_mark_count(VN)
    log("当前 vn_report.xlsx：CN 标记行 = %d，工作表 %d 张" % (vn_cnt, vn_sheets))
    if vn_cnt <= 0:
        log("★ 当前越南报表里没有中国记录 —— 已经拆过了，无需再执行。")
        return 0
    if os.path.isfile(MERGE):
        mg_cnt, _ = cn_mark_count(MERGE)
        log("已存在的 merge_report.xlsx：CN 标记行 = %d（本次会被覆盖）" % mg_cnt)

    # ---- 1) 挑一份干净的备份 ----
    cands = []
    if os.path.isdir(BAK):
        for fn in sorted(os.listdir(BAK)):
            if fn.startswith("vn_report_") and fn.endswith(".bak") and "_before_merge" in fn:
                cands.append(os.path.join(BAK, fn))
    log("候选 before_merge 备份 %d 份：" % len(cands))
    clean = []
    for p in cands:
        n, sh = cn_mark_count(p)
        log("   %-58s CN 标记 = %-3s 工作表 = %s" % (os.path.basename(p), n, sh))
        if n == 0:
            clean.append(p)
    if not clean:
        log("")
        log("✗ 没有任何一份备份是【干净】的（CN 标记都为 0 的才可用）。")
        log("  为避免误删数据，本次不执行任何写盘。请人工指定要用哪份备份。")
        return 2
    pick = clean[-1]          # 多份干净的取最新的一份
    log("")
    log("→ 选用的干净备份：%s" % os.path.basename(pick))

    # ---- 2) 目标状态预览 ----
    log("")
    log("将要执行：")
    log("   a) %s  →  备份为 backups/vn_report_<ts>_merged_result.xlsx.bak"
        % os.path.basename(VN))
    log("   b) %s  →  merge_report.xlsx（=%d 行中国记录，供「合并财务总报表」使用）"
        % (os.path.basename(VN), vn_cnt))
    log("   c) %s  →  vn_report.xlsx（干净版，0 行中国记录）" % os.path.basename(pick))
    log("   d) 删除 vn 快照缓存键 → 刷新后越南报表立刻恢复为纯净版")

    if not apply:
        log("")
        log("=== 预演结束：加 --apply 才会写盘 ===")
        return 0

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    os.makedirs(BAK, exist_ok=True)

    # ---- 3) 写盘 ----
    keep = os.path.join(BAK, "vn_report_%s_merged_result.xlsx.bak" % ts)
    shutil.copyfile(VN, keep)
    log("✔ a) 合并结果已备份 → %s" % os.path.basename(keep))

    tmp = MERGE + ".split.tmp"
    shutil.copyfile(VN, tmp)
    os.replace(tmp, MERGE)
    log("✔ b) merge_report.xlsx 已生成（来源 = 合并后的越南报表）")

    tmp2 = VN + ".restore.tmp"
    shutil.copyfile(pick, tmp2)
    os.replace(tmp2, VN)
    log("✔ c) vn_report.xlsx 已还原为干净版")

    try:
        os.remove(os.path.join(FIN, "vn_report.snapshot.key"))
        log("✔ d) 已清掉 vn 快照缓存键")
    except OSError:
        log("   d) 快照缓存键本就不存在（下次会自动重建）")

    # ---- 4) 校验 ----
    v2, s2 = cn_mark_count(VN)
    m2, s3 = cn_mark_count(MERGE)
    log("")
    log("校验：vn_report.xlsx CN 标记 = %d（期望 0）工作表 %d 张" % (v2, s2))
    log("校验：merge_report.xlsx CN 标记 = %d（期望 >0）工作表 %d 张" % (m2, s3))
    if v2 == 0 and m2 > 0:
        log("✅ 拆分成功：越南报表纯净，合并报表含中国记录。")
        return 0
    log("✗ 校验未通过，请检查上面的输出。")
    return 3


if __name__ == "__main__":
    sys.exit(main())
