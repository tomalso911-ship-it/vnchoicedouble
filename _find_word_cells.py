#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""定位 Choice VN / T.L / A.C 这些格子在哪个报表、哪张表、哪些单元格。"""
import os
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(ROOT, "finance_data")

WANT = ("choice vn", "choicevn", "t.l", "a.c", "tl", "ac", "dabaco")

for fn in ("equity_report.xlsx", "cn_report.xlsx", "vn_report.xlsx", "merge_report.xlsx"):
    p = os.path.join(FIN, fn)
    if not os.path.isfile(p):
        print("(缺) " + fn)
        continue
    print("=" * 74)
    print(fn, os.path.getsize(p), "字节")
    try:
        wb = openpyxl.load_workbook(p, data_only=False)
    except Exception as e:
        print("  打不开:", e)
        continue
    for ws in wb.worksheets:
        hits = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, 400),
                               max_col=min(ws.max_column or 0, 80)):
            for c in row:
                v = c.value
                if not isinstance(v, str):
                    continue
                s = v.strip()
                if s.lower() in WANT:
                    hits.append("%s%d=%r" % (c.column_letter, c.row, s))
        if hits:
            print("  [%s] %d 处: %s" % (ws.title, len(hits), ", ".join(hits[:30])))
    wb.close()
print("=" * 74)
