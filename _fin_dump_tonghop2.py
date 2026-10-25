# -*- coding: utf-8 -*-
"""看 TỔNG HỢP 是否含 中文块(K..)/英文块(U..)：把 r4..r22 的前 24 列全打出来。"""
import io, sys
from openpyxl import load_workbook
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
wb = load_workbook(r'finance_data/vn_report.xlsx', data_only=False)
ws = wb['TỔNG HỢP']
print('max_col =', ws.max_column)
for r in range(4, 23):
    cells = []
    for c in range(1, 25):
        v = ws.cell(row=r, column=c).value
        if v is None or v == '':
            continue
        if isinstance(v, str):
            v = v.replace('\n', '\\n')
            if len(v) > 40:
                v = v[:40] + '…'
        cells.append('%s%d=%r' % (chr(64 + c) if c <= 26 else '?', r, v))
    if cells:
        print(' r%-3d %s' % (r, ' | '.join(cells)))
