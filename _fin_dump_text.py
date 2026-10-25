# -*- coding: utf-8 -*-
"""列出 finance_data/vn_report.xlsx 里所有「文字类」单元格，用来评估多语言改造量。

输出：每张表的名字/尺寸、所有字符串单元格(坐标+值)、去重后的文字清单（按出现次数）、
      数字/公式单元格计数、合并区域数。只读，不修改任何文件。
用法: python _fin_dump_text.py [xlsx路径]
"""
import collections
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "finance_data", "vn_report.xlsx")

import openpyxl  # noqa: E402

wb = openpyxl.load_workbook(path, data_only=False)
print("文件: %s" % path)
print("工作表 %d 张: %s" % (len(wb.sheetnames), wb.sheetnames))
print()

all_texts = collections.Counter()
for ws in wb.worksheets:
    strings = []
    nfmt = nnum = nformula = nempty = 0
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if v is None:
                nempty += 1
            elif isinstance(v, str):
                s = v.strip()
                if not s:
                    nempty += 1
                elif s.startswith("="):
                    nformula += 1
                else:
                    strings.append((c.coordinate, s))
                    all_texts[s] += 1
            elif isinstance(v, (int, float)):
                nnum += 1
            else:
                nfmt += 1
    merges = list(ws.merged_cells.ranges)
    print("── 工作表 %r  dims=%s  max_row=%s max_col=%s" % (ws.title, ws.dimensions, ws.max_row, ws.max_column))
    print("   文字格 %d · 数字 %d · 公式 %d · 其它 %d · 空 %d · 合并区 %d"
          % (len(strings), nnum, nformula, nfmt, nempty, len(merges)))
    for coord, s in strings[:400]:
        print("     %-6s %s" % (coord, s[:160]))
    if len(strings) > 400:
        print("     … 还有 %d 个文字格" % (len(strings) - 400))
    print()

print("=" * 70)
print("去重文字共 %d 条（总出现 %d 次）——这就是需要翻译的「标签/表头」规模："
      % (len(all_texts), sum(all_texts.values())))
for s, n in all_texts.most_common(200):
    print("   %3d×  %s" % (n, s[:150]))
