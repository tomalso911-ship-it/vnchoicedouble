# -*- coding: utf-8 -*-
"""看看 Chi tiết thu-chi 的隐藏列（N..Z）在 xlsx 里到底有没有内容/公式。"""
import io, sys
from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
wb = load_workbook('finance_data/vn_report.xlsx', data_only=False)
ws = wb['Chi tiết thu-chi']
cols = list(range(9, 27))   # I..Z
cnt = {c: 0 for c in cols}
samples = {}
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=9, max_col=26):
    for cell in row:
        if cell.value is not None and cell.value != '':
            cnt[cell.column] += 1
            if cell.column not in samples:
                samples[cell.column] = (cell.coordinate, cell.value)
print('列(9=I..26=Z) 非空格数 / 首个样例：')
for c in cols:
    print('  col%-3d %-14s %s' % (c, cnt[c], samples.get(c, '')))
print()
for r in [5, 6, 7, 95, 530, 533, 568, 569, 570]:
    print('row%-4d' % r, [(chr(64 + c), ws.cell(row=r, column=c).value) for c in cols if ws.cell(row=r, column=c).value is not None])
