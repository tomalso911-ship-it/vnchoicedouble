# -*- coding: utf-8 -*-
"""逐行比对：ledger 公式重算的余额 vs 文件里的静态余额（或缓存的公式值）。
输出不一致的行，供判断"是文件笔误还是我们算错"。"""
import io, sys, datetime, re
from openpyxl import load_workbook
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
wbf = load_workbook('finance_data/vn_report.xlsx', data_only=False)
wbv = load_workbook('finance_data/vn_report.xlsx', data_only=True)
wsf, wsv = wbf['Chi tiết thu-chi'], wbv['Chi tiết thu-chi']


def is_date(v):
    return isinstance(v, (datetime.datetime, datetime.date))


def num(r, c):
    """按【服务端口径】取值：纯数字常量公式（如 '=150000000'）算它的数字（服务端就是这么做的）。"""
    v = wsv.cell(row=r, column=c).value
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    f = wsf.cell(row=r, column=c).value
    if isinstance(f, str) and f.startswith('='):
        try:
            return float(f[1:].replace(',', '').strip())
        except ValueError:
            return 0.0
    return 0.0


REGIONS = [('VPbank', 5, 6, 77), ('MBbank', 79, 80, 292), ('ACB', 293, 294, 299),
           ('Tiền mặt 1', 300, 302, 568)]
for (name, hdr, lo, hi) in REGIONS:
    # 期初：上一行是静态数字余额才引用
    prev_f = wsf.cell(row=lo - 1, column=6).value
    run = float(prev_f) if isinstance(prev_f, (int, float)) and not isinstance(prev_f, bool) else 0.0
    diffs = []
    rows = 0
    for r in range(lo, hi + 1):
        if not is_date(wsv.cell(row=r, column=1).value):
            continue
        rows += 1
        run += num(r, 4) - num(r, 5)
        fv = wsv.cell(row=r, column=6).value
        if isinstance(fv, (int, float)) and not isinstance(fv, bool):
            if abs(float(fv) - run) > 0.5:
                diffs.append((r, fv, run))
    print('%-11s 行%d..%d  交易 %d 笔  期初=%s  公式末值=%.0f' % (
        name, lo, hi, rows, (('%g' % float(prev_f)) if isinstance(prev_f, (int, float)) else '0'), run))
    if diffs:
        print('   ★与文件静态值不一致 %d 行（前 8 行）：' % len(diffs))
        for (r, fv, run2) in diffs[:8]:
            print('      F%-4d 文件=%-16g 公式=%-16.0f 差=%g' % (r, fv, run2, run2 - float(fv)))
    else:
        print('   ✓ 与文件静态值完全一致')
