# -*- coding: utf-8 -*-
"""扫描 vn_report.xlsx：哪些格的缓存值是 Excel 错误（#VALUE!...），并打出其公式。
   用 read_only 流式读，且只读前 30 列（'Chi thực tế' 的 max_col 有 16381，全读会卡死）。"""
import io, sys
from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
ERRS = ('#VALUE!', '#REF!', '#DIV/0!', '#N/A', '#NAME?', '#NULL!', '#NUM!', '#ERROR!')
MAXC = 30

wbf = load_workbook('finance_data/vn_report.xlsx', data_only=False, read_only=True)
wbv = load_workbook('finance_data/vn_report.xlsx', data_only=True, read_only=True)
print('工作表:', wbf.sheetnames)
total = 0
for name in wbf.sheetnames:
    wsf = wbf[name]
    wsv = wbv[name]
    fmap = {}
    for row in wsf.iter_rows(min_col=1, max_col=MAXC):
        for c in row:
            if c.value is not None:
                fmap[c.coordinate] = c.value
    hits = []
    for row in wsv.iter_rows(min_col=1, max_col=MAXC):
        for c in row:
            if isinstance(c.value, str) and c.value.strip() in ERRS:
                hits.append((c.coordinate, fmap.get(c.coordinate)))
    print('\n== %-22s 错误格 %d 个' % (name, len(hits)))
    total += len(hits)
    for coord, f in hits[:25]:
        print('    %-6s 公式 = %r' % (coord, f))
print('\n错误格合计 =', total)
