# -*- coding: utf-8 -*-
"""只把单元格里【单独一个】中文「我」改成拉丁字母「I」。
   不碰描述句子里作为人称代词的「我」。"""
import datetime
import os
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
OUT = os.path.join(HERE, "_fix_wo_i_out.txt")
L = []


def log(s=""):
    L.append(str(s))
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    print(s)


def replace_standalone_wo(path, tag):
    wb = openpyxl.load_workbook(path, data_only=False)
    hits = []
    for sh in wb.worksheets:
        for r in range(1, sh.max_row + 1):
            for c in range(1, sh.max_column + 1):
                cell = sh.cell(r, c)
                v = cell.value
                # 只替换整格内容为 "我"（可能带首尾空白）的单元格
                if isinstance(v, str) and v.strip() == "我":
                    cell.value = "I"
                    coord = "%s!%s" % (sh.title, cell.coordinate)
                    hits.append(coord)
    if hits:
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        bak = os.path.join(FIN, "backups", "%s_%s_before_fix_standalone_wo_i.xlsx.bak" % (tag, ts))
        shutil.copyfile(path, bak)
        tmp = path + ".fixwoi.tmp"
        wb.save(tmp)
        os.replace(tmp, path)
        log("[%s] 替换 %d 处，备份 -> %s" % (tag, len(hits), os.path.basename(bak)))
        for coord in hits:
            log("   %s : 我 -> I" % coord)
    else:
        log("[%s] 未发现单独的「我」" % tag)
    wb.close()
    return len(hits)


log("=== 精确替换单独「我」→「I」 @ %s ===" % datetime.datetime.now().strftime("%Y%m%d_%H%M%S"))
total = 0
for fn, tag in (("vn_report.xlsx", "vn"), ("merge_report.xlsx", "merge"), ("cn_report.xlsx", "cn")):
    p = os.path.join(FIN, fn)
    if os.path.exists(p):
        total += replace_standalone_wo(p, tag)
    else:
        log("[%s] 文件不存在 %s" % (tag, p))

if total:
    for key in ("vn_report.snapshot.key", "merge_report.snapshot.key", "cn_report.snapshot.key"):
        p = os.path.join(FIN, key)
        try:
            os.remove(p)
            log("删快照键 %s" % key)
        except OSError:
            pass

log("共替换 %d 处。done" % total)
