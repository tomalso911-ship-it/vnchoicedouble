#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""中国报表数据安全审计：现状数据量 + 备份/回滚点是否够用。"""
import os, sys, glob
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BASE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(BASE, "finance_data")
CN = os.path.join(FIN, "cn_report.xlsx")
BK = os.path.join(FIN, "backups")

import openpyxl

print("=" * 70)
print("CN_REPORT =", CN)
st = os.stat(CN)
print("  size=%d  mtime=%.19s" % (st.st_size, __import__("datetime").datetime.fromtimestamp(st.st_mtime)))
print("  rev = %d-%d" % (int(st.st_mtime_ns), int(st.st_size)))

wb = openpyxl.load_workbook(CN, data_only=False)
ws = wb[wb.sheetnames[0]]
nonempty = 0
rows_nonempty = 0
maxr = 0
maxc = 0
for r in range(1, ws.max_row + 1):
    has = False
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v is not None and str(v).strip() != "":
            nonempty += 1
            has = True
            if r > maxr: maxr = r
            if c > maxc: maxc = c
    if has:
        rows_nonempty += 1
print("  非空单元格 = %d   有数据行 = %d   max_row=%d max_col=%d" % (nonempty, rows_nonempty, ws.max_row, ws.max_column))
wb.close()

print("-" * 70)
print("备份目录:", BK, "(存在)" if os.path.isdir(BK) else "(不存在)")
if os.path.isdir(BK):
    for pat in ("cn_report_*.xlsx.bak", "cn_report_*.bak", "cn_*.bak"):
        files = sorted(glob.glob(os.path.join(BK, pat)), key=os.path.getmtime)
        print("  %-26s 共 %d 份" % (pat, len(files)))
        for f in files[-6:]:
            s = os.stat(f)
            import datetime as _d
            print("      %s  %8d B  %s" % (os.path.basename(f), s.st_size,
                  _d.datetime.fromtimestamp(s.st_mtime).strftime("%Y-%m-%d %H:%M:%S")))
print("-" * 70)
print("archive 目录:", os.path.join(FIN, "archive"))
ar = os.path.join(FIN, "archive")
if os.path.isdir(ar):
    fs = sorted(glob.glob(os.path.join(ar, "cn_*")), key=os.path.getmtime)
    print("  cn_* 共 %d 份" % len(fs))
    for f in fs[-6:]:
        print("      %s" % os.path.basename(f))
print("=" * 70)
