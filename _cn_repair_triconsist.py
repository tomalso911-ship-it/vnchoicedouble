# -*- coding: utf-8 -*-
"""★一次性数据修复（FIN-CN-TRI-CONSIST-20261013）：
把 cn_report.xlsx 被「人民币换算值」污染的【越南文基准块 D/E】按【中文块 N/O】(与英文块 X/Y 校验)
回填成真实越南盾，使三语块重新同值。

判据（与 app.py::_cn_block_total 同一口径）：三语块存的都是同一个越南盾数值，
基准块只是被显示层的换算值覆盖了（41,400,000 → 10,664.14 → 2.75）。
只改【数值不一致】的金额格；文字/日期/公式/样式/核销列一律不动。
"""
import datetime
import io
import os
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
CN = os.path.join(FIN, "cn_report.xlsx")
OUT = os.path.join(HERE, "_cn_repair_triconsist_out.txt")
L = []


def log(s):
    L.append(str(s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(s)


def total(ws, cols):
    t = 0.0
    for r in range(5, 205):
        for c in cols:
            v = ws.cell(r, c).value
            if isinstance(v, bool):
                continue
            if isinstance(v, (int, float)):
                t += abs(float(v))
    return round(t, 2)


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.replace(",", "").replace("\xa0", "").strip()
        try:
            return float(s)
        except ValueError:
            return None
    return None


wb = openpyxl.load_workbook(CN, data_only=False)
ws = wb[wb.sheetnames[0]]
log("工作表 = %s" % wb.sheetnames)
t_before = (total(ws, (4, 5)), total(ws, (14, 15)), total(ws, (24, 25)))
log("修复前：基准块D+E=%s  中文块N+O=%s  英文块X+Y=%s" % t_before)

bak = os.path.join(FIN, "backups",
                   "cn_report_%s_before_triconsist_repair.xlsx.bak"
                   % datetime.datetime.now().strftime("%Y%m%d_%H%M%S"))
shutil.copyfile(CN, bak)
log("已备份 -> %s" % bak)

fixed = []
skipped = []
for r in range(5, 205):
    for bc, zc, ec in ((4, 14, 24), (5, 15, 25)):
        bv = num(ws.cell(r, bc).value)
        zv = num(ws.cell(r, zc).value)
        ev = num(ws.cell(r, ec).value)
        if bv is None:
            continue
        ref = None
        if zv is not None and ev is not None:
            if abs(zv - ev) > max(abs(zv) * 0.001, 0.01):
                skipped.append("行%d 列%d：中英块本身不一致 zh=%s en=%s" % (r, bc, zv, ev))
                continue
            ref = zv
        elif zv is not None:
            ref = zv
        elif ev is not None:
            ref = ev
        if ref is None:
            continue
        if abs(bv - ref) <= max(abs(ref) * 0.001, 0.01):
            continue
        cell = ws.cell(r, bc)
        fmt = cell.number_format
        cell.value = ref
        try:
            cell.number_format = fmt
        except Exception:
            pass
        fixed.append("行%d 列%d：%s → %s" % (r, bc, bv, ref))

log("改动的格数 = %d" % len(fixed))
for s in fixed[:40]:
    log("   " + s)
if len(fixed) > 40:
    log("   ...（其余 %d 条）" % (len(fixed) - 40))
for s in skipped[:20]:
    log("   跳过 " + s)

if fixed:
    wb.save(CN)
    log("已写盘")
else:
    log("无需改动")

wb2 = openpyxl.load_workbook(CN, data_only=False)
ws2 = wb2[wb2.sheetnames[0]]
log("修复后：基准块D+E=%s  中文块N+O=%s  英文块X+Y=%s"
    % (total(ws2, (4, 5)), total(ws2, (14, 15)), total(ws2, (24, 25))))
log("抽查 E5=%r E8=%r  O5=%r  Y5=%r  AC7=%r"
    % (ws2["E5"].value, ws2["E8"].value, ws2["O5"].value, ws2["Y5"].value, ws2["AC7"].value))
