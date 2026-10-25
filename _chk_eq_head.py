import openpyxl, sys
sys.stdout.reconfigure(encoding="utf-8")
wb = openpyxl.load_workbook("finance_data/equity_report.xlsx")
print("sheets:", wb.sheetnames)
ws = wb["Tổng đầu tư"]
print("dims:", ws.dimensions)
for r in range(1, 6):
    print(r, [ws.cell(row=r, column=c).value for c in range(1, 7)])
