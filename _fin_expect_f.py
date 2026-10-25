# -*- coding: utf-8 -*-
"""源 xlsx 里 F 列的【期望值】（data_only = 上次 Excel 保存时算出的缓存值）作为对照基准。"""
import io, sys
from openpyxl import load_workbook
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
wbf = load_workbook('finance_data/vn_report.xlsx', data_only=False)
wbv = load_workbook('finance_data/vn_report.xlsx', data_only=True)
wsf, wsv = wbf['Chi tiết thu-chi'], wbv['Chi tiết thu-chi']
print('1基行 | F公式 | F缓存值')
for r in (4, 5, 6, 7, 8, 9, 79, 80, 81, 293, 294, 295, 300, 301, 302, 303, 304, 569):
    print('  F%-4d | %-24r | %r' % (r, wsf.cell(row=r, column=6).value, wsv.cell(row=r, column=6).value))
