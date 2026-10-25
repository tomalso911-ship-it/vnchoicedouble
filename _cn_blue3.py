# -*- coding: utf-8 -*-
"""只读验证（改的是临时副本，绝不动 finance_data/cn_report.xlsx）：
   ① 三种语言块（A/B/C、K/L/M、U/V/W）的数据区是不是都已经是蓝色；
   ② 合计行（第 205 行）没有被染蓝；
   ③ 新的 _cn_normalize_style() 是幂等的（第二次返回 False、文件 mtime 不变）。"""
import io
import os
import shutil
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import app  # noqa: E402
import openpyxl  # noqa: E402

TMP = tempfile.mkdtemp(prefix="cnblue3_")
SRC = os.path.join(HERE, "finance_data", "cn_report.xlsx")
DST = os.path.join(TMP, "cn_report.xlsx")
shutil.copyfile(SRC, DST)
app.CN_REPORT = DST
app.CN_TITLES_PATH = os.path.join(TMP, "fin_sheet_titles_cn.json")
app.CN_SNAP_KEY = os.path.join(TMP, "cn_report.snapshot.key")
app.CN_SNAP_PREFIX = "cn_report.snapshot."
app.FIN_DIR = TMP

WANT = "0000FF"


def color(c):
    v = getattr(getattr(c.font, "color", None), "rgb", None)
    return str(v).upper() if isinstance(v, str) else None


def isblue(c):
    v = color(c)
    return bool(v and v[-6:] == WANT)


def report(tag):
    wb = openpyxl.load_workbook(DST)
    ws = wb[wb.sheetnames[0]]
    print("")
    print("---- %s (sheets=%s dims=%s max_row=%d) ----" % (tag, wb.sheetnames, ws.dimensions, ws.max_row))
    for r in (5, 6, 100, 204):
        cells = ["A", "B", "C", "D", "E", "K", "L", "M", "N", "O", "U", "V", "W", "X", "Y", "AC"]
        flags = " ".join("%s=%s" % (col, "BLUE" if isblue(ws["%s%d" % (col, r)]) else str(color(ws["%s%d" % (col, r)])))
                         for col in cells)
        print("   行%-4d %s" % (r, flags))
    print("   行205(合计) C=%r blue=%s | E=%r blue=%s | M=%r blue=%s | W=%r blue=%s" % (
        ws["C205"].value, isblue(ws["C205"]), ws["E205"].value, isblue(ws["E205"]),
        ws["M205"].value, isblue(ws["M205"]), ws["W205"].value, isblue(ws["W205"])))
    # 统计
    n = 0
    nonblue = []
    for row in ws.iter_rows(min_row=5, max_row=204, min_col=1, max_col=28):
        for c in row:
            n += 1
            if not isblue(c):
                nonblue.append(c.coordinate)
    print("   数据区格子=%d 非蓝=%d %s" % (n, len(nonblue), nonblue[:8]))


report("改之前")

m0 = os.stat(DST).st_mtime_ns
r1 = app._cn_normalize_style()
m1 = os.stat(DST).st_mtime_ns
print("")
print("第 1 次 _cn_normalize_style() -> %s ; mtime 变了=%s" % (r1, m0 != m1))
r2 = app._cn_normalize_style()
m2 = os.stat(DST).st_mtime_ns
print("第 2 次 _cn_normalize_style() -> %s ; mtime 变了=%s  (幂等应为 False/False)" % (r2, m1 != m2))

report("改之后")

try:
    shutil.rmtree(TMP, ignore_errors=True)
except Exception:
    pass
print("")
print("临时目录已清理")
