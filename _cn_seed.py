# -*- coding: utf-8 -*-
"""临时：生成中国报表种子文件 + 冒烟自检（结果写 _cn_seed_out.txt）。"""
import io
import os
import sys
import traceback

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

out = []


def log(s):
    out.append(str(s))


try:
    import app
    log("import app: OK")
    log("CN_REPORT = %s" % app.CN_REPORT)
    # 重新生成：先删掉旧文件与旧快照，确保用的是最新种子结构
    for p in (app.CN_REPORT, app.CN_TITLES_PATH):
        if os.path.isfile(p):
            os.remove(p)
    app._cn_cache_invalidate()
    app._cn_seed_if_missing()
    log("seed done; exists=%s size=%s" % (os.path.isfile(app.CN_REPORT),
                                          os.path.getsize(app.CN_REPORT) if os.path.isfile(app.CN_REPORT) else -1))
    log("titles file exists=%s" % os.path.isfile(app.CN_TITLES_PATH))

    ok, detail = app._cn_cache_build(force=True)
    log("cache build -> ok=%s detail=%s" % (ok, detail))
    log("cache current -> %r" % (app._cn_cache_current(),))

    import openpyxl
    wb = openpyxl.load_workbook(app.CN_REPORT)
    ws = wb[app.CN_SHEET_NAME] if app.CN_SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
    log("sheets=%r dims=%s max_row=%d" % (wb.sheetnames, ws.dimensions, ws.max_row))
    log("AC4=%r AC5=%r A4=%r C205/E205=%r/%r" % (ws["AC4"].value, ws["AC5"].value, ws["A4"].value,
                                                 ws["C205"].value, ws["E205"].value))

    # 快照冒烟：确认 xlsx→Univer 快照能跑通，且三语块都在
    import univer_convert
    snap = univer_convert.xlsx_to_univer(app.CN_REPORT)
    sid = snap["sheetOrder"][0]
    log("snapshot sheet=%r columnCount=%s styles=%s" % (snap["sheets"][sid]["name"],
                                                        snap["sheets"][sid].get("columnCount"),
                                                        len(snap.get("styles") or {})))
    with io.open(app.CN_TITLES_PATH, encoding="utf-8") as fh:
        log("cn titles: %s" % fh.read().strip())
except Exception as e:
    log("ERR %s" % e)
    log(traceback.format_exc())

io.open(os.path.join(HERE, "_cn_seed_out.txt"), "w", encoding="utf-8").write("\n".join(out) + "\n")
print("done")
