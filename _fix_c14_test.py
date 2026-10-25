# -*- coding: utf-8 -*-
"""验收 FIN-TRANS-MISS-20261013：用【用户真实的 cn_report.xlsx 的临时副本】复现并验证修复。
   （真实文件只读复制，全程不写回 finance_data/）

   预期：切到中国报表、看 zh 视图时，第 14 行摘要从越南文
   'Thanh toán tiền vật chất cho Công ty Huifeng' 自动翻成 '向惠丰公司支付物资款项'；
   并最终被自动保存写进文件的中文块 M14。
"""
import datetime
import io
import os
import shutil
import socket
import sys
import tempfile
import threading
import time

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "_fix_c14_out.txt")
LOG = []
VI = "Thanh toán tiền vật chất cho Công ty Huifeng"
ZH = "向惠丰公司支付物资款项"


def log(s):
    LOG.append("[%s] %s" % (datetime.datetime.now().strftime("%H:%M:%S"), s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(LOG) + "\n")


def port_open(port):
    s = socket.socket()
    s.settimeout(0.6)
    try:
        s.connect(("127.0.0.1", port))
        return True
    except Exception:
        return False
    finally:
        s.close()


import app  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402

TMP = tempfile.mkdtemp(prefix="fixc14_")
REAL = os.path.join(HERE, "finance_data", "cn_report.xlsx")
app.FIN_DIR = TMP
app.CN_REPORT = os.path.join(TMP, "cn_report.xlsx")
app.CN_SNAP_KEY = os.path.join(TMP, "cn_report.snapshot.key")
app.CN_SNAP_PREFIX = "cn_report.snapshot."
app.CN_TITLES_PATH = os.path.join(TMP, "fin_sheet_titles_cn.json")
shutil.copy(REAL, app.CN_REPORT)                      # ★用真实报表的副本
_sh = os.path.join(HERE, "finance_data", "fin_sheet_titles_cn.json")
if os.path.isfile(_sh):
    shutil.copy(_sh, app.CN_TITLES_PATH)
log("临时报表（真实文件副本）= %s" % app.CN_REPORT)

PORT = 5173                                            # 换个端口，避开历史遗留的服务进程
BASE = "http://127.0.0.1:%d/lite" % PORT
if not port_open(PORT):
    threading.Thread(target=lambda: app.app.run(host="127.0.0.1", port=PORT, debug=False,
                                               use_reloader=False, threaded=True),
                     daemon=True).start()
    for _ in range(60):
        if port_open(PORT):
            break
        time.sleep(0.3)

CELL = """(p) => {
  try {
    var sn = fin_currentSnapshot();
    var sid = sn.sheetOrder[0];
    var cc = ((sn.sheets[sid].cellData || {})[String(p.r)] || {})[String(p.c)];
    return cc ? cc.v : null;
  } catch(e){ return 'ERR:' + e.message; }
}"""

errors = []
tr_reqs = []
with sync_playwright() as p:
    b = p.chromium.launch(channel="chrome", headless=True, args=["--no-sandbox"])
    pg = b.new_page(viewport={"width": 1500, "height": 950})
    pg.on("pageerror", lambda e: errors.append("pageerror: " + str(e)[:200]))
    pg.on("console", lambda m: errors.append("console.error: " + m.text[:200]) if m.type == "error" else None)
    pg.on("request", lambda r: tr_reqs.append(r.url) if "/api/translate?" in r.url else None)

    pg.goto(BASE, wait_until="domcontentloaded")
    pg.wait_for_timeout(1200)
    pg.fill("#loginUser", "tom")
    pg.fill("#loginPwd", "66668888")
    pg.click("#loginBtn")
    pg.wait_for_selector("#topbar:not(.hidden)", timeout=25000)
    pg.evaluate("() => { try { switchView('finance'); } catch(e){} }")
    pg.wait_for_timeout(800)
    pg.evaluate("() => { try { setLangDebug('zh'); } catch(e){} }")   # 与用户截图一致：中文界面
    pg.wait_for_timeout(1500)

    log("")
    log("① 切到中国报表（行 14 = 0 基 r13，摘要 = 0 基 c2）")
    pg.evaluate("() => { var t = document.querySelector('#fin-tabs .fin-tab[data-key=\\'cn\\']'); if (t) t.click(); }")
    for _ in range(70):
        pg.wait_for_timeout(400)
        if pg.evaluate("() => (typeof finReady !== 'undefined') && !!finReady"):
            break
    pg.wait_for_timeout(3000)
    got = pg.evaluate(CELL, {"r": 13, "c": 2})
    log("   刚载入时 C14 = %r" % got)

    log("")
    log("② 等自动翻译 + 屏幕补写（最多 25s）")
    ok = False
    for _ in range(100):
        pg.wait_for_timeout(250)
        got = pg.evaluate(CELL, {"r": 13, "c": 2})
        if isinstance(got, str) and got.strip() == ZH:
            ok = True
            break
    log("   C14 = %r" % got)
    log("   ★屏幕是否已变成中文？ %s" % ("是 ✔" if ok else "否 ✗（BUG 没修好）"))

    log("")
    log("③ 是否真的发了翻译请求（/api/translate 含该越南文）")
    hit = [u for u in tr_reqs if "Huifeng" in u]
    log("   总翻译请求 %d 条；命中该句 %d 条" % (len(tr_reqs), len(hit)))
    for u in hit[:4]:
        log("     %s" % u[:150])

    log("")
    log("④ 等自动保存把译文写回文件（最多 25s）")
    import openpyxl
    m14 = None
    for _ in range(100):
        pg.wait_for_timeout(250)
        try:
            wb = openpyxl.load_workbook(app.CN_REPORT, data_only=True)
            ws = wb[wb.sheetnames[0]]
            m14 = ws["M14"].value
            w14 = ws["W14"].value
            wb.close()
        except Exception as e:
            log("   读文件失败：%s" % e)
            break
        if m14:
            break
    log("   文件里 M14（中文块）= %r" % m14)
    log("   文件里 W14（英文块）= %r" % (w14 if 'w14' in dir() else None))
    log("   ★中文是否已落盘？ %s" % ("是 ✔" if m14 else "否（自动保存可能还没轮到）"))

    log("")
    log("页面报错 = %s" % errors[:6])
    pg.screenshot(path=os.path.join(HERE, "_fix_c14_after.png"))
    b.close()

shutil.rmtree(TMP, ignore_errors=True)
log("临时目录已清理（真实文件未改动）")
print("\n".join(LOG))
