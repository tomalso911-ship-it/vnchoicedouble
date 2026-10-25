# -*- coding: utf-8 -*-
"""扫描 cn_report.xlsx / merge_report.xlsx：找出缓存值为 Excel 错误的格及其公式。"""
import io, sys, os, glob
from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
ERRS = ('#VALUE!', '#REF!', '#DIV/0!', '#N/A', '#NAME?', '#NULL!', '#NUM!', '#ERROR!')
MAXC = 32

files = [p for p in ['finance_data/cn_report.xlsx', 'finance_data/merge_report.xlsx'] if os.path.exists(p)]
files += glob.glob('finance_data/*report.xlsx')
seen = []
for p in files:
    if p not in seen:
        seen.append(p)
print('待查文件:', seen)

for path in seen:
    print('\n########## %s ##########' % path)
    try:
        wbf = load_workbook(path, data_only=False, read_only=True)
        wbv = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        print('  打不开:', e)
        continue
    print('  工作表:', wbf.sheetnames)
    for name in wbf.sheetnames:
        fmap = {}
        for row in wbf[name].iter_rows(min_col=1, max_col=MAXC):
            for c in row:
                if c.value is not None:
                    fmap[c.coordinate] = c.value
        hits = []
        for row in wbv[name].iter_rows(min_col=1, max_col=MAXC):
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ERRS:
                    hits.append((c.coordinate, fmap.get(c.coordinate)))
        if hits:
            print('  == %-22s 错误格 %d 个' % (name, len(hits)))
            for coord, f in hits[:25]:
                print('      %-6s 公式 = %r' % (coord, f))
    print('  (完)')
