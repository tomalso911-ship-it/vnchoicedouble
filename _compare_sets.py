# -*- coding: utf-8 -*-
"""只读验证：当前 vn_report.xlsx「Chi thực tế」的记录集合
   ==  before_merge 备份的记录集合  ∪  {11 行中国记录}  ？

若「当前有、备份无」的恰好只有那 11 行中国记录，且「备份有、当前无」为空，
则说明【还原只会去掉中国行，不会抹掉任何别的数据】。
"""
import datetime
import io
import os
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
VN = os.path.join(FIN, "vn_report.xlsx")
BAK = os.path.join(FIN, "backups", "vn_report_20261014_055253_before_merge.xlsx.bak")
SHEET = "Chi th\u1ef1c t\u1ebf"
MARK_COL = 40
# 与 _merge_cn_into_vn.py 一致：只看三块真正的数据列
DATA_COLS = list(range(1, 6)) + list(range(11, 16)) + list(range(21, 26))
OUT = os.path.join(HERE, "_compare_sets_out.txt")
L = []


def log(s):
    L.append(str(s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(s)


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return round(v, 2)
    return str(v).strip()


def rows_of(path):
    probe = None
    src = path
    if not path.lower().endswith(".xlsx"):
        probe = path + ".probe.xlsx"
        shutil.copyfile(path, probe)
        src = probe
    wb = openpyxl.load_workbook(src, data_only=False)
    ws = wb[SHEET]
    recs = []
    for r in range(5, 420):
        vals = [ws.cell(r, c).value for c in DATA_COLS]
        joined = " ".join(str(v) for v in vals if v is not None)
        if not any(str(v).strip() for v in vals if v is not None):
            continue
        if "T\u1ed4NG C\u1ed8NG" in joined or any(str(v).startswith("=SUM") for v in vals if isinstance(v, str)):
            continue   # 合计行不算记录
        mk = ws.cell(r, MARK_COL).value
        tag = "CN" if (isinstance(mk, str) and mk.startswith("CN:")) else "VN"
        recs.append((
            tuple(norm(ws.cell(r, c).value) for c in DATA_COLS),
            tag,
            r,
            norm(ws.cell(r, 1).value),
            norm(ws.cell(r, 2).value),
            norm(ws.cell(r, 5).value),
        ))
    wb.close()
    if probe:
        try:
            os.remove(probe)
        except OSError:
            pass
    return recs


A = rows_of(VN)
B = rows_of(BAK)
log("当前 vn_report.xlsx「Chi thực tế」记录数 = %d（其中 CN 标记 %d）"
    % (len(A), sum(1 for x in A if x[1] == "CN")))
log("备份 055253        记录数 = %d" % len(B))
log("")

setA = set(x[0] for x in A)
setB = set(x[0] for x in B)
onlyA = [x for x in A if x[0] not in setB]
onlyB = [x for x in B if x[0] not in setA]

log("【当前有、备份没有】的记录 = %d 条" % len(onlyA))
for x in onlyA:
    log("    [%s] 日期=%s 凭证=%s 金额=%s" % (x[1], x[3], x[4], x[5]))
log("")
log("【备份有、当前没有】的记录 = %d 条" % len(onlyB))
for x in onlyB:
    log("    [%s] 日期=%s 凭证=%s 金额=%s" % (x[1], x[3], x[4], x[5]))
log("")

cn_only = all(x[1] == "CN" for x in onlyA)
log("结论：")
log("  · 当前比备份多出的记录，是否全是中国记录(CN) ：%s" % ("是" if cn_only else "否 ← 有越南记录会被抹掉！"))
log("  · 备份里有、当前却缺失的记录条数              ：%d（期望 0）" % len(onlyB))
log("  → 还原是否安全：%s" % ("安全（只去掉中国行）" if (cn_only and not onlyB) else "不安全，需人工检查"))
log("done")
