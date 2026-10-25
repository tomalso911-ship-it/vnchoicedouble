# -*- coding: utf-8 -*-
import os, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl
FIN = os.path.join(os.path.dirname(os.path.abspath(__file__)), "finance_data")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_split_sum_out.txt")
L = []
def log(s=""):
    L.append(str(s)); open(OUT,"w",encoding="utf-8").write("\n".join(L)+"\n"); print(s)

def load(fn): return openpyxl.load_workbook(os.path.join(FIN, fn), data_only=True)

def vn_sum():
    ws = load("vn_report.xlsx")["Chi thực tế"]
    t = 0.0; n = 0
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 5).value
        if isinstance(v, (int, float)) and r != 3:
            lbl = ws.cell(r, 3).value
            if isinstance(lbl, str) and ("TỔNG" in lbl or "ổng" in lbl): continue
            t += v; n += 1
    log("vn  Chi thực tế E 数据求和 = %s （%d 行）" % (t, n))
    return t

def merge_split():
    ws = load("merge_report.xlsx")["Chi thực tế"]
    cn = 0.0; vn = 0.0; ncn = 0; nvn = 0; other = 0.0
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 5).value
        if not isinstance(v, (int, float)) or r == 3: continue
        lbl = ws.cell(r, 3).value
        if isinstance(lbl, str) and ("TỔNG" in lbl or "ổng" in lbl): continue
        mk = ws.cell(r, 40).value   # MARK_COL
        if isinstance(mk, str) and mk.startswith("CN:"):
            cn += v; ncn += 1
        else:
            vn += v; nvn += 1
    log("merge Chi thực tế  E: 中国行(n=%d) 合计=%s | 越南行(n=%d) 合计=%s | 合计=%s"
        % (ncn, cn, nvn, vn, cn + vn))
    log("  => 越南行合计 - vn单独合计 的差额 = %s" % (vn - vn_sum()))

vn_sum()
merge_split()
log("done")
