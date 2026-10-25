# -*- coding: utf-8 -*-
"""只读验证：修复后的 cn_report.xlsx —— 金额是数字+¥格式、三块语言正确、网格线齐全。"""
import io
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

OUT = io.open("c:/vnchoicedouble/_verify_cn_repair_out.txt", "w", encoding="utf-8", newline="\n")


def log(s):
    OUT.write(str(s) + "\n")
    print(s)


CJK = lambda s: any("\u3400" <= ch <= "\u9fff" for ch in str(s or ""))

wb = openpyxl.load_workbook("c:/vnchoicedouble/finance_data/cn_report.xlsx")
ws = wb[wb.sheetnames[0]]
log("工作表 = %s  尺寸 = %s" % (wb.sheetnames, ws.dimensions))

log("-" * 60)
log("① 金额列：类型(n=数字/s=文本) 与数字格式")
bad = 0
for coord in ("D5", "E5", "D6", "E6", "N5", "O5", "X5", "Y5", "E15", "E205"):
    c = ws[coord]
    ok = (c.data_type == "n" or str(c.value or "").startswith("="))
    fmtok = (str(c.number_format or "").strip() == '"¥"#,##0')
    if not (ok and fmtok):
        bad += 1
    log("   %-5s type=%-3s fmt=%-12r v=%-22r %s" % (coord, c.data_type, c.number_format, c.value,
                                                   "" if (ok and fmtok) else "  ← 有问题"))
log("   金额列不合格格数 = %d（期望 0）" % bad)
# 整列扫一遍
tot = 0
for r in range(5, 206):
    for col in "DENOXY":
        cc = ws["%s%d" % (col, r)]
        if cc.data_type == "s":
            tot += 1
log("   行5-205 的 D/E/N/O/X/Y 里仍是【文本】的格 = %d（期望 0）" % tot)

log("-" * 60)
log("② 三语块语言检查（行 5-60）")
zh_in_en = zh_in_vi = 0
for r in range(5, 61):
    for c in range(21, 26):
        v = ws.cell(r, c).value
        if isinstance(v, str) and CJK(v):
            zh_in_en += 1
            log("     英文块仍有中文: %s = %r" % (ws.cell(r, c).coordinate, v))
    for c in range(1, 6):
        v = ws.cell(r, c).value
        if isinstance(v, str) and CJK(v):
            zh_in_vi += 1
            log("     越南文块仍有中文: %s = %r" % (ws.cell(r, c).coordinate, v))
log("   英文块含中文 = %d（期望 0）；越南文块含中文 = %d（期望 0）" % (zh_in_en, zh_in_vi))
log("   样例 r5: C=%r | M=%r | W=%r | F=%r | P=%r | Z=%r"
    % (ws["C5"].value, ws["M5"].value, ws["W5"].value, ws["F5"].value, ws["P5"].value, ws["Z5"].value))

log("-" * 60)
log("③ 黑色网格线（四边）")
lack = []
for coord in ("A1", "E1", "A2", "A4", "B4", "C4", "D4", "E4", "F4", "G4", "H4", "I4", "J4",
              "A5", "F5", "G5", "H5", "I5", "J5", "AB5", "AC5", "A205", "Y205", "AC205"):
    b = ws[coord].border
    sides = [k for k in ("left", "right", "top", "bottom") if getattr(b, k, None) and getattr(b, k).style]
    if len(sides) < 4:
        lack.append("%s(%d边)" % (coord, len(sides)))
log("   抽查 24 格，四边不全的 = %d %s" % (len(lack), lack[:8]))

OUT.close()
