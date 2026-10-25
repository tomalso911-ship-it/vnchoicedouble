# -*- coding: utf-8 -*-
"""修正中国报表 cn_report.xlsx 第 9 行（凭证 UNC00056 / 9,688,000 VND）的摘要。

现状：三语摘要被占位符「实际支出」顶替（越 Chi tiêu thực tế / 英 Actual expenditure），
      不是这笔业务的真实摘要。之前的 _fix_cn_desc.py 把它当成"已有摘要"跳过了。

原始出处：
  · FinanceReport/Choice Slats VN - Financial Report.xlsx  → Chi thực tế r63
      C = "Chi tiền cho CÔNG TY TNHH MTV XUÂN HÒA HT"  E = 9,688,000
  · finance_data/vn_report.xlsx → Chi thực tế r103（同笔，同越南文）

按与行10/行11/行14 完全相同的「报表口径」（_polish_cn_desc.py 确立的润色规则：
Công ty TNHH <lĩnh vực> <thương hiệu> / <品牌名>有限公司 / <Brand> ... Co., Ltd.）落地。

只改文字；金额/日期/凭证号/核销列一律不碰。先备份，再原子落盘，最后失效快照缓存。
用法: python _cn_fix_row9.py [--dry-run]
"""
import datetime
import glob
import os
import shutil
import sys

import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
CN = os.path.join(BASE, "finance_data", "cn_report.xlsx")
BKD = os.path.join(BASE, "finance_data", "backups")

ROW = 9
NEW = {
    3: "Công ty TNHH MTV Xuân Hòa HT",          # C 摘要 - 越南文
    13: "XUÂN HÒA HT 有限公司",                  # M 摘要 - 中文
    23: "Xuan Hoa HT Co., Ltd.",                 # W 摘要 - 英文
}
# 期望的当前值（防误改：不符就只报告不落盘）
EXPECT = {3: "Chi tiêu thực tế", 13: "实际支出", 23: "Actual expenditure"}
AMT_COLS = (5, 15, 25)


def total(ws):
    t = 0
    for r in range(5, 205):
        for c in AMT_COLS:
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                t += v
    return t


def main():
    dry = "--dry-run" in sys.argv
    wb = openpyxl.load_workbook(CN, data_only=False)
    ws = wb.worksheets[0]

    print("凭证行 %d: 日期=%s 凭证=%s 金额(E)=%s" % (
        ROW, ws.cell(ROW, 1).value, ws.cell(ROW, 2).value, ws.cell(ROW, 5).value))
    todo = []
    for c, new in NEW.items():
        cur = ws.cell(ROW, c).value
        cl = ws.cell(ROW, c).column_letter
        print("  %s: 当前=%r" % (cl, cur))
        if cur == new:
            print("    已是目标值，跳过")
            continue
        if str(cur or "").strip() != EXPECT[c]:
            print("    !!! 当前值与预期 %r 不符，跳过（可能已被手动改过）" % EXPECT[c])
            continue
        todo.append((c, cur, new))

    print("待改格数:", len(todo))
    if dry or not todo:
        print("(dry-run，未修改文件)" if dry else "(无可改，未落盘)")
        return

    t_before = total(ws)
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = os.path.join(BKD, "cn_report_%s_before_row9fix.xlsx.bak" % ts)
    shutil.copyfile(CN, bak)
    print("已备份 ->", os.path.basename(bak))

    for c, _old, new in todo:
        cell = ws.cell(ROW, c)
        cell.value = new
        if cell.number_format not in ("General", "@"):
            cell.number_format = "General"

    tmp = CN + ".row9.tmp"
    wb.save(tmp)
    wb.close()
    os.replace(tmp, CN)
    print("已原子落盘 ->", CN)

    n = 0
    for p in glob.glob(os.path.join(BASE, "finance_data", "cn_report.snapshot.*")) + \
             [os.path.join(BASE, "finance_data", "cn_report.snapshot.key")]:
        try:
            os.remove(p)
            n += 1
        except OSError:
            pass
    print("已失效快照缓存文件数:", n)

    ws2 = openpyxl.load_workbook(CN, data_only=False).worksheets[0]
    t_after = total(ws2)
    print("=== 复核 ===")
    for c, cl in ((3, "C"), (13, "M"), (23, "W")):
        print("  r%d%s = %r" % (ROW, cl, ws2.cell(ROW, c).value))
    print("  支出合计(E/O/Y) 改前=%s 改后=%s %s" % (
        format(t_before, ","), format(t_after, ","),
        "一致 ✓" if t_before == t_after else "!!! 不一致 !!!"))


if __name__ == "__main__":
    main()
