# -*- coding: utf-8 -*-
"""中国报表后端接口冒烟测试（结果写 _cn_api_test_out.txt）。
覆盖：越南链路未受影响 / 中国快照 / 中国表名 / 手动保存往返（含核销 ☑）/ 备份 / 只有一张表。"""
import io
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

out = []
fails = []


def log(s):
    out.append(str(s))


def ck(name, cond, extra=""):
    log(("  OK   " if cond else "  FAIL ") + name + (" | " + str(extra) if extra else ""))
    if not cond:
        fails.append(name)


import app  # noqa: E402

cli = app.app.test_client()

log("=== 1) 越南链路（改动不该影响它）")
r = cli.get("/api/finance/vn-report/snapshot")
ck("VN snapshot 200", r.status_code == 200, r.status_code)
jv = r.get_json() or {}
ck("VN snapshot id", jv.get("id") == "vn-report", jv.get("id"))
ck("VN 8 张表", len(jv.get("sheets") or {}) == 8, list((jv.get("sheets") or {}).values())[:2])
r = cli.get("/api/finance/sheet-titles")
jt = (r.get_json() or {}).get("titles") or {}
ck("VN 表名仍含 Thu Dabaco", "Thu Dabaco" in jt, list(jt.keys())[:3])
ck("VN 表名 zh 仍含 Dabaco", "dabaco" in (jt.get("Thu Dabaco", {}).get("zh", "").lower()),
   jt.get("Thu Dabaco"))

log("=== 2) 中国报表：快照")
r = cli.get("/api/finance/cn-report/snapshot")
ck("CN snapshot 200", r.status_code == 200, r.status_code)
jc = r.get_json() or {}
ck("CN snapshot id=cn-report", jc.get("id") == "cn-report", jc.get("id"))
sids = list((jc.get("sheets") or {}).keys())
ck("CN 只有一张表", len(sids) == 1, sids)
sid = sids[0] if sids else None
sh = (jc.get("sheets") or {}).get(sid, {})
ck("CN 工作表名=实际支出", sh.get("name") == "\u5b9e\u9645\u652f\u51fa", sh.get("name"))
cd = sh.get("cellData") or {}
ck("CN AC4 = 核销", (cd.get("3", {}).get("28", {}) or {}).get("v") == "\u6838\u9500",
   (cd.get("3", {}).get("28", {}) or {}).get("v"))
ck("CN AC5 = ☐", (cd.get("4", {}).get("28", {}) or {}).get("v") == "\u2610",
   (cd.get("4", {}).get("28", {}) or {}).get("v"))
ck("CN 三语表头都在", all((cd.get("3", {}).get(str(c), {}) or {}).get("v") for c in (0, 10, 20)),
   [(cd.get("3", {}).get(str(c), {}) or {}).get("v") for c in (0, 10, 20)])
ck("CN 合计行公式在", "=SUM" in str((cd.get("204", {}).get("4", {}) or {}).get("f", "")),
   (cd.get("204", {}).get("4", {}) or {}).get("f"))

log("=== 3) 中国报表：表名三语")
r = cli.get("/api/finance/sheet-titles?kind=cn")
jtc = (r.get_json() or {}).get("titles") or {}
ck("CN 表名文件有实际支出", "\u5b9e\u9645\u652f\u51fa" in jtc, jtc)
ck("CN 表名 zh=实际支出", jtc.get("\u5b9e\u9645\u652f\u51fa", {}).get("zh") == "\u5b9e\u9645\u652f\u51fa", jtc)
ck("CN 表名 vi=Chi thực tế", jtc.get("\u5b9e\u9645\u652f\u51fa", {}).get("vi") == "Chi th\u1ef1c t\u1ebf", jtc)
ck("CN 表名不影响越南那份", "Thu Dabaco" in ((cli.get("/api/finance/sheet-titles").get_json() or {}).get("titles") or {}))

log("=== 4) 手动保存往返（含核销 ☑）")
# 模拟前端 finCnBuildSavePayload 的结果：在快照上写字 + 勾核销
import copy  # noqa: E402
pay = copy.deepcopy(jc)
pcd = pay["sheets"][sid]["cellData"]
pcd.setdefault("4", {})["0"] = {"v": 46000, "t": 2}       # A5 日期（Excel 序列号）
pcd["4"]["1"] = {"v": "PC-TEST-001"}                       # B5 凭证号
pcd["4"]["2"] = {"v": "测试付款 China Report"}              # C5 摘要
pcd["4"]["4"] = {"v": 1234567}                             # E5 支出
pcd["4"]["28"] = {"v": "\u2611"}                           # AC5 核销=已勾选
pcd.setdefault("5", {})["28"] = {"v": "\u2611"}            # AC6 已勾选
r = cli.post("/api/finance/cn-report/save", json=pay)
ck("CN save 200", r.status_code == 200, r.get_json())
ck("CN save ok", (r.get_json() or {}).get("ok") is True, r.get_json())

import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(app.CN_REPORT)
ck("保存后仍只有一张表", wb.sheetnames == [app.CN_SHEET_NAME], wb.sheetnames)
ws = wb[app.CN_SHEET_NAME]
ck("B5 写回", ws["B5"].value == "PC-TEST-001", ws["B5"].value)
ck("C5 写回", ws["C5"].value == "测试付款 China Report", ws["C5"].value)
ck("E5 写回", ws["E5"].value == 1234567, ws["E5"].value)
ck("D5 写回", ws["D5"].value is None, ws["D5"].value)
ck("AC5 核销=☑", ws["AC5"].value == "\u2611", ws["AC5"].value)
ck("AC6 核销=☑", ws["AC6"].value == "\u2611", ws["AC6"].value)
ck("AC4 表头还在", ws["AC4"].value == "\u6838\u9500", ws["AC4"].value)
ck("合计公式还在", "=SUM" in str(ws["E205"].value), ws["E205"].value)
ck("三语表头没被毁", (ws["K4"].value or "").strip() != "" and (ws["U4"].value or "").strip() != "",
   (ws["K4"].value, ws["U4"].value))

log("=== 5) 保存后快照缓存已失效并可按需重建")
sm = cli.get("/api/finance/cn-report/snapshot")
ck("重取快照 200（自动重建）", sm.status_code == 200, sm.status_code)
j2 = sm.get_json() or {}
ck("重建后看得到刚写的值", (j2.get("sheets", {}).get(sid, {}).get("cellData", {})
                            .get("4", {}).get("1", {}) or {}).get("v") == "PC-TEST-001",
   (j2.get("sheets", {}).get(sid, {}).get("cellData", {}).get("4", {}).get("1", {}) or {}).get("v"))

log("=== 6) 每日备份")
bdir = os.path.join(app.FIN_DIR, "backups")
bks = [f for f in os.listdir(bdir) if f.startswith("cn_report_")] if os.path.isdir(bdir) else []
ck("生成了中国报表当日备份", len(bks) >= 1, bks)

log("=== 7) 导出接口")
r = cli.get("/api/finance/cn-report")
ck("CN 下载 xlsx 200", r.status_code == 200 and len(r.data) > 1000, (r.status_code, len(r.data)))

log("")
log("失败项: %d" % len(fails))
for f in fails:
    log("  - %s" % f)

io.open(os.path.join(HERE, "_cn_api_test_out.txt"), "w", encoding="utf-8").write("\n".join(out) + "\n")

# 清理：把中国报表恢复成干净种子（测试写入的数据不要留在里面）
try:
    os.remove(app.CN_REPORT)
    app._cn_cache_invalidate()
    app._cn_seed_if_missing()
    log("[cleanup] 已恢复干净种子报表")
except Exception as e:
    log("[cleanup] 失败: %s" % e)
io.open(os.path.join(HERE, "_cn_api_test_out.txt"), "w", encoding="utf-8").write("\n".join(out) + "\n")
print("done, fails=%d" % len(fails))
