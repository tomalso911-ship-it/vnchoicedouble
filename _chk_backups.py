# -*- coding: utf-8 -*-
"""只读对比：当前 vn_report.xlsx 与两份 before_merge 备份的差异（CN 行 / 合计 / 跨表引用）。"""
import io
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
SHEET = "Chi th\u1ef1c t\u1ebf"
MARK_COL = 40
OUT = os.path.join(HERE, "_chk_backups_out.txt")
L = []


def log(s):
    L.append(str(s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(s)


FILES = [
    ("当前 vn_report.xlsx", os.path.join(FIN, "vn_report.xlsx")),
    ("备份 A 20261014_052849", os.path.join(FIN, "backups", "vn_report_20261014_052849_before_merge.xlsx.bak")),
    ("备份 B 20261014_055253", os.path.join(FIN, "backups", "vn_report_20261014_055253_before_merge.xlsx.bak")),
]

for name, path in FILES:
    log("=== %s" % name)
    if not os.path.isfile(path):
        log("   （文件不存在）")
        continue
    log("   大小 = %d bytes" % os.path.getsize(path))
    wb = openpyxl.load_workbook(path, data_only=False)
    log("   工作表 %d 张：%s" % (len(wb.sheetnames), " | ".join(wb.sheetnames)))
    ws = wb[SHEET]
    cn, last = 0, 0
    for r in range(5, 420):
        vals = [ws.cell(r, c).value for c in (1, 2, 3, 4, 5)]
        if any(v is not None and str(v).strip() != "" for v in vals):
            last = r
        mk = ws.cell(r, MARK_COL).value
        if isinstance(mk, str) and mk.startswith("CN:"):
            cn += 1
    log("   Chi thực tế：CN 标记行 = %d；最后有数据行 = %d" % (cn, last))
    log("   行3 合计：E(Chi VI)=%s  O(Chi ZH)=%s  Y(Chi EN)=%s"
        % (ws.cell(3, 5).value, ws.cell(3, 15).value, ws.cell(3, 25).value))
    for r in range(5, 420):
        if str(ws.cell(r, 3).value or "").strip() == "T\u1ed4NG C\u1ed8NG":
            log("   底部 TỔNG CỘNG 在第 %d 行，E=%s" % (r, ws.cell(r, 5).value))
            break
    refs = []
    for sh in wb.worksheets:
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 400), max_col=min(sh.max_column, 40)):
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("=") and SHEET in v:
                    refs.append("%s!%s=%s" % (sh.title, cell.coordinate, v))
    log("   引用「Chi thực tế」的公式 %d 处：%s" % (len(refs), refs[:5]))
    wb.close()
log("done")
