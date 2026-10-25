# -*- coding: utf-8 -*-
"""① 备份 vn_report.xlsx；② 只改 ACB 区段的 F295 笔误 642451 → 643451；
   ③ 【外科手术式】xlsx 内 zip 定点改一处 XML，不整本重写（避免丢公式缓存值/格式）；
   ④ 改前后做全量单元格摘要比对，证明"只差这一格"；⑤ 原子替换。"""
import io, sys, os, re, json, time, shutil, zipfile, hashlib, posixpath
from openpyxl import load_workbook
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

XLSX = r'C:\vnchoicedouble\finance_data\vn_report.xlsx'
SHEET = 'Chi tiết thu-chi'
COORD = 'F295'
OLD, NEW = 642451, 643451


def cell_digest(path, tag):
    wb = load_workbook(path, data_only=False)
    h = hashlib.sha256()
    per = []
    for ws in wb.worksheets:
        sub = hashlib.sha256()
        for (r, c), cell in sorted(ws._cells.items()):
            line = '%d|%d|%r\n' % (r, c, cell.value)
            sub.update(line.encode('utf-8'))
            h.update(('%s|' % ws.title).encode('utf-8') + line.encode('utf-8'))
        per.append('%s:%d 格 %s' % (ws.title, len(ws._cells), sub.hexdigest()[:10]))
    wb.close()
    print('   [%s] 总摘要 %s' % (tag, h.hexdigest()[:20]))
    for p in per:
        print('        ', p)
    return h.hexdigest()


def zip_digest(path, tag, skip=()):
    d = {}
    with zipfile.ZipFile(path) as z:
        for i in z.infolist():
            if i.filename in skip:
                continue
            d[i.filename] = hashlib.sha256(z.read(i.filename)).hexdigest()
    print('   [%s] zip 条目 %d 个' % (tag, len(d)))
    return d


print('=== 0) 改前状态断言 ===')
wb0 = load_workbook(XLSX, data_only=False)
ws0 = wb0[SHEET]
assert ws0['F294'].value == 16000000, ws0['F294'].value
assert ws0['F293'].value == '=D293-E293', ws0['F293'].value
assert ws0['F295'].value == OLD, ws0['F295'].value
assert ws0['F296'].value == 643452, ws0['F296'].value
print('   F293=%r  F294=%r  F295=%r  F296=%r  ✓ 与预期一致'
      % (ws0['F293'].value, ws0['F294'].value, ws0['F295'].value, ws0['F296'].value))
wb0.close()

print('\n=== 1) 备份 ===')
bakdir = r'C:\vnchoicedouble\backups'
os.makedirs(bakdir, exist_ok=True)
bak = os.path.join(bakdir, 'vn_report_%s_before_F295fix.xlsx.bak' % time.strftime('%Y%m%d_%H%M%S'))
shutil.copy2(XLSX, bak)
print('   ->', bak)

before_cells = cell_digest(XLSX, '改前 单元格')
before_zip = zip_digest(XLSX, '改前 zip')
st0 = os.stat(XLSX)
print('   改前 mtime_ns=%d size=%d' % (st0.st_mtime_ns, st0.st_size))

print('\n=== 2) 定位工作表 XML ===')
with zipfile.ZipFile(XLSX) as z:
    wbxml = z.read('xl/workbook.xml').decode('utf-8')
    rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
m = re.search(r'<sheet[^>]*name="%s"[^>]*/?>' % re.escape(SHEET), wbxml)
assert m, 'workbook.xml 里没找到该表'
rid = re.search(r'r:id="([^"]+)"', m.group(0)).group(1)
# 属性顺序无关：先取出含该 Id 的 <Relationship .../>，再从里面取 Target
mrel = None
for cand in re.finditer(r'<Relationship\b[^>]*/>', rels):
    if re.search(r'Id="%s"' % re.escape(rid), cand.group(0)):
        mrel = cand
        break
assert mrel, 'rels 里没找到 Id=%s（rels 片段：%s）' % (rid, rels[:400])
target = re.search(r'Target="([^"]+)"', mrel.group(0)).group(1)
if target.startswith('/'):
    part = target.lstrip('/')
else:
    part = posixpath.normpath(posixpath.join('xl', target))
print('   %s -> r:id=%s -> %s' % (SHEET, rid, part))
print('     关系元素: %s' % mrel.group(0))

print('\n=== 3) 定点改 XML ===')
with zipfile.ZipFile(XLSX) as z:
    infos = z.infolist()
    blobs = {i.filename: z.read(i.filename) for i in infos}
xml = blobs[part].decode('utf-8')
pat = re.compile(r'<c r="%s"(?P<attrs>[^>/]*)>(?P<inner>.*?)</c>' % COORD, re.S)
mm = pat.search(xml)
assert mm, '单元格 %s 在 XML 里没找到' % COORD
attrs, inner = mm.group('attrs'), mm.group('inner')
print('   原 XML: <c r="%s"%s>%s</c>' % (COORD, attrs, inner))
assert 'f>' not in inner, '这格是公式，不该走这条路：%r' % inner
assert re.fullmatch(r'<v>%d</v>' % OLD, inner.strip()), '值不是 %d：%r' % (OLD, inner)
new_inner = '<v>%d</v>' % NEW
xml2 = xml[:mm.start()] + '<c r="%s"%s>%s</c>' % (COORD, attrs, new_inner) + xml[mm.end():]
assert xml2.count(new_inner) == 1
blobs[part] = xml2.encode('utf-8')
print('   新 XML: <c r="%s"%s>%s</c>' % (COORD, attrs, new_inner))

print('\n=== 4) 原子写回 ===')
tmp = XLSX + '.tmp'
with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zo:
    for i in infos:
        zi = zipfile.ZipInfo(i.filename, date_time=i.date_time)
        zi.compress_type = i.compress_type
        zi.external_attr = i.external_attr
        zi.internal_attr = i.internal_attr
        zi.create_system = i.create_system
        zo.writestr(zi, blobs[i.filename])
os.replace(tmp, XLSX)
st1 = os.stat(XLSX)
print('   已写回  mtime_ns=%d size=%d' % (st1.st_mtime_ns, st1.st_size))

print('\n=== 5) 改后校验 ===')
wb1 = load_workbook(XLSX, data_only=False)
ws1 = wb1[SHEET]
assert wb1.sheetnames == load_workbook(bak).sheetnames
print('   表名一致 ✓  %s' % wb1.sheetnames)
print('   F293=%r F294=%r F295=%r F296=%r F297=%r'
      % (ws1['F293'].value, ws1['F294'].value, ws1['F295'].value, ws1['F296'].value, ws1['F297'].value))
assert ws1['F295'].value == NEW
wb1.close()
wbv = load_workbook(XLSX, data_only=True)
print('   （data_only 读）F295=%r  F303=%r  F304=%r   ← 公式缓存值必须还在'
      % (wbv[SHEET]['F295'].value, wbv[SHEET]['F303'].value, wbv[SHEET]['F304'].value))
wbv.close()

after_cells = cell_digest(XLSX, '改后 单元格')
after_zip = zip_digest(XLSX, '改后 zip')
changed = [k for k in before_zip if before_zip[k] != after_zip.get(k)]
print('\n   zip 内容发生变化的条目: %s' % changed)
print('   单元格摘要 改前=%s 改后=%s  %s' % (
    before_cells, after_cells,
    '（值变了一格，摘要本来就会变，下面用逐格差异确认只差 F295）'))

print('\n=== 6) 逐格差异（应只有 F295）===')
wa = load_workbook(bak, data_only=False)
wb2 = load_workbook(XLSX, data_only=False)
diffs = []
for title in wa.sheetnames:
    A, B = wa[title], wb2[title]
    keys = set(A._cells) | set(B._cells)
    for k in sorted(keys):
        va = A._cells.get(k).value if k in A._cells else None
        vb = B._cells.get(k).value if k in B._cells else None
        if va != vb:
            diffs.append((title, k[0], k[1], va, vb))
print('   差异条数 = %d' % len(diffs))
for d in diffs[:10]:
    print('      %s 行%s 列%s : %r -> %r' % d)
assert len(diffs) == 1 and diffs[0][0] == SHEET and diffs[0][3] == OLD and diffs[0][4] == NEW
print('\n✅ 全库只改了 %s!%s : %d -> %d' % (SHEET, COORD, OLD, NEW))
