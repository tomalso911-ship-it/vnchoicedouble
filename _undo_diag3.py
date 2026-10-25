# -*- coding: utf-8 -*-
"""看清"红+粗"到底出现在哪些行、是不是日期行 —— 决定清理规则会不会误伤原始格式。"""
import os
import sys
import shutil
import datetime
import re
import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
FD = os.path.join(BASE, "finance_data")
DATE_RE = re.compile(r"^\s*(\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4})")
DATEPAT_RE = re.compile(r"[ymd]", re.I)


def loader(path):
    if not path.lower().endswith(".xlsx"):
        tmp = os.path.join(BASE, "_tmp_probe3.xlsx")
        shutil.copyfile(path, tmp)
        path = tmp
    return openpyxl.load_workbook(path)


def is_date_row(ws, r):
    v = ws.cell(r, 1).value
    if isinstance(v, (datetime.datetime, datetime.date)):
        return True
    if isinstance(v, (int, float)):
        pat = ws.cell(r, 1).number_format or ""
        return bool(DATEPAT_RE.search(pat))
    if isinstance(v, str) and DATE_RE.match(v):
        return True
    return False


def scan(path, label):
    print("=" * 84)
    print(label, "  ", os.path.basename(path))
    wb = loader(path)
    for ws in wb.worksheets:
        items = []
        for r in range(1, min(ws.max_row, 900) + 1):
            hits = []
            for c in range(1, 11):
                f = ws.cell(r, c).font
                rgb = ""
                try:
                    rgb = str(getattr(f.color, "rgb", "") or "").upper()
                except Exception:
                    rgb = ""
                if "FF0000" in rgb and f.bold:
                    hits.append(c)
            if hits:
                dr = is_date_row(ws, r)
                a = ws.cell(r, 1).value
                vals = [str(ws.cell(r, c).value)[:14] if ws.cell(r, c).value not in (None, "") else "" for c in range(1, 9)]
                items.append((r, hits, dr, a, vals))
        if not items:
            continue
        print("  --- %s" % ws.title)
        for r, hits, dr, a, vals in items:
            print("    r%-4d 红+粗列=%s 日期行=%-5s A=%r | %s" % (
                r, hits, dr, str(a)[:26], " | ".join(vals)))


scan(os.path.join(FD, "vn_report.xlsx.bak_before_postfix"), "【旧·昨天可用】")
scan(os.path.join(FD, "vn_report.xlsx"), "【新·今天上传】")
