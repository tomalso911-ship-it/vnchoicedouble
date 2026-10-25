import sys, re, json
sys.path.insert(0, r'c:\vnchoicedouble')
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass
import openpyxl
wb = openpyxl.load_workbook(r'c:\vnchoicedouble\finance_data\vn_report.xlsx', data_only=False)
ws = wb['Chi tiết thu-chi']
print('sheet:', ws.title, 'dims', ws.dimensions)
from collections import Counter
c = Counter()
samples = {}
for r in range(1, ws.max_row + 1):
    cell = ws.cell(row=r, column=1)
    if cell.value is None:
        continue
    c[cell.number_format] += 1
    if cell.number_format not in samples:
        samples[cell.number_format] = (cell.coordinate, repr(cell.value))
for pat, n in c.most_common(12):
    print('  pattern=%-16r count=%-5d sample=%s' % (pat, n, samples[pat]))
