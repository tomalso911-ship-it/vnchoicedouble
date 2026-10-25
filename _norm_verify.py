# -*- coding: utf-8 -*-
"""校验：日期格式是否统一 dd-mm-yyyy、泄漏的红+粗是否被清、原始小计行的红+粗是否保住。"""
import os
import sys
import shutil
import datetime
import re
import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "finance_data", "vn_report.xlsx")
if not path.lower().endswith(".xlsx"):
    tmp = os.path.join(BASE, "_tmp_v.xlsx")
    shutil.copyfile(path, tmp)
    path = tmp

DATE_RE = re.compile(r"^\s*\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}")
wb = openpyxl.load_workbook(path)
print("文件: %s  (%d bytes)" % (os.path.basename(sys.argv[1] if len(sys.argv) > 1 else path), os.path.getsize(path)))
for ws in wb.worksheets:
    fmts = {}
    bad = 0
    for r in range(1, min(ws.max_row, 900) + 1):
        for c in range(1, 29):
            cc = ws.cell(r, c)
            v = cc.value
            if isinstance(v, (datetime.datetime, datetime.date)):
                fmts[cc.number_format] = fmts.get(cc.number_format, 0) + 1
                if cc.number_format != "dd-mm-yyyy":
                    bad += 1
            elif isinstance(v, (int, float)) and not isinstance(v, bool) and re.search(r"yy|dd|mmm", str(cc.number_format or ""), re.I):
                bad += 1
    reds = []
    for r in range(1, min(ws.max_row, 900) + 1):
        hit = []
        for c in range(1, 11):
            f = ws.cell(r, c).font
            try:
                rgb = str(getattr(f.color, "rgb", "") or "").upper()
            except Exception:
                rgb = ""
            if "FF0000" in rgb and f.bold:
                hit.append(c)
        if hit:
            a = ws.cell(r, 1).value
            dr = isinstance(a, (datetime.datetime, datetime.date)) or (isinstance(a, str) and DATE_RE.match(a))
            reds.append("r%d%s" % (r, "(日期行!)" if dr else ""))
    print("  %-22s 日期格格式=%s 非标准=%d  红+粗行=%s" % (ws.title, fmts or "-", bad, ", ".join(reds) or "无"))
