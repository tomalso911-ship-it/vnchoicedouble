# -*- coding: utf-8 -*-
"""临时：逐表统计 ORIG / ARCHIVE / CURRENT 的 边框 / 填充 / 数字格式 / 粗体 单元格数，
定位到底哪些表的哪些样式被“保存回写”弄丢了。"""
import sys
import openpyxl

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

FILES = [
    ('ORIG   ', r'c:/vnchoicedouble/FinanceReport/Choice Slats VN - Financial Report.xlsx'),
    ('ARCHIVE', r'c:/vnchoicedouble/finance_data/archive/vn_2026-05.xlsx'),
    ('CURRENT', r'c:/vnchoicedouble/finance_data/vn_report.xlsx'),
]

data = {}
order = None
for tag, p in FILES:
    wb = openpyxl.load_workbook(p, data_only=False)
    if order is None:
        order = wb.sheetnames
    stats = {}
    for ws in wb.worksheets:
        bd = fill = nf = bold = cells = 0
        for row in ws.iter_rows():
            for cl in row:
                if cl.value is None:
                    continue
                cells += 1
                b = cl.border
                if any(getattr(b, s).style for s in ('left', 'right', 'top', 'bottom')):
                    bd += 1
                if cl.fill is not None and getattr(cl.fill, 'patternType', None):
                    fill += 1
                if cl.number_format and cl.number_format != 'General':
                    nf += 1
                if cl.font is not None and cl.font.b:
                    bold += 1
        stats[ws.title] = (cells, bd, fill, nf, bold)
    data[tag] = stats
    print('=== %s  (%s)' % (tag.strip(), p))

names = order
print()
print('%-22s %-26s %-26s %-26s' % ('sheet', 'ORIG  cells/bd/fill/nf/bold', 'ARCHIVE', 'CURRENT'))
for nm in names:
    row = ['%-22s' % nm]
    for tag, _ in FILES:
        s = data[tag.strip()].get(nm)
        row.append('%-26s' % ('/' .join(str(x) for x in s) if s else '-'))
    print(' '.join(row))
