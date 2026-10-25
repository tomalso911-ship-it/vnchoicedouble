# -*- coding: utf-8 -*-
import sys
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass
from openpyxl import load_workbook
wb = load_workbook(r"finance_data/cn_report.xlsx")
ws = wb.worksheets[0]
print("=== 数据区字色 / 斜体（按行） ===")
for r in range(5, 15):
    ac = ws.cell(r, 29).value
    cols = []
    for c in (1, 2, 3, 5, 6, 11, 13, 15, 21, 23, 25):
        f = ws.cell(r, c).font
        rgb = (f.color.rgb if f.color else None)
        cols.append("%s:%s%s" % (ws.cell(r, c).column_letter, str(rgb)[-6:], "I" if f.italic else ""))
    print("r%-3d AC=%s | %s" % (r, ac, " ".join(cols)))
print()
print("AC 列自身字色 r5 =", (ws.cell(5, 29).font.color.rgb if ws.cell(5, 29).font.color else None))
print()
print("=== 合计公式 ===")
for r in (3, 205):
    for c in ("C", "E", "M", "O", "W", "Y"):
        v = ws["%s%d" % (c, r)].value
        if v is not None:
            print("  %s%d = %r  (fmt=%r)" % (c, r, v, ws["%s%d" % (c, r)].number_format))
print()
print("=== 金额是否原样（越南盾） ===")
t = 0
for r in range(5, 205):
    for c in (5, 15, 25):
        v = ws.cell(r, c).value
        if isinstance(v, (int, float)):
            t += v
print("  三块支出合计 = %s（每块 %s）" % (format(t, ","), format(t // 3, ",")))
checked = sum(1 for r in range(5, 205) if ws.cell(r, 29).value == "☑")
print("  已勾选行数 =", checked)
wb.close()
