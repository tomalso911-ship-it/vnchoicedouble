# -*- coding: utf-8 -*-
"""dump 「TỔNG HỢP」表：每行标签 + C 列金额（越南 / 合并两份），
   以及资金流水账「Chi tiết thu-chi」最后一笔交易日期（做"截至"用）。"""
import io, sys, datetime
from openpyxl import load_workbook
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

for path in (r'finance_data/vn_report.xlsx', r'finance_data/merge_report.xlsx'):
    print('\n########## %s ##########' % path)
    wb = load_workbook(path, data_only=False)
    print('  表:', wb.sheetnames)
    ws = wb['TỔNG HỢP']
    print('  TỔNG HỢP max_row=%d max_col=%d' % (ws.max_row, ws.max_column))
    for r in range(1, min(ws.max_row, 50) + 1):
        vals = []
        for c in range(1, 6):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            if isinstance(v, str):
                v = v.strip()
                if len(v) > 58:
                    v = v[:58] + '…'
            vals.append('%s=%r' % (chr(64 + c), v))
        if vals:
            print('   r%-3d %s' % (r, ' | '.join(vals)))
    # 资金流水账最后一笔日期
    if 'Chi tiết thu-chi' in wb.sheetnames:
        w = wb['Chi tiết thu-chi']
        last = None
        for r in range(1, w.max_row + 1):
            v = w.cell(row=r, column=1).value
            if isinstance(v, (datetime.datetime, datetime.date)):
                last = v
        print('  Chi tiết thu-chi 最后一笔日期 =', last)
    wb.close()
