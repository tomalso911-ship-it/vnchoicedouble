# -*- coding: utf-8 -*-
"""① 逐格证明 vn_report.xlsx 改了 F295 而已（并对齐备份）；
   ② 校验公式缓存值仍在；
   ③ 作废旧快照缓存并按新键重建（快照→指纹→键）；
   ④ 清掉用户要求删除的两份 .bak_before_* 旧污染快照。"""
import io, sys, os, re, glob, json, shutil, zipfile, hashlib, time
from openpyxl import load_workbook
sys.path.insert(0, r'C:\vnchoicedouble')
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

XLSX = r'C:\vnchoicedouble\finance_data\vn_report.xlsx'
SHEET, COORD = 'Chi tiết thu-chi', 'F295'
BAKLIST = sorted(glob.glob(r'C:\vnchoicedouble\backups\vn_report_*_before_F295fix.xlsx.bak'))
print('=== 1) 逐格比对（新文件 vs 备份）===')
print('   备份:', [os.path.basename(b) for b in BAKLIST])
# 备份去掉 .bak 后缀才能被 openpyxl 打开（它按扩展名判类型）
tmpbak = r'C:\vnchoicedouble\backups\_cmp_before.xlsx'
shutil.copy2(BAKLIST[0], tmpbak)
A = load_workbook(tmpbak, data_only=False)
B = load_workbook(XLSX, data_only=False)
assert A.sheetnames == B.sheetnames, (A.sheetnames, B.sheetnames)
print('   表名一致 ✓', B.sheetnames)
diffs = []
for t in A.sheetnames:
    for k in sorted(set(A[t]._cells) | set(B[t]._cells)):
        va = A[t]._cells[k].value if k in A[t]._cells else None
        vb = B[t]._cells[k].value if k in B[t]._cells else None
        if va != vb:
            diffs.append((t, k[0], k[1], va, vb))
print('   差异条数 = %d' % len(diffs))
for d in diffs:
    print('      %s 行%s 列%s : %r -> %r' % d)
assert len(diffs) == 1 and diffs[0][0] == SHEET and diffs[0][3] == 642451 and diffs[0][4] == 643451
print('   ✅ 全库只改了 %s!%s : 642451 -> 643451' % (SHEET, COORD))
A.close(); B.close(); os.remove(tmpbak)

print('\n=== 2) 公式缓存值仍在？（外科式改法不应动别的字节）===')
v = load_workbook(XLSX, data_only=True)[SHEET]
print('   F295=%r（新值）  F303=%r  F304=%r  F294=%r' % (
    v['F295'].value, v['F303'].value, v['F304'].value, v['F294'].value))
assert v['F295'].value == 643451 and v['F303'].value == 117 and v['F304'].value == 30000117

print('\n=== 3) zip 条目级差异（应只 1 个）===')
def zd(p):
    with zipfile.ZipFile(p) as z:
        return {i.filename: hashlib.sha256(z.read(i.filename)).hexdigest() for i in z.infolist()}
dst = r'C:\vnchoicedouble\backups\_cmp_before.xlsx'
shutil.copy2(BAKLIST[0], dst)
za, zb = zd(dst), zd(XLSX)
os.remove(dst)
ch = [k for k in za if za[k] != zb.get(k)]
print('   变化条目:', ch)
assert ch == ['xl/worksheets/sheet2.xml'], ch

print('\n=== 4) 作废旧快照缓存（键已随 mtime 失效）===')
FIN = r'C:\vnchoicedouble\finance_data'
removed = []
for p in [os.path.join(FIN, 'vn_report.snapshot.key'), os.path.join(FIN, 'vn_report.snapshot.fp')] + \
         glob.glob(os.path.join(FIN, 'vn_report.snapshot.*.json')):
    try:
        os.remove(p); removed.append(os.path.basename(p))
    except OSError:
        pass
print('   已删:', removed)

print('\n=== 5) 按新键重建快照 + 写指纹 ===')
import univer_convert
st = os.stat(XLSX)
key = "%d:%d:%s" % (st.st_mtime_ns, st.st_size, XLSX)
path = os.path.join(FIN, 'vn_report.snapshot.%s.json' % hashlib.sha1(key.encode('utf-8')).hexdigest()[:16])
snap = univer_convert.xlsx_to_univer(XLSX)
with open(path + '.part', 'w', encoding='utf-8') as fh:
    json.dump(snap, fh, ensure_ascii=False, separators=(",", ":"))
os.replace(path + '.part', path)
h = hashlib.sha256()
with open(path, 'rb') as fh:
    for c in iter(lambda: fh.read(1 << 20), b''):
        h.update(c)
with open(os.path.join(FIN, 'vn_report.snapshot.fp') + '.tmp', 'w', encoding='utf-8') as fh:
    fh.write(h.hexdigest())
os.replace(os.path.join(FIN, 'vn_report.snapshot.fp') + '.tmp', os.path.join(FIN, 'vn_report.snapshot.fp'))
with open(os.path.join(FIN, 'vn_report.snapshot.key') + '.tmp', 'w', encoding='utf-8') as fh:
    fh.write(key)
os.replace(os.path.join(FIN, 'vn_report.snapshot.key') + '.tmp', os.path.join(FIN, 'vn_report.snapshot.key'))
print('   新快照:', os.path.basename(path))
print('   指纹已写 =', os.path.exists(os.path.join(FIN, 'vn_report.snapshot.fp')))

print('\n=== 6) 清理旧污染快照 .bak_before_* ===')
old = sorted(glob.glob(os.path.join(FIN, 'vn_report.snapshot.*.json.bak_before_*')))
print('   找到:', [os.path.basename(p) for p in old])
for p in old:
    os.remove(p)
print('   已删除 %d 个' % len(old))
left = [os.path.basename(p) for p in glob.glob(os.path.join(FIN, '*snapshot*'))]
print('   finance_data 里剩下的快照相关文件:', left)
