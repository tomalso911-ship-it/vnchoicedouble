# -*- coding: utf-8 -*-
"""验证：中国报表数据区字体已是蓝色（文件 + 服务端快照），核销列仍居中。"""
import json
import os
import sys
import urllib.request

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
P = os.path.join(HERE, "finance_data", "cn_report.xlsx")


def fx(cell):
    c = getattr(cell.font, "color", None)
    rgb = getattr(c, "rgb", None) if c else None
    return {"v": cell.value,
            "rgb": (str(rgb) if rgb is not None else None),
            "h": (cell.alignment.horizontal if cell.alignment else None),
            "bold": bool(cell.font.bold)}


wb = openpyxl.load_workbook(P)
ws = wb[wb.sheetnames[0]]
print("XLSX sheets =", json.dumps(wb.sheetnames, ensure_ascii=True))
for a in ("A5", "B5", "C5", "E5", "F5", "J5", "K5", "AB5", "AC4", "AC5", "A4", "E205"):
    print("  %-5s = %s" % (a, json.dumps(fx(ws[a]), ensure_ascii=True)))

# 数字格式是否被保留（蓝色不能把格式冲掉）
print("E5 number_format =", json.dumps(ws["E5"].number_format), "| A5 =", json.dumps(ws["A5"].number_format))

try:
    with urllib.request.urlopen("http://127.0.0.1:5050/api/finance/cn-report/snapshot", timeout=240) as r:
        snap = json.loads(r.read().decode("utf-8"))
    sid = snap["sheetOrder"][0]
    cd = snap["sheets"][sid]["cellData"]
    out = {}
    for rr, cc in (("4", 0), ("4", 4), ("4", 5), ("4", 9), ("4", 28), ("3", 28)):
        cell = (cd.get(rr, {}) or {}).get(str(cc))
        out["r%s c%s" % (rr, cc)] = {"v": (cell or {}).get("v"), "s": (cell or {}).get("s")}
    print("SNAPSHOT =", json.dumps(out, ensure_ascii=True))
except Exception as e:
    print("SNAPSHOT GET ERR:", str(e)[:200])
