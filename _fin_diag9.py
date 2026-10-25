# -*- coding: utf-8 -*-
"""直接读源 xlsx：Chi tiết thu-chi 的 F 列（第6列）在若干行的真实公式，判断 xlsx 是否已被写坏。"""
import io, sys
from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
wb = load_workbook('finance_data/vn_report.xlsx', data_only=False)
ws = wb['Chi tiết thu-chi']
print('xlsx: Chi tiết thu-chi  max_row=%s' % ws.max_row)
print('row(1基) | A | D | E | F')
for r in [4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 79, 80, 81, 294, 295, 300, 301, 302, 303, 568]:
    vals = []
    for c in (1, 4, 5, 6):
        cell = ws.cell(row=r, column=c)
        vals.append(repr(cell.value))
    print('%6d | %s' % (r, ' | '.join(vals)))
