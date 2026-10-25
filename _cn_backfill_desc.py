# -*- coding: utf-8 -*-
"""回填中国报表 cn_report.xlsx「实际支出」表被写空/写脏的摘要（C/M/W 三语）。

背景：
  2026-10-14 版（cn_report_20261014_095127_before_datefmt.xlsx.bak）是最后一次完好状态；
  其后的回归把第 8 / 10 / 14 行（凭证 UNC00055 / UNC00057 / UNC00102）的
  C(摘要-越) / M(摘要-中) / W(摘要-英) 写成了数字 0，随后又被 2026-10-15 的
  junk-clean 清成了空。除这 9 格外，当前文件与那份完好版【逐格完全一致】。

做法：
  1) 先备份当前文件；
  2) 只写回这 9 格文字（顺带把 3 个 C 格里残留的金额格式 '"¥"#,##0' 恢复成 General）；
  3) 原子落盘（.tmp + os.replace，避免读写并发拿到坏 zip）；
  4) 失效中国报表快照缓存，让前端下次必定从新文件重建。

用法: python _cn_backfill_desc.py [--dry-run]
"""
import glob
import os
import shutil
import sys
import time

import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
CN = os.path.join(BASE, "finance_data", "cn_report.xlsx")
GOOD = os.path.join(BASE, "finance_data", "backups",
                    "cn_report_20261014_095127_before_datefmt.xlsx.bak")

# 需要回填的行 -> 三语摘要来源列（C=3 / M=13 / W=23）
ROWS = [8, 10, 14]
COLS = [3, 13, 23]

# 兜底：若好版本文件缺失，用下面写死的值（已与 FinanceReport 原始报表逐笔核对）
FALLBACK = {
    (8, 3): "Trả tiền thuê cần cẩu của Công ty Beng Fook",
    (8, 13): "向明福公司支付起重机租金",
    (8, 23): "Payment for the crane rental from Beng Fook Company",
    (10, 3): "Công ty TNHH Thương mại Điện máy Zhongli",
    (10, 13): "中力电机贸易有限公司",
    (10, 23): "Zhongli Electrical Machinery Trading Co., Ltd.",
    (14, 3): "Thanh toán tiền vật tư Công ty Huy Phượng",
    (14, 13): "向 Huy Phượng 公司支付材料款",
    (14, 23): "Payment for materials to Huy Phuong Company",
}


def load_values():
    if os.path.isfile(GOOD):
        tmp = GOOD + ".dxl.tmp.xlsx"
        shutil.copyfile(GOOD, tmp)
        try:
            ws = openpyxl.load_workbook(tmp, data_only=False).worksheets[0]
            out = {}
            for r in ROWS:
                for c in COLS:
                    v = ws.cell(r, c).value
                    if v is None or str(v).strip() == "":
                        return None
                    out[(r, c)] = v
            return out
        finally:
            try:
                os.remove(tmp)
            except OSError:
                pass
    return None


def main():
    dry = "--dry-run" in sys.argv
    vals = load_values() or FALLBACK
    print("来源:", "备份 " + os.path.basename(GOOD) if os.path.isfile(GOOD) else "内置兜底值")

    wb = openpyxl.load_workbook(CN)
    ws = wb.worksheets[0]
    changed = []
    for (r, c), newv in sorted(vals.items()):
        oldv = ws.cell(r, c).value
        if oldv == newv:
            continue
        changed.append((r, c, oldv, newv))
    print("待写入格数:", len(changed))
    for r, c, oldv, newv in changed:
        cl = ws.cell(r, c).column_letter
        print("  r%d%s: %r -> %r" % (r, cl, oldv, newv))

    if dry or not changed:
        print("(dry-run，未修改文件)")
        return

    ts = time.strftime("%Y%m%d_%H%M%S")
    bak = CN + ".bak_before_descbackfill." + ts
    shutil.copyfile(CN, bak)
    print("已备份 ->", os.path.basename(bak))

    for (r, c), newv in vals.items():
        cell = ws.cell(r, c)
        cell.value = newv
        # 摘要列是文字列：清掉残留的金额格式，避免显示成 ¥ 数字
        if cell.number_format not in ("General", "@"):
            cell.number_format = "General"

    tmp = CN + ".backfill.tmp"
    wb.save(tmp)
    os.replace(tmp, CN)
    print("已写回 ->", CN)

    # 失效中国报表快照缓存（键 + 旧快照），前端下次必定重建
    n = 0
    for p in glob.glob(os.path.join(BASE, "finance_data", "cn_report.snapshot.*")) + \
             [os.path.join(BASE, "finance_data", "cn_report.snapshot.key")]:
        try:
            os.remove(p)
            n += 1
        except OSError:
            pass
    print("已失效中国报表快照缓存文件数:", n)

    # 校验
    ws2 = openpyxl.load_workbook(CN, data_only=False).worksheets[0]
    ok = all(ws2.cell(r, c).value == v for (r, c), v in vals.items())
    print("校验回填结果:", "OK" if ok else "FAIL")


if __name__ == "__main__":
    main()
