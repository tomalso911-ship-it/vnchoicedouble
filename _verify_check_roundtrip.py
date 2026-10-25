# -*- coding: utf-8 -*-
"""验证用户最初报的那件事：**在屏幕上勾「核销」→ 自动保存 → 文件里真的变了，且金额不被破坏**。
做法（安全）：对第 6 行（Excel AC6）勾上 → 等自动保存 → 读文件确认 ☑ → 再取消 → 等自动保存 →
读文件确认回到原值；并复核三块金额仍是 245,025,360、E5 仍是 41,400,000。
只动一格复选框，最终状态与原状一致。"""
import datetime
import io
import json
import os
import shutil
import sys
import time

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
CN = os.path.join(FIN, "cn_report.xlsx")
OUT = os.path.join(HERE, "_verify_check_roundtrip_out.txt")
L = []


def log(s):
    L.append(s)
    io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(s)


def read_state():
    wb = openpyxl.load_workbook(CN, data_only=False)
    ws = wb[wb.sheetnames[0]]
    tot = {}
    for tag, cols in (("base", (4, 5)), ("zh", (14, 15)), ("en", (24, 25))):
        t = 0.0
        for r in range(5, 205):
            for c in cols:
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    t += abs(float(v))
        tot[tag] = round(t, 2)
    st = {"AC5": ws["AC5"].value, "AC6": ws["AC6"].value, "AC7": ws["AC7"].value,
          "E5": ws["E5"].value, "E8": ws["E8"].value,
          "base": tot["base"], "zh": tot["zh"], "en": tot["en"]}
    wb.close()
    return st


safety = os.path.join(FIN, "backups",
                      "cn_report_%s_before_check_rt.xlsx.bak"
                      % datetime.datetime.now().strftime("%Y%m%d_%H%M%S"))
shutil.copyfile(CN, safety)
before = read_state()
log("① 初始（文件）：%s" % json.dumps(before, ensure_ascii=False, default=str))
log("   备份 -> %s" % os.path.basename(safety))

with sync_playwright() as p:
    b = p.chromium.launch(channel="chrome", headless=True, args=["--no-sandbox"])
    pg = b.new_page(viewport={"width": 1500, "height": 900})
    errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto("http://127.0.0.1:5050/lite", wait_until="domcontentloaded")
    pg.wait_for_timeout(1500)
    pg.fill("#loginUser", "tom")
    pg.fill("#loginPwd", "66668888")
    pg.click("#loginBtn")
    pg.wait_for_selector("#topbar:not(.hidden)", timeout=20000)
    pg.evaluate("() => { try { switchView('finance'); } catch(e){} }")
    pg.wait_for_timeout(1000)
    pg.evaluate("() => { var t = document.querySelector('#fin-tabs .fin-tab[data-key=\\'cn\\']'); if (t) t.click(); }")
    for _ in range(80):
        pg.wait_for_timeout(500)
        if pg.evaluate("() => (typeof finReady !== 'undefined') && !!finReady"):
            break
    pg.wait_for_timeout(2500)

    def toggle_and_read(tag):
        r = pg.evaluate("""async () => {
          function g(rr, cc){ return String(finUniverAPI.getActiveWorkbook().getActiveSheet().getRange(rr, cc).getValue()); }
          var before = g(5, 28);
          finCnCheckApply({ startRow: 5, endRow: 5, startColumn: 28, endColumn: 28 });
          await new Promise(function(res){ setTimeout(res, 6000); });   // 等自动保存（挂起窗口最多 ~3s + 防抖 0.7s）
          var st = document.getElementById('fin_status');
          return { before: before, after: g(5, 28), status: st ? st.textContent : '?' };
        }""")
        log("②%s 屏幕上 AC6：%s → %s（状态栏=%s）" % (tag, r["before"], r["after"], r["status"]))
        time.sleep(1.5)
        f = read_state()
        log("   文件里 AC6=%r ；三块合计 base=%s zh=%s en=%s ；E5=%r"
            % (f["AC6"], f["base"], f["zh"], f["en"], f["E5"]))
        return r, f

    r1, f1 = toggle_and_read(" 勾上")
    ok1 = (f1["AC6"] == "\u2611")
    log("   → 勾上后【文件里真的是 ☑】？ %s" % ("✔ 是" if ok1 else "✘ 否"))
    r2, f2 = toggle_and_read(" 取消")
    ok2 = (f2["AC6"] == "\u2610")
    log("   → 取消后【文件里回到 ☐】？ %s" % ("✔ 是" if ok2 else "✘ 否"))
    ok3 = (f2["base"] == f2["zh"] == f2["en"] == 245025360 and f2["E5"] == 41400000)
    log("   → 金额未被破坏（三块 245,025,360 / E5 41,400,000）？ %s" % ("✔ 是" if ok3 else "✘ 否"))
    log("③ 页面报错 = %s" % (errs[:5] if errs else "无"))
    b.close()
