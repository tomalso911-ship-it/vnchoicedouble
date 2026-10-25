# -*- coding: utf-8 -*-
"""逐个工作区核对：文件里的【静态余额】与【Σ(D-E)】是否自洽；
   找出 ledger 用公式重算后与文件静态值不一致的地方。"""
import io, sys
from openpyxl import load_workbook
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
wbv = load_workbook('finance_data/vn_report.xlsx', data_only=True)
ws = wbv['Chi tiết thu-chi']


def num(r, c):
    v = ws.cell(row=r, column=c).value
    return v if isinstance(v, (int, float)) else 0


print('== ACB 区段明细（1基 293..300，D/E/F）==')
for r in range(293, 301):
    print('  r%-4d A=%-14r D=%-14s E=%-14s F=%-14s' % (
        r, str(ws.cell(row=r, column=1).value)[:14], ws.cell(row=r, column=4).value,
        ws.cell(row=r, column=5).value, ws.cell(row=r, column=6).value))

REGIONS = [('VPbank', 6, 77), ('MBbank', 80, 292), ('ACB', 294, 299), ('Tiền mặt 1', 302, 568)]
print()
print('== 各区段 Σ(D-E) 与文件末行余额对照 ==')
for (name, lo, hi) in REGIONS:
    tot = 0.0
    last_f = None
    for r in range(lo, hi + 1):
        tot += num(r, 4) - num(r, 5)
        last_f = ws.cell(row=r, column=6).value
    hdr = ws.cell(row=lo - 1, column=6).value
    print('  %-10s 行%-4d..%-4d  Σ(D-E)=%-16.0f 末行F=%-16r 区头F=%-18r 差=%s' % (
        name, lo, hi, tot, last_f, hdr,
        ('' if not isinstance(last_f, (int, float)) else '%.0f' % (tot - last_f))))
