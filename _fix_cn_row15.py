# -*- coding: utf-8 -*-
"""清除 cn_report.xlsx「实际支出」的重复行 15（2026-07-19 | UNC00101 | 9,810,200）。
证据：凭证 UNC00101 + 金额 9,810,200 在 原始报表 / vn_report / archive 5 份快照里都只有 1 笔（2026-06-19），
      2026-07-19 在源表是另外两笔（UNC00136 35,640,000 内、无凭证 4,750,000 电机维修）。
      行13 与行15 凭证号、金额完全相同 = 账务上不可能，属重复行，令中方报表每块多算 9,810,200。
安全闸：先核对证据，任一不符就【不落盘】；落盘前备份 + 原子替换。
"""
import datetime
import os
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
CN = os.path.join(FIN, "cn_report.xlsx")
VN = os.path.join(FIN, "vn_report.xlsx")
BKD = os.path.join(FIN, "backups")
OUT = os.path.join(HERE, "_fix_cn_row15_out.txt")
DUPPY = 15
KEEP = 13
BLOCKS = ((1, 2, 3, 4, 5), (11, 12, 13, 14, 15), (21, 22, 23, 24, 25))
AMTS = (5, 15, 25)
L = []


def log(s=""):
    L.append(str(s))
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    print(s)


def fmt(v):
    if v is None:
        return "-"
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


def dstr(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        return v.strip()
    return ""


def total(ws):
    return sum(v for r in range(5, 205) for c in AMTS
               if isinstance((v := ws.cell(r, c).value), (int, float)))


wbc = openpyxl.load_workbook(CN, data_only=False)
wsc = wbc[wbc.sheetnames[0]]
wbv = openpyxl.load_workbook(VN, data_only=False)
wsv = wbv["Chi th\u1ef1c t\u1ebf"]

# ---------- 安全闸 ----------
log("=== 落盘前证据核对 ===")
ok = True
r15 = [wsc.cell(DUPPY, c).value for c in (1, 2, 5)]
r13 = [wsc.cell(KEEP, c).value for c in (1, 2, 5)]
log("cn 行13: %s | %s | %s" % tuple(fmt(x) for x in r13))
log("cn 行15: %s | %s | %s" % tuple(fmt(x) for x in r15))

if not (dstr(r13[0]) == "2026-06-19" and r13[1] == "UNC00101" and r13[2] == 9810200):
    log("!! 行13 不再是 2026-06-19 UNC00101 9,810,200 -> 中止"); ok = False
if not (dstr(r15[0]) == "2026-07-19" and r15[1] == "UNC00101" and r15[2] == 9810200):
    log("!! 行15 不再是 2026-07-19 UNC00101 9,810,200 -> 中止"); ok = False

vn_amt = [(r, dstr(wsv.cell(r, 1).value), wsv.cell(r, 2).value)
          for r in range(5, 420) if wsv.cell(r, 5).value == 9810200]
vn_dup = [(r, dstr(wsv.cell(r, 1).value), wsv.cell(r, 2).value)
          for r in range(5, 420) if str(wsv.cell(r, 2).value or "").strip().upper() == "UNC00101"]
log("vn 侧 金额=9,810,200 的行 = %s" % vn_amt)
log("vn 侧 凭证=UNC00101 的行 = %s" % vn_dup)
if len(vn_amt) != 1 or len(vn_dup) != 1:
    log("!! vn 侧不是唯一一笔 -> 中止（不能删）"); ok = False
if vn_amt and dstr(wsv.cell(vn_amt[0][0], 1).value) != "2026-06-19":
    log("!! vn 那笔不是 2026-06-19 -> 中止"); ok = False

# cn 表内重复凭证体检
seen = {}
for r in range(5, 205):
    b = str(wsc.cell(r, 2).value or "").strip().upper()
    if b:
        seen.setdefault(b, []).append(r)
dups = {k: v for k, v in seen.items() if len(v) > 1}
log("cn 表内重复凭证号 = %s" % (dups or "无"))
if set(dups.get("UNC00101", [])) != {KEEP, DUPPY}:
    log("!! 重复凭证情况和预期不符 -> 中止"); ok = False

if not ok:
    log("")
    log("已中止，未改动任何文件。")
    wbc.close(); wbv.close()
    sys.exit(1)

# ---------- 执行 ----------
before = total(wsc)
log("")
log("=== 修改前：支出合计(三块相加) = %s  （每块 %s）===" % (format(before, ","), format(before // 3, ",")))

ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
bak = os.path.join(BKD, "cn_report_%s_before_dup15fix.xlsx.bak" % ts)
shutil.copy2(CN, bak)
log("已备份 -> %s" % bak)

for c in sum(BLOCKS, ()):
    wsc.cell(DUPPY, c).value = None
wsc.cell(DUPPY, 29).value = None      # 核销列 AC
log("已清空行%d 的 A-E / K-O / U-Y / AC" % DUPPY)

tmp = CN + ".fix.tmp"
wbc.save(tmp)
wbc.close()
wbv.close()
os.replace(tmp, CN)
log("已原子落盘：%s" % CN)

# ---------- 复核 ----------
wb2 = openpyxl.load_workbook(CN, data_only=False)
ws2 = wb2[wb2.sheetnames[0]]
after = total(ws2)
log("")
log("=== 复核 ===")
log("支出合计 改前=%s  改后=%s  差=%s" % (format(before, ","), format(after, ","), format(before - after, ",")))
log("每块合计 = %s" % format(after // 3, ","))
log("行3 合计公式 = %s" % ws2.cell(3, 5).value)
for r in range(5, 205):
    if ws2.cell(r, 2).value is not None or ws2.cell(r, 3).value is not None:
        log("  行%-3d %s | %s | %s | D=%s E=%s | N=%s O=%s | X=%s Y=%s"
            % (r, dstr(ws2.cell(r, 1).value), fmt(ws2.cell(r, 2).value), str(ws2.cell(r, 3).value)[:30],
               fmt(ws2.cell(r, 4).value), fmt(ws2.cell(r, 5).value),
               fmt(ws2.cell(r, 14).value), fmt(ws2.cell(r, 15).value),
               fmt(ws2.cell(r, 24).value), fmt(ws2.cell(r, 25).value)))
log("行15 现在：%s" % [fmt(ws2.cell(15, c).value) for c in (1, 2, 3, 5, 13, 15, 23, 25)])
wb2.close()
log("done")
