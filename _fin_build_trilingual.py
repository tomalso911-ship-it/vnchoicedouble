# -*- coding: utf-8 -*-
"""把 vn_report.xlsx 改造成三语并列结构（方案 B）：
   VI 块 A..(源列) | ZH 块 K.. | EN 块 U..
   - 文字列（A 标题 / C 摘要 / G 备注）调 Google 翻译填充 ZH/EN
   - 数字列（D/E/F）与代号列（B/H）原样复制（语言无关）
   保留原 A..H 不动；写入 K+/U+ 列。带持久术语缓存，可重跑补全。
   同时生成 finance_data/fin_sheet_titles.json（各 sheet 三语标签名）。
"""
import openpyxl, json, os, sys, time
from urllib.parse import quote, urlencode
from urllib.request import Request, urlopen

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "finance_data", "vn_report.xlsx")
CACHE_PATH = os.path.join(BASE, "finance_data", "fin_trans_cache.json")
TITLES_PATH = os.path.join(BASE, "finance_data", "fin_sheet_titles.json")

# 源列 -> 角色：translate=翻译, copy=原样复制
ROLE = {1: "translate", 2: "copy", 3: "translate", 4: "copy",
        5: "copy", 6: "copy", 7: "translate", 8: "copy"}
VI_START, ZH_START, EN_START = 1, 11, 21  # A, K, U


def load_cache():
    if os.path.isfile(CACHE_PATH):
        try:
            return json.load(open(CACHE_PATH, encoding="utf-8"))
        except Exception:
            pass
    return {}


CACHE = load_cache()


def save_cache():
    try:
        json.dump(CACHE, open(CACHE_PATH, "w", encoding="utf-8"),
                  ensure_ascii=False, indent=1)
    except Exception as e:
        print("  [cache] save failed:", e)


# ---- 翻译链（与 app.py 的 fin_translate 一致）：Google(主) -> DeepL(备) -> MyMemory(末) ----
_DEEPL_KEY = os.environ.get("DEEPL_API_KEY") or ""
if not _DEEPL_KEY:
    try:
        _kf = os.path.join(BASE, "finance_data", ".deepl_key")
        if os.path.isfile(_kf):
            _DEEPL_KEY = open(_kf, encoding="utf-8").read().strip()
    except Exception:
        _DEEPL_KEY = ""
_DEEPL_ENDPOINT = "https://api-free.deepl.com/v2/translate"
_DEEPL_LANG = {"zh": "ZH", "en": "EN", "vi": "VI"}


def _g_google(text, target):
    try:
        u = ("https://translate.googleapis.com/translate_a/single?client=gtx&sl=auto&tl="
             + quote(target) + "&dt=t&q=" + quote(text))
        req = Request(u, headers={"User-Agent": "Mozilla/5.0"})
        with urlopen(req, timeout=10) as r:
            data = json.loads(r.read().decode("utf-8"))
        return "".join(seg[0] for seg in (data[0] or []) if seg and seg[0])
    except Exception:
        return ""


def _g_deepl(text, target):
    if not _DEEPL_KEY:
        return ""
    try:
        body = urllib.parse.urlencode({"text": text, "target_lang": _DEEPL_LANG[target]}).encode("utf-8")
        req = Request(_DEEPL_ENDPOINT, data=body, method="POST")
        req.add_header("Authorization", "DeepL-Auth-Key " + _DEEPL_KEY)
        req.add_header("Content-Type", "application/x-www-form-urlencoded")
        req.add_header("User-Agent", "Mozilla/5.0")
        with urlopen(req, timeout=12) as r:
            j = json.loads(r.read().decode("utf-8"))
        return (j.get("translations") or [{}])[0].get("text") or ""
    except Exception:
        return ""


def mtrans(text, target):
    """末备翻译源 MyMemory（免费、无 key、与 Google 不同服务）。
       源语言固定 vi（本工作簿为越南文）。"""
    try:
        u = ("https://api.mymemory.translated.net/get?q=" + quote(text)
             + "&langpair=vi|" + quote(target))
        req = Request(u, headers={"User-Agent": "Mozilla/5.0"})
        with urlopen(req, timeout=12) as r:
            data = json.loads(r.read().decode("utf-8"))
        t = (data.get("responseData") or {}).get("translatedText") or ""
        if "MYMEMORY WARNING" in t or not t:
            return ""
        return t
    except Exception:
        return ""


def gtrans(text, target):
    text = (text or "").strip()
    if not text:
        return ""
    key = text + "|" + target
    if key in CACHE:
        return CACHE[key]
    t = ""
    for fn in (_g_google, _g_deepl, mtrans):
        try:
            t = fn(text, target)
        except Exception:
            t = ""
        if t:
            break
    if t:
        CACHE[key] = t
        time.sleep(0.1)
    return t


def col_letter(n):
    return openpyxl.utils.get_column_letter(n)


def main():
    if not os.path.isfile(SRC):
        print("源文件不存在:", SRC)
        return
    # 连通性预检：不通则不碰文件直接退出
    print("预检翻译网络…")
    probe = gtrans("test", "en")
    if not probe:
        print("✗ 翻译 API 不可达（可能服务器无外网）。未修改文件，请检查网络后重跑。")
        return
    print("✓ 翻译网络可用")

    # 备份
    bak = SRC + ".bak_pre_trilingual"
    if not os.path.isfile(bak):
        import shutil
        shutil.copyfile(SRC, bak)
        print("已备份 ->", bak)

    wb = openpyxl.load_workbook(SRC)
    titles = {}
    total_api = [0]

    for ws in wb.worksheets:
        name = ws.title
        print("=== 处理 sheet:", name, "(max_row=%d)" % ws.max_row)
        # 源列：1..8 中任一格非空的列
        src_cols = [c for c in range(1, 9)
                    if any(ws.cell(r, c).value not in (None, "") for r in range(1, min(ws.max_row, 400) + 1))]
        if not src_cols:
            continue
        zh_off = ZH_START - 1
        en_off = EN_START - 1
        for c in src_cols:
            role = ROLE.get(c, "copy")
            zh_c = c + zh_off
            en_c = c + en_off
            for r in range(1, ws.max_row + 1):
                v = ws.cell(r, c).value
                if v is None or v == "":
                    continue
                if role == "copy" or not isinstance(v, str):
                    ws.cell(r, zh_c).value = v
                    ws.cell(r, en_c).value = v
                else:
                    zh = gtrans(v, "zh") or v   # 失败回退填越南原文，绝不空
                    en = gtrans(v, "en") or v
                    total_api[0] += 2
                    ws.cell(r, zh_c).value = zh
                    ws.cell(r, en_c).value = en
        # sheet 标签名三语
        tv = gtrans(name, "zh"); total_api[0] += 1
        te = gtrans(name, "en"); total_api[0] += 1
        titles[name] = {"vi": name, "zh": tv or name, "en": te or name}
        print("   标签: vi=%s zh=%s en=%s" % (name, titles[name]["zh"], titles[name]["en"]))
        if total_api[0] % 20 == 0:
            time.sleep(0.3)
            save_cache()

    wb.save(SRC)
    save_cache()
    json.dump(titles, open(TITLES_PATH, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("完成。API 调用约 %d 次。缓存存于 %s；标签存于 %s" % (total_api[0], CACHE_PATH, TITLES_PATH))


if __name__ == "__main__":
    main()
