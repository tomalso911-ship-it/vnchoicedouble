# -*- coding: utf-8 -*-
"""确认核销列居中：xlsx 里的对齐 + 服务端快照里带过去的样式（ASCII 输出）。"""
import io
import json
import os
import sys
import urllib.request

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
P = os.path.join(HERE, "finance_data", "cn_report.xlsx")
out = {}

wb = openpyxl.load_workbook(P)
ws = wb[wb.sheetnames[0]]
for a in ("AC4", "AC5", "AC6", "AC204"):
    c = ws[a]
    al = c.alignment
    out[a] = {"v": c.value,
              "h": (al.horizontal if al else None),
              "v_align": (al.vertical if al else None)}
print("XLSX sheets =", json.dumps(wb.sheetnames, ensure_ascii=True))
print("AC align    =", json.dumps(out, ensure_ascii=True))

try:
    with urllib.request.urlopen("http://127.0.0.1:5050/api/finance/cn-report/snapshot", timeout=180) as r:
        snap = json.loads(r.read().decode("utf-8"))
    sid = snap["sheetOrder"][0]
    cd = snap["sheets"][sid]["cellData"]
    style = {}
    for rr in ("3", "4", "5", "6"):
        cell = (cd.get(rr, {}) or {}).get("28")
        style["AC" + str(int(rr) + 1)] = {"v": (cell or {}).get("v"), "s": (cell or {}).get("s")}
    print("SNAPSHOT id =", snap.get("id"))
    print("SNAPSHOT AC =", json.dumps(style, ensure_ascii=True))
except Exception as e:
    print("SNAPSHOT GET ERR:", str(e)[:200])
