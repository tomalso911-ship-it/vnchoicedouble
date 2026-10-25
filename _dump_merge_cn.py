# -*- coding: utf-8 -*-
"""dump：merge 里 10 个 CN: 蓝行的真实值/格式，vs cn 源文件对应行，找出拷贝串列根因。"""
import os
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
OUT = os.path.join(HERE, "_dump_merge_cn_out.txt")
L = []


def log(s=""):
    L.append(str(s))
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    print(s)


def fmtv(v):
    if v is None:
        return "·"
    if isinstance(v, str) and v.startswith("="):
        return "公式:%s" % v[:30]
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


COLS = list(range(1, 6)) + list(range(11, 16)) + list(range(21, 26))  # A-E, K-O, U-Y
COLS = [1, 2, 3, 4, 5, 11, 12, 13, 14, 15, 21, 22, 23, 24, 25]

wb = openpyxl.load_workbook(os.path.join(FIN, "merge_report.xlsx"), data_only=False)
ws = wb["Chi thực tế"]
log("===== merge_report / Chi thực tế 里带 CN: 标记的行 =====")
for r in range(5, ws.max_row + 1):
    mk = ws.cell(r, 40).value
    if isinstance(mk, str) and mk.startswith("CN:"):
        cells = " ".join("%s=%s/%s" % (openpyxl.utils.get_column_letter(c),
                                        fmtv(ws.cell(r, c).value), ws.cell(r, c).number_format)
                         for c in COLS)
        log("行%-3d [%s] %s" % (r, mk, cells))
wb.close()

wbc = openpyxl.load_workbook(os.path.join(FIN, "cn_report.xlsx"), data_only=False)
cws = wbc[wbc.sheetnames[0]]
log("")
log("===== cn_report / 实际支出 里非空的数据行 =====")
for r in range(5, cws.max_row + 1):
    nonempty = any(cws.cell(r, c).value not in (None, "") for c in COLS)
    if not nonempty:
        continue
    cells = " ".join("%s=%s/%s" % (openpyxl.utils.get_column_letter(c),
                                    fmtv(cws.cell(r, c).value), cws.cell(r, c).number_format)
                     for c in COLS)
    log("行%-3d %s" % (r, cells))
wbc.close()
log("done")
