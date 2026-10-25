# -*- coding: utf-8 -*-
"""把【中国财务报表 cn_report.xlsx「实际支出」的每一条记录】按【发生日期的自然顺序】
分散插入到【越南财务报表 vn_report.xlsx 的「Chi thực tế（实际支出）」工作表】里。

规则（用户 2026-10-14 要求）：
  ① 不追加到末尾，而是按日期排进越南记录中间（同日 → 越南在前、中国在后，稳定排序）；
  ② 中国记录整行【蓝字】（与现在中国报表数据区一致，方便区分）；
  ③ 不再用 SUM 公式：表头行 3 与底部「TỔNG CỘNG」行都写成【静态合计值】；
  ④ 幂等：合并行在隐藏列打 CN:<凭证号>:<日期> 标记，重跑不会重复插入（先剔除旧的中国行再重排）；
  ⑤ 写盘前自动备份到 finance_data/backups/，并原子替换（.merge.tmp + os.replace）。

用法：
  python _merge_cn_into_vn.py            # 预演（只看会插入哪些行、插到第几行，不写盘）
  python _merge_cn_into_vn.py --apply    # 真正执行
"""
import datetime
import io
import os
import shutil
import sys
from copy import copy

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402
from openpyxl.styles import Color, Font  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
VN = os.path.join(FIN, "vn_report.xlsx")
CN = os.path.join(FIN, "cn_report.xlsx")
SHEET = "Chi thực tế"

FIRST_DATA_ROW = 5        # 数据区起始行（1-4 是标题/合计/表头）
# ★判断"这一行有没有记录"只看三块的真正数据列：VI(A–E) / ZH(K–O) / EN(U–Y)。
#   绝不能把 AC（核销列，第 29 列）算进去 —— 中国报表给 200 行都预生成了 ☐，
#   否则空行会全被当成"中国记录"（预演实测：200 行里 189 行是空的）。
DATA_COLS = list(range(1, 6)) + list(range(11, 16)) + list(range(21, 26))
COPY_COLS = list(range(1, 29))   # 实际搬运 1..28 列（A..AB），跳过 29(AC 核销) 与之后
MARK_COL = 40                    # 隐藏标记列（AN），只用来做幂等判定，不显示
MARK_PREFIX = "CN:"
BLUE = "FF0000FF"
TOTAL_LABEL = "TỔNG CỘNG"

OUT = os.path.join(HERE, "_merge_cn_into_vn_out.txt")
L = []


def log(s):
    L.append(str(s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(s)


def cell_date_key(v):
    """把单元格值转成可比较的日期键（无日期 → 排在最后）。"""
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, datetime.date):
        return datetime.datetime(v.year, v.month, v.day)
    return None


def row_empty(vals):
    for v in vals:
        if v is not None and str(v).strip() != "":
            return False
    return True


def read_rows(ws, r0, r1, tag):
    rows = []
    for r in range(r0, r1 + 1):
        vals = [ws.cell(r, c).value for c in DATA_COLS]
        if row_empty(vals):
            continue
        # 跳过"合计"行：C 列是 TỔNG CỘNG / 或含 SUM 公式
        joined = " ".join(str(v) for v in vals if v is not None)
        if "TỔNG CỘNG" in joined or any(str(v).startswith("=SUM") for v in vals if isinstance(v, str)):
            continue
        cells = {}
        for c in COPY_COLS:
            src = ws.cell(r, c)
            cells[c] = (src.value, copy(src._style))
        rows.append({
            "tag": tag,
            "src": r,
            "date": cell_date_key(ws.cell(r, 1).value),
            "voucher": str(ws.cell(r, 2).value or "").strip(),
            "cells": cells,
        })
    return rows


def main():
    apply = "--apply" in sys.argv
    log("模式 = %s" % ("★真正写盘" if apply else "预演（不写盘）"))

    wbv = openpyxl.load_workbook(VN, data_only=False)
    if SHEET not in wbv.sheetnames:
        log("✗ 越南报表里找不到工作表 %r（现有：%s）" % (SHEET, wbv.sheetnames))
        return 2
    ws = wbv[SHEET]
    wbc = openpyxl.load_workbook(CN, data_only=False)
    cws = wbc[wbc.sheetnames[0]]

    # ---- 1) 读越南数据行（剔除上一次合并进去的中国行） ----
    vn_rows_all = read_rows(ws, FIRST_DATA_ROW, 400, "vn")
    vn_rows, dropped = [], 0
    for row in vn_rows_all:
        mk = ws.cell(row["src"], MARK_COL).value
        if isinstance(mk, str) and mk.startswith(MARK_PREFIX):
            dropped += 1
            continue
        vn_rows.append(row)
    log("越南记录 = %d 行（其中剔除上次合并的中国行 %d 行）" % (len(vn_rows_all), dropped))

    # ---- 2) 读中国数据行 ----
    cn_rows = read_rows(cws, FIRST_DATA_ROW, 205, "cn")
    log("中国记录 = %d 行" % len(cn_rows))
    for row in cn_rows:
        log("   %s  凭证=%s  金额=%s  %s"
            % (row["date"].strftime("%Y-%m-%d") if row["date"] else "无日期",
               row["voucher"], row["cells"][5][0],
               str(row["cells"][3][0])[:34]))

    # ---- 2.5) 去重：越南表里【已经存在】同一笔（同日期 + 同凭证号 + 同金额）就不再插入 ----
    #     实测：vn_report 的「Chi thực tế」里本来就有 200 多条 UNC… 凭证的中国付款，
    #     直接插会变成同一笔钱记两遍（合计被翻倍）→ 必须按 (日期, 凭证, 金额) 判重。
    # ★用户 2026-10-14 决定：--keep-dup 时【不做判重】，中国报表里每一条都照插
    #   （越南表里已有的同日期+同凭证+同金额的行全部保留 → 合计会多出中国这一份）。
    if "--keep-dup" in sys.argv:
        kept = list(cn_rows)
        log("模式 = --keep-dup：不做判重，中国 %d 条全部照插（合计会重复计入中国这一份）" % len(kept))
    else:
        def key_of(row):
            d = row["date"].strftime("%Y-%m-%d") if row["date"] else ""
            amt = row["cells"][5][0]
            try:
                amt = round(float(amt), 2)
            except (TypeError, ValueError):
                amt = str(amt)
            return (d, row["voucher"].upper(), amt)

        vn_keys = set(key_of(x) for x in vn_rows)
        kept, skipped, seen_cn = [], [], set()
        for row in cn_rows:
            k = key_of(row)
            if k[1] and k in vn_keys:
                skipped.append(row)
                continue
            if k in seen_cn:                 # 中国表内部自身的重复行
                skipped.append(row)
                continue
            seen_cn.add(k)
            kept.append(row)
        log("判重后：中国记录保留 %d 行，跳过 %d 行（越南表里已有同日期+同凭证+同金额，或中国表内重复）"
            % (len(kept), len(skipped)))
        for row in skipped:
            log("   跳过 %s %s %s" % (row["date"].strftime("%Y-%m-%d") if row["date"] else "无日期",
                                      row["voucher"], row["cells"][5][0]))

    # ---- 3) 按日期稳定排序（越南先、中国后 → 同日越南在前） ----
    far = datetime.datetime(2999, 1, 1)
    merged = sorted(vn_rows + kept, key=lambda x: (x["date"] or far))
    n_cn = sum(1 for x in merged if x["tag"] == "cn")
    log("合并后总行数 = %d（越南 %d + 中国 %d）" % (len(merged), len(merged) - n_cn, n_cn))

    # 中国行落在第几行
    pos = []
    for i, row in enumerate(merged):
        if row["tag"] == "cn":
            pos.append(FIRST_DATA_ROW + i)
    log("中国行插入位置（Excel 行号）= %s" % pos)

    if not apply:
        log("")
        log("=== 预演结束：加 --apply 才会写盘 ===")
        return 0

    # ---- 4) 备份 ----
    bak = os.path.join(FIN, "backups",
                       "vn_report_%s_before_merge.xlsx.bak"
                       % datetime.datetime.now().strftime("%Y%m%d_%H%M%S"))
    shutil.copyfile(VN, bak)
    log("已备份 -> %s" % os.path.basename(bak))

    # ---- 4.5) 记住旧的合计行号，并修正其它表的跨表引用 ----
    #     实测：TỔNG HỢP!C21 = 'Chi thực tế'!E384（指向合计格）。
    #     合并后合计行会下移，若不改，汇总表就会读到空/错值 → 改成指向【行 3】的表头合计
    #     （行 3 位置固定，以后再合并也不会失效）。
    old_total_row = None
    for r in range(FIRST_DATA_ROW, 400):
        if str(ws.cell(r, 3).value or "").strip() == TOTAL_LABEL:
            old_total_row = r
            break
    if old_total_row:
        import re as _re
        pat = _re.compile(r"(!\s*\$?)([A-Z]{1,2})\$?" + str(old_total_row) + r"\b")
        fixed = 0
        for sh in wbv.worksheets:
            for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 500),
                                    max_col=min(sh.max_column, 40)):
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.startswith("=") and SHEET in v:
                        nv = pat.sub(lambda m: "%s%s3" % (m.group(1), m.group(2)), v)
                        if nv != v:
                            cell.value = nv
                            fixed += 1
                            log("   修正跨表引用 %s!%s: %s → %s" % (sh.title, cell.coordinate, v, nv))
        log("旧合计行 = 第 %d 行；跨表引用修正 %d 处" % (old_total_row, fixed))

    # ---- 5) 清空数据区（含旧的合计行） ----
    clear_to = FIRST_DATA_ROW + len(merged) + 6
    for r in range(FIRST_DATA_ROW, max(clear_to, 400) + 1):
        for c in COPY_COLS:
            ws.cell(r, c).value = None
        ws.cell(r, MARK_COL).value = None

    # ---- 6) 写回合并后的行 ----
    # ★中国行按列角色显式设定类型与格式（避免 blind 拷贝把金额套成日期格式、日期套成货币格式）
    DATE_COLS = {1, 11, 21}          # A / K / U = 越/中/英 日期列
    AMT_COLS = {4, 5, 14, 15,24, 25} # D/E N/O X/Y = 越/中/英 收支列
    def write_row(r, row):
        is_cn = row["tag"] == "cn"
        for c in COPY_COLS:
            v, st = row["cells"][c]
            dst = ws.cell(r, c)
            if is_cn:
                if c in DATE_COLS:
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        try:                       # 裸序列号 → 真日期
                            v = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=int(v))
                        except Exception:
                            pass
                    dst.value = v
                    dst.number_format = "dd-mm-yyyy"
                elif c in AMT_COLS:
                    if v is not None:
                        try:
                            v = float(v)
                        except (TypeError, ValueError):
                            pass
                    dst.value = v
                    dst.number_format = "#,##0"     # 数值格式，绝不给日期格式
                else:
                    dst.value = v
                    try:
                        dst._style = copy(st)
                    except Exception:
                        pass
                f = copy(dst.font)
                f.color = Color(rgb=BLUE)           # ★中国记录：整行蓝字
                # ★FIN-CN-CHECK-GRAY-20261015（用户 2026-10-15 要求"对合并表没有任何影响"）：
                #   中国报表里「勾了核销的行」在屏幕/文件里是【灰+斜体】；合并表必须完全不受影响，
                #   所以这里连同斜体一起清掉 —— 否则只改颜色会让斜体漏进合并表。
                f.italic = False
                dst.font = f
            else:
                dst.value = v
                dst._style = copy(st)
        if row["tag"] == "cn":
            ws.cell(r, MARK_COL).value = "%s%s:%s" % (
                MARK_PREFIX, row["voucher"],
                row["date"].strftime("%Y-%m-%d") if row["date"] else "")
        else:
            ws.cell(r, MARK_COL).value = None

    for i, row in enumerate(merged):
        write_row(FIRST_DATA_ROW + i, row)

    # 标记列隐藏，界面上看不到
    from openpyxl.utils import get_column_letter  # noqa: E402
    ws.column_dimensions[get_column_letter(MARK_COL)].hidden = True

    # ---- 7) 合计：静态值，不再用 SUM 公式 ----
    last = FIRST_DATA_ROW + len(merged) - 1
    total_row = last + 2
    sums = {}
    for col in (5, 15, 25):          # E / O / Y = 越 / 中 / 英 三块的 Chi
        t = 0.0
        for r in range(FIRST_DATA_ROW, last + 1):
            v = ws.cell(r, col).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                t += v
        sums[col] = t
    # 表头行 3（原来 =SUM(...)）
    for col, val in sums.items():
        ws.cell(3, col).value = val
    # 底部「TỔNG CỘNG」行
    ws.cell(total_row, 3).value = TOTAL_LABEL
    tpl = ws.cell(FIRST_DATA_ROW, 3)      # 用数据行做样式模板
    ws.cell(total_row, 3)._style = copy(tpl._style)
    for col, val in sums.items():
        c = ws.cell(total_row, col)
        c.value = val
        c._style = copy(ws.cell(FIRST_DATA_ROW, col)._style)
    log("合计（静态值，无公式）：E=%s  O=%s  Y=%s  → 写在第 %d 行；行3 同步改为静态值"
        % (sums[5], sums[15], sums[25], total_row))

    # ---- 8) 原子落盘 ----
    tmp = VN + ".merge.tmp"
    wbv.save(tmp)
    os.replace(tmp, VN)
    log("✔ 已写入 %s（原子替换）" % os.path.basename(VN))
    try:
        os.remove(os.path.join(FIN, "vn_report.snapshot.key"))
        log("已清掉快照缓存键 → 重新载入即看到合并结果")
    except Exception:
        pass
    return 0


if __name__ == "__main__":
    sys.exit(main())
