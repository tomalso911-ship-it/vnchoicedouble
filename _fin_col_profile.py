# -*- coding: utf-8 -*-
"""列画像（收窄版）：每 sheet 前 40 列中，有内容的列才打印。"""
import openpyxl, sys
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

CAP = 40
wb = openpyxl.load_workbook('finance_data/vn_report.xlsx', data_only=True)
for ws in wb.worksheets:
    print("=== SHEET:", ws.title, "| max_row:", ws.max_row)
    rows = min(ws.max_row, 300)
    printed = 0
    for c in range(1, CAP + 1):
        text = num = 0
        sample = None
        hdr = ws.cell(1, c).value
        for r in range(1, rows + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            if isinstance(v, (int, float)):
                num += 1
            else:
                text += 1
                if sample is None:
                    sample = str(v)[:24]
        if text == 0 and num == 0:
            continue
        col = openpyxl.utils.get_column_letter(c)
        print("  %s(%d): text=%-4d num=%-4d hdr=%-22s sample=%s" % (col, c, text, num, repr(hdr)[:22] if hdr is not None else 'None', sample))
        printed += 1
    if printed == 0:
        print("  (前 %d 列均无内容)" % CAP)
    print()
