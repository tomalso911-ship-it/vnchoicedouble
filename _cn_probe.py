# -*- coding: utf-8 -*-
"""临时探查：VN 报表 8 张表的结构（重点 'Chi thực tế' 实际支出），结果写入 _cn_probe_out.txt。"""
import io
import os

import openpyxl
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
P = os.path.join(BASE, "finance_data", "vn_report.xlsx")
OUT = os.path.join(BASE, "_cn_probe_out.txt")

wb = openpyxl.load_workbook(P, data_only=False)
L = ["=== sheetnames: %r" % (wb.sheetnames,)]

for nm in wb.sheetnames:
    ws = wb[nm]
    last = 0
    for r in range(1, min(ws.max_row, 12) + 1):
        for c in range(1, ws.max_column + 1):
            if ws.cell(row=r, column=c).value is not None:
                last = max(last, c)
    L.append("")
    L.append("=== SHEET %r  max_row=%d max_col=%d last_nonempty_col(first12rows)=%d merges=%d"
             % (nm, ws.max_row, ws.max_column, last, len(ws.merged_cells.ranges)))
    L.append("    merges: %s" % ([str(m) for m in list(ws.merged_cells.ranges)[:12]],))
    for r in range(1, min(ws.max_row, 6) + 1):
        cells = []
        for c in range(1, min(last, 60) + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                cells.append("%s%d=%r" % (get_column_letter(c), r, v))
        L.append("    R%d: %s" % (r, " | ".join(cells)))

ws = wb["Chi th\u1ec7t t\u1ebf"] if "Chi th\u1ec7t t\u1ebf" in wb.sheetnames else None
if ws is None:
    for nm in wb.sheetnames:
        if nm.startswith("Chi ") and nm != "Chi n\u1ed9i b\u1ed9":
            ws = wb[nm]
L.append("")
L.append("=== TARGET %r  max_row=%d" % (ws.title, ws.max_row))
for r in range(1, min(ws.max_row, 8) + 1):
    cells = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            cells.append("%s%d=%r" % (get_column_letter(c), r, v))
    L.append("  R%d: %s" % (r, " | ".join(cells)))
for r in range(max(1, ws.max_row - 4), ws.max_row + 1):
    cells = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            cells.append("%s%d=%r" % (get_column_letter(c), r, v))
    L.append("  T%d: %s" % (r, " | ".join(cells)))
w = {}
for c in range(1, 40):
    d = ws.column_dimensions.get(get_column_letter(c))
    if d is not None and d.width:
        w[get_column_letter(c)] = round(d.width, 1)
L.append("  col widths: %r" % (w,))
# 数字格式与样式抽样（前 6 行）
fmts = {}
for r in range(1, min(ws.max_row, 6) + 1):
    for c in range(1, 40):
        cell = ws.cell(row=r, column=c)
        if cell.value is not None:
            fmts["%s%d" % (get_column_letter(c), r)] = (cell.number_format, cell.font.b, cell.alignment.horizontal)
L.append("  numfmt sample: %r" % (fmts,))
L.append("  freeze=%r sheet_state=%r" % (ws.freeze_panes, ws.sheet_state))

io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
print("written", OUT)
