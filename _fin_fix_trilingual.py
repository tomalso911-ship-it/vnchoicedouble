# -*- coding: utf-8 -*-
"""补齐 vn_report.xlsx 的 ZH/EN 块里「没翻译 / 机翻回退成越南文」的文字格。

背景：`_fin_build_trilingual.py` 当初按「列角色」决定翻译哪些列（B/D/E/F 当"代号/数字"直接复制），
于是像表头 "Số chứng từ"、"Thu/Chi/Tồn" 这类**文字**表头被原样复制，加上部分长句翻译失败后
`or v` 回退成越南文，导致了「还有很多没翻译的」。

本脚本换一套判定：
  · 任何**字符串**格（数字/日期不动）都要有对应语言的译文；
  · 但「单号/代号」类（含数字且无空格，如 NTTK00001 / UNC00004）不翻译，原样复制；
  · 只改 ZH(K)/EN(U) 两块（列 +10 / +20），**VI 原块一个字节都不动** → VI 公式引用不受影响；
  · 已翻译好的（译文与原文不同且非空）保持不动，只补/重译「缺失 或 仍等于越南原文」的格；
  · 翻译走与 app.py 同一套链路（Google → DeepL → MyMemory）+ 持久缓存 fin_trans_cache.json。

用法：
  python _fin_fix_trilingual.py --dry      # 只统计，不改文件（默认）
  python _fin_fix_trilingual.py --apply    # 真正翻译并写回（会先备份）
"""
import json
import os
import re
import shutil
import sys
import time
from urllib.parse import quote, urlencode
from urllib.request import Request, urlopen

import openpyxl

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "finance_data", "vn_report.xlsx")
# 允许用位置参数指定文件（供上传流水线 fin_trilingual.py 调用 / 自测用；不传则就是线上文件）
_ARGP = [a for a in sys.argv[1:] if not a.startswith("--")]
if _ARGP:
    SRC = os.path.abspath(_ARGP[0])
CACHE_PATH = os.path.join(BASE, "finance_data", "fin_trans_cache.json")
ZH_OFF, EN_OFF = 10, 20
SRC_COLS = range(1, 9)          # A..H（VI 基准块）

_DEEPL_KEY = os.environ.get("DEEPL_API_KEY") or ""
if not _DEEPL_KEY:
    try:
        _kf = os.path.join(BASE, "finance_data", ".deepl_key")
        if os.path.isfile(_kf):
            _DEEPL_KEY = open(_kf, encoding="utf-8").read().strip()
    except Exception:
        _DEEPL_KEY = ""
_DEEPL_ENDPOINT = "https://api-free.deepl.com/v2/translate"
_DEEPL_LANG = {"zh": "ZH", "en": "EN"}


def load_cache():
    if os.path.isfile(CACHE_PATH):
        try:
            return json.load(open(CACHE_PATH, encoding="utf-8"))
        except Exception:
            pass
    return {}


CACHE = load_cache()


def save_cache():
    try:
        json.dump(CACHE, open(CACHE_PATH, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    except Exception as e:
        print("  [cache] save failed:", e)


def _g_google(text, target):
    try:
        u = ("https://translate.googleapis.com/translate_a/single?client=gtx&sl=auto&tl="
             + quote(target) + "&dt=t&q=" + quote(text))
        req = Request(u, headers={"User-Agent": "Mozilla/5.0"})
        with urlopen(req, timeout=10) as r:
            data = json.loads(r.read().decode("utf-8"))
        return "".join(seg[0] for seg in (data[0] or []) if seg and seg[0])
    except Exception:
        return ""


def _g_deepl(text, target):
    if not _DEEPL_KEY:
        return ""
    try:
        body = urlencode({"text": text, "target_lang": _DEEPL_LANG[target]}).encode("utf-8")
        req = Request(_DEEPL_ENDPOINT, data=body, method="POST")
        req.add_header("Authorization", "DeepL-Auth-Key " + _DEEPL_KEY)
        req.add_header("Content-Type", "application/x-www-form-urlencoded")
        req.add_header("User-Agent", "Mozilla/5.0")
        with urlopen(req, timeout=12) as r:
            j = json.loads(r.read().decode("utf-8"))
        return (j.get("translations") or [{}])[0].get("text") or ""
    except Exception:
        return ""


def mtrans(text, target):
    try:
        u = ("https://api.mymemory.translated.net/get?q=" + quote(text)
             + "&langpair=vi|" + quote(target))
        req = Request(u, headers={"User-Agent": "Mozilla/5.0"})
        with urlopen(req, timeout=12) as r:
            data = json.loads(r.read().decode("utf-8"))
        t = (data.get("responseData") or {}).get("translatedText") or ""
        if "MYMEMORY WARNING" in t or not t:
            return ""
        return t
    except Exception:
        return ""


# ★FIN-TOTAL-LABEL-LOCK-20261017（用户 2026-10-17：「全部的，改成：总计（锁定这个中文翻译）」）
# 合计行标签三语写死，与 fin_trilingual.py 的 TR_TOTAL_LOCK 保持一致。
# 谷歌词典把 "TỔNG CỘNG" 译成「全部的」，这条通道也必须先拦，否则补译又写回「全部的」。
TOTAL_LOCK = {"zh": "总计", "en": "TOTAL", "vi": "TỔNG CỘNG"}
TOTAL_SRC = frozenset(x.upper() for x in (
    "TỔNG CỘNG", "Tổng cộng", "Tổng", "TOTAL", "TOTAL CỘNG",
    "全部的", "全部", "合计", "总计", "共計", "合計",
))


def total_lock(text, target):
    t = (text or "").strip()
    if not t:
        return ""
    suf = ""
    if t.endswith(":") or t.endswith("："):
        suf = t[-1]
        t = t[:-1].strip()
    if t.upper() not in TOTAL_SRC:
        return ""
    base = TOTAL_LOCK.get(target)
    return (base + suf) if base else ""


def gtrans(text, target):
    text = (text or "").strip()
    if not text:
        return ""
    key = text + "|" + target
    _lk = total_lock(text, target)
    if _lk:
        CACHE[key] = _lk
        return _lk
    if key in CACHE:
        return CACHE[key]
    t = ""
    for fn in (_g_google, _g_deepl, mtrans):
        try:
            t = fn(text, target)
        except Exception:
            t = ""
        if t:
            break
    if t:
        CACHE[key] = t
        time.sleep(0.05)
    return t


CODE_RE = re.compile(r"^[A-Z0-9][A-Z0-9\-_./]*$")          # 全大写字母数字（单号/代号）
NUMLIKE_RE = re.compile(r"^[\d\s.,%\-+()]+$")               # 纯数字样式

# ---------------------------------------------------------------------------
# ★财务专用术语表（原样译文，优先于机翻；键用小写比对）
# 为什么必须要有：机翻对单字表头会闹笑话 —— "Thu"→"星期四"、"Chi"→"志"、"Tồn"→"存在"，
# 分类码 "chi tt"→"奇特"、"thu nb"→"收藏NB"。这些是报表骨架文字，必须人工定死。
# ---------------------------------------------------------------------------
GLOSSARY = {
    # ---- 表头 ----
    "ngày hạch toán": ("记账日期", "Accounting date"),
    "số chứng từ": ("凭证号", "Document No."),
    "diễn giải": ("摘要", "Description"),
    "thu": ("收入", "Receipts"),
    "chi": ("支出", "Payments"),
    "tồn": ("结余", "Balance"),
    "ghi chú": ("备注", "Note"),
    "stt": ("序号", "No."),
    "description/nội dung": ("摘要/内容", "Description/CONTENT"),
    "amount (vnd)/số tiền": ("金额（VND）", "Amount (VND)"),
    "remarks/ghi chú": ("备注", "Remarks"),
    # ---- 表标题 / 副标题 ----
    "sổ chi tiết thu chi": ("收支明细账簿", "DETAILED CASH BOOK"),
    "sổ chi tiết thu nội bộ": ("内部收入明细账簿", "INTERNAL RECEIPTS DETAIL"),
    "sổ chi tiết thu tiền mr cường": ("何明强先生收款明细账簿", "RECEIPTS DETAIL - MR. CUONG"),
    "sổ chi tiết thu tiền khách hàng": ("客户收款明细账簿", "CUSTOMER RECEIPTS DETAIL"),
    "sổ chi tiết thu khác": ("其他收入明细账簿", "OTHER RECEIPTS DETAIL"),
    "sổ chi tiết chi tiền": ("付款明细账簿", "PAYMENTS DETAIL"),
    "sổ chi tiết chi nội bộ": ("内部支出明细账簿", "INTERNAL PAYMENTS DETAIL"),
    "tiền mặt + tiền gửi ngân hàng": ("现金 + 银行存款", "Cash + Bank deposits"),
    "công ty tnhh choice slats việt nam": ("越南 CHOICE SLATS 有限公司", "CHOICE SLATS VIETNAM CO., LTD"),
    "actual cash receipts and payments report": ("实际现金收支报表", "ACTUAL CASH RECEIPTS AND PAYMENTS REPORT"),
    "ngân hàng vpbank": ("VPbank 银行", "VPbank"),
    # ---- 分类码（H 列）----
    "thu nb": ("内部收入", "Internal receipts"),
    "chi nb": ("内部支出", "Internal payments"),
    "thu tt": ("实际收入", "Actual receipts"),
    "chi tt": ("实际支出", "Actual payments"),
    "thu khác": ("其他收入", "Other receipts"),
    "chi khác": ("其他支出", "Other payments"),
    "thu khách hàng": ("客户收款", "Customer receipts"),
}

# 工作表标签三语（与 finance_data/fin_sheet_titles.json 保持一致）
SHEET_TITLES = {
    "TỔNG HỢP": ("汇总", "SUMMARY"),
    "Chi tiết thu-chi": ("收支明细", "Income & Expenditure Details"),
    "Thu nội bộ": ("内部收入", "Internal Receipts"),
    "Thu thực tế Mr Cường": ("何明强先生实收", "Actual Receipts - Mr. Cường"),
    "Thu Dabaco": ("收款 Dabaco", "Dabaco Receipts"),  # ★DABACO 品牌名不翻译（FIN-NO-TRANS-20261013）
    "Thu khác": ("其他收入", "Other Receipts"),
    "Chi thực tế": ("实际支出", "Actual Payments"),
    "Chi nội bộ": ("内部支出", "Internal Payments"),
}


def gloss(v):
    """命中术语表则返回 (zh, en)，否则 None。"""
    return GLOSSARY.get(str(v).strip().lower())


# ---------------------------------------------------------------------------
# ★译后订正（对人名/固定说法做统一）：机翻会把同一个人名译得前后不一致 ——
#   "Hà Minh Cường" 时而译成"河明强"、时而留拉丁字母 "Ha Minh Cuong"。
#   客户要求中文一律用「何明强」（河内等**地名**不能动，所以只做整词替换）。
# ---------------------------------------------------------------------------
POST_FIX_ZH = [
    ("河明强", "何明强"),
    ("Ha Minh Cuong", "何明强"),
    ("Hà Minh Cường", "何明强"),
]


def post_fix(text, lang):
    if not text or lang != "zh":
        return text
    for a, b in POST_FIX_ZH:
        if a in text:
            text = text.replace(a, b)
    return text


def postfix_sweep(wb, save=True):
    """把译后订正回扫到 ZH/EN 两块（不联网）。"""
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_col=1, max_col=28):
            for cell in row:
                v = cell.value
                if not isinstance(v, str) or not v.strip() or v.strip().startswith("="):
                    continue
                col = cell.column
                lang = "zh" if 11 <= col <= 19 else ("en" if 21 <= col <= 29 else None)
                if lang is None:
                    continue
                nv = post_fix(v, lang)
                if nv != v:
                    cell.value = nv
                    n += 1
    if save:
        wb.save(SRC)
    return n


def is_code(v):
    s = str(v).strip()
    if not s:
        return True
    # ★脏值守卫（FIN-JUNK-GUARD-20261015）：Excel 错误值(#NAME?/#VALUE! 等，
    #   含已被机翻成「＃姓名？」的残渣)一律不翻译 —— 否则中文摘要会变成乱码。
    if re.match(r'^[#＃].*[?!？!]$', s):
        return True
    if NUMLIKE_RE.match(s):
        return True
    if " " in s:
        return False
    # 无空格：既含字母又含数字，或全大写 → 视为单号/代号，不翻译
    if re.search(r"\d", s) and re.search(r"[A-Za-zÀ-ỹ]", s):
        return True
    if CODE_RE.match(s):
        return True
    return False


def needs(cur, v):
    """该格需要（重新）翻译吗？"""
    if cur is None:
        return True
    cur = str(cur).strip()
    if cur == "":
        return True
    return cur == str(v).strip()


def main():
    apply = "--apply" in sys.argv
    if not os.path.isfile(SRC):
        print("源文件不存在:", SRC)
        return

    # 只做"译后订正"回扫（人名统一等），不联网、不重新翻译
    if "--postfix" in sys.argv:
        wb = openpyxl.load_workbook(SRC)
        n = postfix_sweep(wb, save=False)
        if n:
            bak = SRC + ".bak_before_postfix"
            if not os.path.isfile(bak):
                shutil.copyfile(SRC, bak)
                print("已备份 ->", bak)
            wb.save(SRC)
        print("译后订正：改动 %d 格" % n)
        return

    wb = openpyxl.load_workbook(SRC)

    stats = {}
    todo = []          # (sheet, r, target_col, kind, src_text, forced_value)
    for ws in wb.worksheets:
        st = {"cells": 0, "need_zh": 0, "need_en": 0, "code": 0, "gloss": 0}
        for c in SRC_COLS:
            for r in range(1, ws.max_row + 1):
                v = ws.cell(r, c).value
                if v is None or (isinstance(v, str) and v.strip() == ""):
                    continue
                if not isinstance(v, str):
                    continue
                if v.strip().startswith("="):          # 公式字符串不动
                    continue
                st["cells"] += 1
                cz = ws.cell(r, c + ZH_OFF).value
                ce = ws.cell(r, c + EN_OFF).value
                g = gloss(v)
                if g:
                    # ★术语表命中：无条件覆盖（机翻对表头/分类码会错）
                    st["gloss"] += 1
                    if str(cz or "").strip() != g[0]:
                        st["need_zh"] += 1
                        todo.append((ws.title, r, c + ZH_OFF, "zh", v, g[0]))
                    if str(ce or "").strip() != g[1]:
                        st["need_en"] += 1
                        todo.append((ws.title, r, c + EN_OFF, "en", v, g[1]))
                    continue
                if is_code(v):
                    st["code"] += 1
                    continue
                if needs(cz, v):
                    st["need_zh"] += 1
                    todo.append((ws.title, r, c + ZH_OFF, "zh", v, None))
                if needs(ce, v):
                    st["need_en"] += 1
                    todo.append((ws.title, r, c + EN_OFF, "en", v, None))
        stats[ws.title] = st

    print("%-28s %8s %6s %6s %7s %7s" % ("sheet", "strings", "code", "gloss", "needZH", "needEN"))
    for k, st in stats.items():
        print("%-28s %8d %6d %6d %7d %7d" % (k, st["cells"], st["code"], st["gloss"], st["need_zh"], st["need_en"]))
    print("总计需要写入的格数：", len(todo))
    miss = set()
    for (_s, _r, _c, kind, text, forced) in todo:
        if forced is None and (text.strip() + "|" + kind) not in CACHE:
            miss.add((text.strip(), kind))
    print("其中缓存里还没有、必须联网翻译的条数：", len(miss))
    for t, k in list(miss)[:10]:
        print("    -", k, t[:60])

    if not apply:
        print("\n（dry-run，未修改文件。加 --apply 才会真正翻译并写回）")
        return

    bak = SRC + ".bak_before_retrans"
    if not os.path.isfile(bak):
        shutil.copyfile(SRC, bak)
        print("已备份 ->", bak)

    done = 0
    for (sname, r, target_col, kind, text, forced) in todo:
        ws = wb[sname]
        if forced is not None:
            t = forced
        else:
            t = gtrans(text, kind)
            if not t:
                t = text            # 翻译失败：保留原文，绝不写空
        t = post_fix(t, kind)       # ★人名等译后统一
        if ws.cell(r, target_col).value != t:
            ws.cell(r, target_col).value = t
        done += 1
        if done % 50 == 0:
            save_cache()
            print("   ...", done, "/", len(todo))

    save_cache()
    wb.save(SRC)
    print("完成：写入 %d 格 -> %s" % (done, SRC))

    # ---- 工作表标签三语：同时刷新 finance_data/fin_sheet_titles.json ----
    titles = {}
    for name, (zh, en) in SHEET_TITLES.items():
        titles[name] = {"vi": name, "zh": zh, "en": en}
    tp = os.path.join(BASE, "finance_data", "fin_sheet_titles.json")
    json.dump(titles, open(tp, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("已刷新工作表标签三语 ->", tp)


if __name__ == "__main__":
    main()
