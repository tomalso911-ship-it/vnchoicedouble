# -*- coding: utf-8 -*-
"""对比「收入列」在几份文件里的状态，判断 0 是不是原始数据：
   · FinanceReport/Choice Slats VN - Financial Report.xlsx（用户手里的原始报表）
   · finance_data/vn_report.xlsx 的「Chi thực tế」
   · finance_data/cn_report.xlsx
"""
import io
import os
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
ORIG = os.path.join(HERE, "FinanceReport", "Choice Slats VN - Financial Report.xlsx")
OUT = os.path.join(HERE, "_diag_cn_zero3_out.txt")
L = []


def log(s):
    L.append(str(s))
    io.open(OUT, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print(s)


def stat(path, sheet, dcol, ecol, r0=5, r1=210):
    p = path
    tmp = None
    if not path.lower().endswith(".xlsx"):
        tmp = path + ".probe.xlsx"
        shutil.copyfile(path, tmp)
        p = tmp
    try:
        wb = openpyxl.load_workbook(p, data_only=False)
        if sheet not in wb.sheetnames:
            log("   ✗ 没有工作表 %r（现有：%s）" % (sheet, wb.sheetnames))
            wb.close()
            return
        ws = wb[sheet]
        none_d = zero_d = num_d = 0
        none_e = zero_e = num_e = 0
        for r in range(r0, r1 + 1):
            for col, acc in ((dcol, "d"), (ecol, "e")):
                v = ws.cell(r, col).value
                if v is None or str(v).strip() == "":
                    if acc == "d":
                        none_d += 1
                    else:
                        none_e += 1
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    if float(v) == 0.0:
                        if acc == "d":
                            zero_d += 1
                        else:
                            zero_e += 1
                    else:
                        if acc == "d":
                            num_d += 1
                        else:
                            num_e += 1
        log("   %s「%s」 行%d-%d：" % (os.path.basename(path), sheet, r0, r1))
        log("      收入列(第%d列)：空=%d  等于0=%d  有值=%d" % (dcol, none_d, zero_d, num_d))
        log("      支出列(第%d列)：空=%d  等于0=%d  有值=%d" % (ecol, none_e, zero_e, num_e))
        wb.close()
    except Exception as e:
        log("   ✗ %s 读取失败：%s" % (os.path.basename(path), e))
    finally:
        if tmp:
            try:
                os.remove(tmp)
            except OSError:
                pass


log("=== 1) 用户原始报表（FinanceReport/） ===")
if os.path.isfile(ORIG):
    stat(ORIG, "Chi th\u1ef1c t\u1ebf", 4, 5)          # D=Thu 收入, E=Chi 支出
    wb = openpyxl.load_workbook(ORIG, data_only=False)
    log("   该文件全部工作表：%s" % wb.sheetnames)
    wb.close()
else:
    log("   （文件不存在）")

log("")
log("=== 2) finance_data/vn_report.xlsx ===")
stat(os.path.join(FIN, "vn_report.xlsx"), "Chi th\u1ef1c t\u1ebf", 4, 5)

log("")
log("=== 3) finance_data/cn_report.xlsx（唯一的工作表） ===")
wb = openpyxl.load_workbook(os.path.join(FIN, "cn_report.xlsx"), data_only=False)
log("   工作表：%s" % wb.sheetnames)
wb.close()
stat(os.path.join(FIN, "cn_report.xlsx"), "实际支出", 4, 5)
log("done")
