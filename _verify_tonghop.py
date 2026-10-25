# -*- coding: utf-8 -*-
import os, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl
FIN = os.path.join(os.path.dirname(os.path.abspath(__file__)), "finance_data")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_verify_tonghop_out.txt")
L = []
def log(s=""):
    L.append(str(s)); open(OUT,"w",encoding="utf-8").write("\n".join(L)+"\n"); print(s)

wb = openpyxl.load_workbook(os.path.join(FIN,"merge_report.xlsx"), data_only=True)
# TỔNG HỢP 表：找含 TỔNG CỘNG/合计 的行，打印附近金额
ws = wb["TỔNG HỢP"]
log("===== merge / TỔNG HỢP 含'合计'字样行 =====")
for r in range(1, ws.max_row+1):
    a = ws.cell(r,1).value
    if isinstance(a,str) and ("TỔNG" in a or "ồng" in a or "合计" in a or "TOTAL" in a.upper()):
        log("行%-3d A=%r C=%s E=%s" % (r, a, ws.cell(r,3).value, ws.cell(r,5).value))
# 也打印 C21 引用的来源（Chi thực tế!E3）的值
c3 = openpyxl.load_workbook(os.path.join(FIN,"merge_report.xlsx"), data_only=False)["Chi thực tế"]
log("")
log("Chi thực tế!E3 (TỔNG HỢP 引用的静态合计) = %s" % c3.cell(3,5).value)
wb.close()
log("done")
