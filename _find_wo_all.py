# -*- coding: utf-8 -*-
"""在所有行、序号列（A/K/U）里找含「我」的单元格，也顺便全表搜索。"""
import os, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

FIN = os.path.join(os.path.dirname(os.path.abspath(__file__)), "finance_data")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_find_wo_all_out.txt")
L = []
def log(s=""):
    L.append(str(s)); open(OUT,"w",encoding="utf-8").write("\n".join(L)+"\n"); print(s)

for fn, tag in (("vn_report.xlsx","vn"), ("merge_report.xlsx","merge")):
    p = os.path.join(FIN, fn)
    wb = openpyxl.load_workbook(p, data_only=False)
    log("=== %s ===" % tag)
    for sh in wb.worksheets:
        seq_hits = []
        all_hits = []
        for r in range(1, sh.max_row + 1):
            for c in range(1, sh.max_column + 1):
                v = sh.cell(r, c).value
                if isinstance(v, str) and "我" in v:
                    coord = "%s!%s" % (sh.title, sh.cell(r,c).coordinate)
                    all_hits.append((coord, repr(v)))
                    if c in (1, 11, 21):  # A/K/U 序号列
                        seq_hits.append((coord, repr(v)))
        if seq_hits:
            log("[%s] 序号列(A/K/U)含「我」:" % sh.title)
            for coord, rv in seq_hits:
                log("   %s -> %s" % (coord, rv))
        if all_hits and not seq_hits:
            log("[%s] 全表含「我」但不在序号列（共%d处）:" % (sh.title, len(all_hits)))
            for coord, rv in all_hits[:5]:
                log("   %s -> %s" % (coord, rv))
    wb.close()
log("done")
