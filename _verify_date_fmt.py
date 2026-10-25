# -*- coding: utf-8 -*-
"""逐格对比「改日期格式前的备份」与「当前文件」：
   除了日期列，其它任何格子都不允许有任何差异。"""
import datetime
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
FIN = os.path.join(HERE, "finance_data")
BKD = os.path.join(FIN, "backups")
OUT = os.path.join(HERE, "_verify_date_fmt_out.txt")
L = []
PAIRS = [
    ("vn_report.xlsx", "vn_report_20261014_094954_before_datefmt.xlsx.bak", (1, 11, 21)),
    ("cn_report.xlsx", "cn_report_20261014_094956_before_datefmt.xlsx.bak", (1, 11, 21)),
]


def log(s=""):
    L.append(str(s))
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    print(s)


def norm(v):
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v):
        return int(v)
    return v


import shutil
import tempfile


def as_xlsx(path):
    """openpyxl 不认 .bak 扩展名，复制成临时 .xlsx 再读。"""
    tmp = os.path.join(tempfile.gettempdir(), os.path.basename(path) + ".verify.xlsx")
    shutil.copy2(path, tmp)
    return tmp


for cur_name, bak_name, datecols in PAIRS:
    cur = os.path.join(FIN, cur_name)
    bak = as_xlsx(os.path.join(BKD, bak_name))
    log("=" * 70)
    log("%s   vs   %s" % (cur_name, bak_name))
    log("=" * 70)
    wb1 = openpyxl.load_workbook(cur, data_only=False)
    wb2 = openpyxl.load_workbook(bak, data_only=False)
    names = [n for n in wb1.sheetnames if n in wb2.sheetnames]
    if set(names) != set(wb1.sheetnames) | set(wb2.sheetnames):
        log("!! 工作表集合不同：%s vs %s" % (wb1.sheetnames, wb2.sheetnames))
    total_diff = 0
    bad = []
    date_diff = 0
    fmt_diff = 0
    for sn in names:
        ws1, ws2 = wb1[sn], wb2[sn]
        for r in range(1, 801):
            for c in range(1, 61):
                v1, v2 = ws1.cell(r, c).value, ws2.cell(r, c).value
                n1, n2 = norm(v1), norm(v2)
                f1 = (ws1.cell(r, c).number_format or "").strip().lower()
                f2 = (ws2.cell(r, c).number_format or "").strip().lower()
                if n1 != n2 or f1 != f2:
                    total_diff += 1
                    is_date_col = c in datecols
                    if n1 != n2:
                        if is_date_col:
                            date_diff += 1
                        else:
                            bad.append((sn, r, c, n2, n1))
                    else:
                        if is_date_col and f1 == "dd-mm-yyyy":
                            fmt_diff += 1          # 日期列的格式统一（含空格套格式）
                        else:
                            bad.append((sn, r, c, "格式 %r -> %r" % (f2, f1), ""))
    log("差异格总数 = %d" % total_diff)
    log("  其中日期列的值变化（文本日期转真日期）= %d" % date_diff)
    log("  日期列的格式变化 = %d" % fmt_diff)
    log("  ★非日期列/非格式的意外差异 = %d %s" % (len(bad), "✓" if not bad else "!! 必须为 0"))
    for b in bad[:25]:
        log("     [%s] 行%d 列%d  %r -> %r" % b)
    wb1.close()
    wb2.close()
    log("")

# 用户截图那一格的前后
log("=" * 70)
log("截图相关格子的前后值（vn_report.xlsx）")
log("=" * 70)
cur = openpyxl.load_workbook(os.path.join(FIN, "vn_report.xlsx"), data_only=False)
bak = openpyxl.load_workbook(as_xlsx(os.path.join(BKD, "vn_report_20261014_094954_before_datefmt.xlsx.bak")),
                             data_only=False)
for sn, rows in (("Chi th\u1ef1c t\u1ebf", [212, 213]), ("Chi ti\u1ebft thu-chi", [213, 294, 295])):
    for r in rows:
        for c in (1, 11, 21):
            cl = openpyxl.utils.get_column_letter(c)
            log("  [%s] %s%d  改前=%r(%s)   改后=%r(%s)"
                % (sn, cl, r,
                   bak[sn].cell(r, c).value, (bak[sn].cell(r, c).number_format or "").strip(),
                   cur[sn].cell(r, c).value, (cur[sn].cell(r, c).number_format or "").strip()))
cur.close()
bak.close()
log("done")
